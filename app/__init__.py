import os
import csv
from io import StringIO, BytesIO
from datetime import datetime, timedelta, date

from dotenv import load_dotenv
from flask import (
    Flask,
    render_template,
    request,
    redirect,
    url_for,
    session,
    jsonify,
    send_file
)
from flask_sqlalchemy import SQLAlchemy
from sqlalchemy import or_
from werkzeug.security import generate_password_hash, check_password_hash

try:
    from openpyxl import load_workbook, Workbook
except ImportError:
    load_workbook = None
    Workbook = None


# --------------------------------------------------
# LOAD ENVIRONMENT VARIABLES
# --------------------------------------------------

load_dotenv()


# --------------------------------------------------
# DATABASE
# --------------------------------------------------

db = SQLAlchemy()


class WellnessScan(db.Model):

    __tablename__ = "wellness_scans"

    id = db.Column(
        db.Integer,
        primary_key=True
    )

    personnel_id = db.Column(
        db.String(50),
        nullable=False
    )

    heart_rate = db.Column(
        db.Integer,
        nullable=True
    )

    breathing_rate = db.Column(
        db.Integer,
        nullable=True
    )

    signal_quality = db.Column(
        db.String(30),
        nullable=True
    )

    confidence = db.Column(
        db.String(20),
        nullable=True
    )

    measurement_time = db.Column(
        db.String(20),
        nullable=True
    )

    scan_duration = db.Column(
        db.String(30),
        nullable=True
    )

    saved_at = db.Column(
        db.DateTime,
        server_default=db.func.current_timestamp()
    )


class WellnessCheckIn(db.Model):

    __tablename__ = "wellness_checkins"

    id = db.Column(
        db.Integer,
        primary_key=True
    )

    personnel_id = db.Column(
        db.String(50),
        nullable=False,
        index=True
    )

    workload_pressure = db.Column(
        db.String(30),
        nullable=False
    )

    recovery_time = db.Column(
        db.String(30),
        nullable=False
    )

    duty_adjustment = db.Column(
        db.String(30),
        nullable=False
    )

    rest_quality = db.Column(
        db.String(30),
        nullable=False
    )

    low_energy = db.Column(
        db.String(30),
        nullable=False
    )

    family_connection = db.Column(
        db.String(30),
        nullable=False
    )

    family_separation = db.Column(
        db.String(30),
        nullable=False
    )

    focus = db.Column(
        db.String(30),
        nullable=False
    )

    mental_tiredness = db.Column(
        db.String(30),
        nullable=False
    )

    pattern_change = db.Column(
        db.String(30),
        nullable=False
    )

    support_request = db.Column(
        db.String(30),
        nullable=False,
        default="Not Answered"
    )

    optional_note = db.Column(
        db.Text,
        nullable=True
    )

    concern_score = db.Column(
        db.Integer,
        nullable=False,
        default=0
    )

    concern_count = db.Column(
        db.Integer,
        nullable=False,
        default=0
    )

    high_concern_count = db.Column(
        db.Integer,
        nullable=False,
        default=0
    )

    summary_label = db.Column(
        db.String(80),
        nullable=False,
        default="Few self-reported concerns"
    )

    submitted_at = db.Column(
        db.DateTime,
        nullable=False,
        default=datetime.now,
        index=True
    )


class SupportRequest(db.Model):

    __tablename__ = "support_requests"

    id = db.Column(
        db.Integer,
        primary_key=True
    )

    personnel_id = db.Column(
        db.String(50),
        nullable=False
    )

    support_category = db.Column(
        db.String(80),
        nullable=False
    )

    contact_method = db.Column(
        db.String(50),
        nullable=False
    )

    priority = db.Column(
        db.String(30),
        nullable=False
    )

    message = db.Column(
        db.Text,
        nullable=True
    )

    status = db.Column(
        db.String(30),
        nullable=False,
        default="Open"
    )

    created_at = db.Column(
        db.DateTime,
        nullable=False,
        default=datetime.now
    )

    updated_at = db.Column(
        db.DateTime,
        nullable=False,
        default=datetime.now,
        onupdate=datetime.now
    )




class WelfareOfficer(db.Model):

    __tablename__ = "welfare_officers"

    id = db.Column(
        db.Integer,
        primary_key=True
    )

    officer_id = db.Column(
        db.String(80),
        nullable=False,
        unique=True,
        index=True
    )

    full_name = db.Column(
        db.String(120),
        nullable=False,
        default="Welfare Officer"
    )

    designation = db.Column(
        db.String(120),
        nullable=True
    )

    assigned_unit = db.Column(
        db.String(120),
        nullable=True
    )

    email = db.Column(
        db.String(160),
        nullable=True
    )

    password_hash = db.Column(
        db.String(255),
        nullable=False
    )

    account_status = db.Column(
        db.String(30),
        nullable=False,
        default="Active",
        index=True
    )

    last_login_at = db.Column(
        db.DateTime,
        nullable=True
    )

    created_at = db.Column(
        db.DateTime,
        nullable=False,
        default=datetime.now
    )

    updated_at = db.Column(
        db.DateTime,
        nullable=False,
        default=datetime.now,
        onupdate=datetime.now
    )


class WelfareOfficerPreference(db.Model):

    __tablename__ = "welfare_officer_preferences"

    id = db.Column(
        db.Integer,
        primary_key=True
    )

    officer_id = db.Column(
        db.String(80),
        nullable=False,
        unique=True,
        index=True
    )

    notify_welfare_alerts = db.Column(
        db.Boolean,
        nullable=False,
        default=True
    )

    notify_support_requests = db.Column(
        db.Boolean,
        nullable=False,
        default=True
    )

    notify_case_updates = db.Column(
        db.Boolean,
        nullable=False,
        default=True
    )

    updated_at = db.Column(
        db.DateTime,
        nullable=False,
        default=datetime.now,
        onupdate=datetime.now
    )


class Personnel(db.Model):

    __tablename__ = "personnel"

    id = db.Column(
        db.Integer,
        primary_key=True
    )

    personnel_id = db.Column(
        db.String(50),
        nullable=False,
        unique=True,
        index=True
    )

    full_name = db.Column(
        db.String(120),
        nullable=False
    )

    rank_designation = db.Column(
        db.String(100),
        nullable=True
    )

    unit = db.Column(
        db.String(100),
        nullable=True
    )

    department = db.Column(
        db.String(100),
        nullable=True
    )

    posting_location = db.Column(
        db.String(120),
        nullable=True
    )

    joining_date = db.Column(
        db.Date,
        nullable=True
    )

    email = db.Column(
        db.String(160),
        nullable=True
    )

    phone = db.Column(
        db.String(30),
        nullable=True
    )

    password_hash = db.Column(
        db.String(255),
        nullable=False
    )

    employment_status = db.Column(
        db.String(30),
        nullable=False,
        default="Active",
        index=True
    )

    created_by = db.Column(
        db.String(80),
        nullable=True
    )

    updated_by = db.Column(
        db.String(80),
        nullable=True
    )

    created_at = db.Column(
        db.DateTime,
        nullable=False,
        default=datetime.now
    )

    updated_at = db.Column(
        db.DateTime,
        nullable=False,
        default=datetime.now,
        onupdate=datetime.now
    )


class LeaveRequest(db.Model):

    __tablename__ = "leave_requests"

    id = db.Column(
        db.Integer,
        primary_key=True
    )

    request_code = db.Column(
        db.String(40),
        nullable=True,
        unique=True,
        index=True
    )

    personnel_id = db.Column(
        db.String(50),
        nullable=False,
        index=True
    )

    leave_type = db.Column(
        db.String(50),
        nullable=False
    )

    start_date = db.Column(
        db.Date,
        nullable=False,
        index=True
    )

    end_date = db.Column(
        db.Date,
        nullable=False,
        index=True
    )

    requested_days = db.Column(
        db.Integer,
        nullable=False
    )

    reason_category = db.Column(
        db.String(80),
        nullable=False
    )

    reason = db.Column(
        db.Text,
        nullable=False
    )

    additional_note = db.Column(
        db.Text,
        nullable=True
    )

    status = db.Column(
        db.String(20),
        nullable=False,
        default="Pending",
        index=True
    )

    decision_reason = db.Column(
        db.String(160),
        nullable=True
    )

    decision_note = db.Column(
        db.Text,
        nullable=True
    )

    decided_by = db.Column(
        db.String(80),
        nullable=True
    )

    decided_at = db.Column(
        db.DateTime,
        nullable=True
    )

    cancelled_at = db.Column(
        db.DateTime,
        nullable=True
    )

    created_at = db.Column(
        db.DateTime,
        nullable=False,
        default=datetime.now,
        index=True
    )

    updated_at = db.Column(
        db.DateTime,
        nullable=False,
        default=datetime.now,
        onupdate=datetime.now
    )


class OperationalRecord(db.Model):

    __tablename__ = "operational_records"

    id = db.Column(
        db.Integer,
        primary_key=True
    )

    personnel_id = db.Column(
        db.String(50),
        nullable=False,
        index=True
    )

    record_date = db.Column(
        db.Date,
        nullable=False,
        index=True
    )

    duty_hours = db.Column(
        db.Float,
        nullable=True
    )

    rest_hours = db.Column(
        db.Float,
        nullable=True
    )

    context_type = db.Column(
        db.String(40),
        nullable=False,
        default="Station Duty"
    )

    deployment_active = db.Column(
        db.Boolean,
        nullable=False,
        default=False
    )

    training_hours = db.Column(
        db.Float,
        nullable=True
    )

    leave_status = db.Column(
        db.String(40),
        nullable=True
    )

    source_type = db.Column(
        db.String(40),
        nullable=False,
        default="Prototype Import"
    )

    source_record_id = db.Column(
        db.String(80),
        nullable=True
    )

    is_valid = db.Column(
        db.Boolean,
        nullable=False,
        default=True
    )

    created_at = db.Column(
        db.DateTime,
        nullable=False,
        default=datetime.now
    )


class AuditEvent(db.Model):

    __tablename__ = "audit_events"

    id = db.Column(
        db.Integer,
        primary_key=True
    )

    actor_role = db.Column(
        db.String(40),
        nullable=False,
        index=True
    )

    actor_id = db.Column(
        db.String(80),
        nullable=True,
        index=True
    )

    event_type = db.Column(
        db.String(80),
        nullable=False,
        index=True
    )

    personnel_id = db.Column(
        db.String(50),
        nullable=True,
        index=True
    )

    resource_type = db.Column(
        db.String(80),
        nullable=True
    )

    resource_id = db.Column(
        db.String(80),
        nullable=True
    )

    event_summary = db.Column(
        db.Text,
        nullable=False
    )

    source_type = db.Column(
        db.String(60),
        nullable=False,
        default="VeerPulse Application"
    )

    created_at = db.Column(
        db.DateTime,
        nullable=False,
        default=datetime.now,
        index=True
    )


class AssessmentEvidenceSnapshot(db.Model):

    __tablename__ = "assessment_evidence_snapshots"

    id = db.Column(
        db.Integer,
        primary_key=True
    )

    assessment_id = db.Column(
        db.Integer,
        nullable=False,
        index=True
    )

    personnel_id = db.Column(
        db.String(50),
        nullable=False,
        index=True
    )

    assessment_version = db.Column(
        db.Integer,
        nullable=False
    )

    rules_version = db.Column(
        db.String(40),
        nullable=False
    )

    baseline_version = db.Column(
        db.String(40),
        nullable=True
    )

    baseline_context = db.Column(
        db.String(40),
        nullable=True
    )

    evidence_status = db.Column(
        db.String(60),
        nullable=False
    )

    current_record_count = db.Column(
        db.Integer,
        nullable=True
    )

    baseline_record_count = db.Column(
        db.Integer,
        nullable=True
    )

    duty_delta = db.Column(
        db.Float,
        nullable=True
    )

    rest_delta = db.Column(
        db.Float,
        nullable=True
    )

    recovery_day_delta = db.Column(
        db.Integer,
        nullable=True
    )

    concerning_factor_count = db.Column(
        db.Integer,
        nullable=True
    )

    severity_score = db.Column(
        db.Integer,
        nullable=True
    )

    resulting_attention_level = db.Column(
        db.String(30),
        nullable=False
    )

    explanation_summary = db.Column(
        db.Text,
        nullable=True
    )

    created_at = db.Column(
        db.DateTime,
        nullable=False,
        default=datetime.now
    )


class AssessmentCheckInEvidence(db.Model):

    __tablename__ = "assessment_checkin_evidence"

    id = db.Column(
        db.Integer,
        primary_key=True
    )

    assessment_id = db.Column(
        db.Integer,
        nullable=False,
        index=True
    )

    personnel_id = db.Column(
        db.String(50),
        nullable=False,
        index=True
    )

    assessment_version = db.Column(
        db.Integer,
        nullable=False
    )

    checkin_id = db.Column(
        db.Integer,
        nullable=False,
        index=True
    )

    checkin_submitted_at = db.Column(
        db.DateTime,
        nullable=False
    )

    concern_score = db.Column(
        db.Integer,
        nullable=False
    )

    concern_count = db.Column(
        db.Integer,
        nullable=False
    )

    high_concern_count = db.Column(
        db.Integer,
        nullable=False
    )

    support_requested = db.Column(
        db.Boolean,
        nullable=False,
        default=False
    )

    summary_label = db.Column(
        db.String(80),
        nullable=False
    )

    created_at = db.Column(
        db.DateTime,
        nullable=False,
        default=datetime.now
    )


class AssessmentScanEvidence(db.Model):

    __tablename__ = "assessment_scan_evidence"

    id = db.Column(
        db.Integer,
        primary_key=True
    )

    assessment_id = db.Column(
        db.Integer,
        nullable=False,
        index=True
    )

    personnel_id = db.Column(
        db.String(50),
        nullable=False,
        index=True
    )

    assessment_version = db.Column(
        db.Integer,
        nullable=False
    )

    scan_id = db.Column(
        db.Integer,
        nullable=False,
        index=True
    )

    scan_saved_at = db.Column(
        db.DateTime,
        nullable=False
    )

    heart_rate = db.Column(
        db.Integer,
        nullable=True
    )

    breathing_rate = db.Column(
        db.Integer,
        nullable=True
    )

    signal_quality = db.Column(
        db.String(30),
        nullable=True
    )

    confidence = db.Column(
        db.String(20),
        nullable=True
    )

    created_at = db.Column(
        db.DateTime,
        nullable=False,
        default=datetime.now
    )


class WelfareAnalysisResult(db.Model):

    __tablename__ = "welfare_analysis_results"

    id = db.Column(
        db.Integer,
        primary_key=True
    )

    personnel_id = db.Column(
        db.String(50),
        nullable=False,
        index=True
    )

    assessment_id = db.Column(
        db.Integer,
        nullable=False,
        unique=True,
        index=True
    )

    assessment_version = db.Column(
        db.Integer,
        nullable=False,
        index=True
    )

    # --------------------------------------------------
    # CORE SCORES
    # --------------------------------------------------

    operational_score = db.Column(
        db.Float,
        nullable=True
    )

    checkin_score = db.Column(
        db.Float,
        nullable=True
    )

    trend_score = db.Column(
        db.Float,
        nullable=True
    )

    overall_score = db.Column(
        db.Float,
        nullable=True
    )

    # --------------------------------------------------
    # DOMAIN SCORES
    # --------------------------------------------------

    workload_score = db.Column(
        db.Float,
        nullable=True
    )

    recovery_score = db.Column(
        db.Float,
        nullable=True
    )

    connection_score = db.Column(
        db.Float,
        nullable=True
    )

    cognitive_score = db.Column(
        db.Float,
        nullable=True
    )

    # --------------------------------------------------
    # ANALYSIS OUTPUT
    # --------------------------------------------------

    attention_level = db.Column(
        db.String(30),
        nullable=False,
        default="Insufficient Data",
        index=True
    )

    trend_direction = db.Column(
        db.String(40),
        nullable=True
    )

    trend_strength = db.Column(
        db.String(30),
        nullable=True
    )

    signal_agreement = db.Column(
        db.String(30),
        nullable=True
    )

    confidence = db.Column(
        db.String(30),
        nullable=True
    )

    analysis_reason = db.Column(
        db.Text,
        nullable=True
    )

    recommendation = db.Column(
        db.Text,
        nullable=True
    )

    # --------------------------------------------------
    # TRACEABILITY
    # --------------------------------------------------

    rules_version = db.Column(
        db.String(40),
        nullable=False,
        default="VP-BRAIN-1.0"
    )

    operational_snapshot_id = db.Column(
        db.Integer,
        nullable=True,
        index=True
    )

    checkin_evidence_id = db.Column(
        db.Integer,
        nullable=True,
        index=True
    )

    scan_evidence_id = db.Column(
        db.Integer,
        nullable=True,
        index=True
    )

    is_current = db.Column(
        db.Boolean,
        nullable=False,
        default=True,
        index=True
    )

    created_at = db.Column(
        db.DateTime,
        nullable=False,
        default=datetime.now,
        index=True
    )

    updated_at = db.Column(
        db.DateTime,
        nullable=False,
        default=datetime.now,
        onupdate=datetime.now
    )


class AnalysisFactor(db.Model):

    __tablename__ = "analysis_factors"

    id = db.Column(
        db.Integer,
        primary_key=True
    )

    analysis_result_id = db.Column(
        db.Integer,
        nullable=False,
        index=True
    )

    assessment_id = db.Column(
        db.Integer,
        nullable=False,
        index=True
    )

    personnel_id = db.Column(
        db.String(50),
        nullable=False,
        index=True
    )

    factor_key = db.Column(
        db.String(80),
        nullable=False
    )

    factor_label = db.Column(
        db.String(120),
        nullable=False
    )

    source_type = db.Column(
        db.String(50),
        nullable=False
    )

    current_value = db.Column(
        db.Float,
        nullable=True
    )

    baseline_value = db.Column(
        db.Float,
        nullable=True
    )

    deviation_pct = db.Column(
        db.Float,
        nullable=True
    )

    factor_score = db.Column(
        db.Float,
        nullable=True
    )

    severity = db.Column(
        db.String(30),
        nullable=True
    )

    direction = db.Column(
        db.String(30),
        nullable=True
    )

    explanation = db.Column(
        db.Text,
        nullable=True
    )

    created_at = db.Column(
        db.DateTime,
        nullable=False,
        default=datetime.now
    )

    __table_args__ = (
        db.UniqueConstraint(
            "analysis_result_id",
            "factor_key",
            name="uq_analysis_result_factor"
        ),
    )


class WelfareAction(db.Model):

    __tablename__ = "welfare_actions"

    id = db.Column(
        db.Integer,
        primary_key=True
    )

    personnel_id = db.Column(
        db.String(50),
        nullable=False,
        index=True
    )

    action_type = db.Column(
        db.String(80),
        nullable=False
    )

    source_assessment_id = db.Column(
        db.Integer,
        nullable=True
    )

    source_assessment_version = db.Column(
        db.Integer,
        nullable=True
    )

    planner_option_id = db.Column(
        db.String(10),
        nullable=True
    )

    status = db.Column(
        db.String(40),
        nullable=False,
        default="Planned"
    )

    officer_note = db.Column(
        db.Text,
        nullable=True
    )

    approved_at = db.Column(
        db.DateTime,
        nullable=True
    )

    delivered_at = db.Column(
        db.DateTime,
        nullable=True
    )

    interrupted_at = db.Column(
        db.DateTime,
        nullable=True
    )

    verification_note = db.Column(
        db.Text,
        nullable=True
    )

    created_at = db.Column(
        db.DateTime,
        nullable=False,
        default=datetime.now
    )

    updated_at = db.Column(
        db.DateTime,
        nullable=False,
        default=datetime.now,
        onupdate=datetime.now
    )


class RecoveryRosterSlot(db.Model):

    __tablename__ = "recovery_roster_slots"

    id = db.Column(
        db.Integer,
        primary_key=True
    )

    personnel_id = db.Column(
        db.String(50),
        nullable=False,
        index=True
    )

    start_time = db.Column(
        db.DateTime,
        nullable=False,
        index=True
    )

    end_time = db.Column(
        db.DateTime,
        nullable=False,
        index=True
    )

    commitment_type = db.Column(
        db.String(40),
        nullable=False
    )

    mandatory = db.Column(
        db.Boolean,
        nullable=False,
        default=False
    )

    reassignable = db.Column(
        db.Boolean,
        nullable=False,
        default=False
    )

    constraint_type = db.Column(
        db.String(60),
        nullable=True
    )

    constraint_note = db.Column(
        db.String(200),
        nullable=True
    )

    source_type = db.Column(
        db.String(50),
        nullable=False,
        default="Prototype Demonstration Data"
    )

    is_valid = db.Column(
        db.Boolean,
        nullable=False,
        default=True
    )

    created_at = db.Column(
        db.DateTime,
        nullable=False,
        default=datetime.now
    )


class WelfareAssessment(db.Model):

    __tablename__ = "welfare_assessments"

    id = db.Column(
        db.Integer,
        primary_key=True
    )

    personnel_id = db.Column(
        db.String(50),
        nullable=False,
        index=True
    )

    assessment_version = db.Column(
        db.Integer,
        nullable=False,
        default=1
    )

    attention_level = db.Column(
        db.String(30),
        nullable=False,
        default="Insufficient Data"
    )

    evidence_status = db.Column(
        db.String(40),
        nullable=False,
        default="Insufficient Data"
    )

    confidence = db.Column(
        db.String(20),
        nullable=True
    )

    baseline_context = db.Column(
        db.String(40),
        nullable=True
    )

    baseline_version = db.Column(
        db.String(40),
        nullable=True
    )

    rules_version = db.Column(
        db.String(40),
        nullable=False,
        default="VP-RULES-1.0"
    )

    explanation_summary = db.Column(
        db.Text,
        nullable=True
    )

    is_current = db.Column(
        db.Boolean,
        nullable=False,
        default=True
    )

    assessment_time = db.Column(
        db.DateTime,
        nullable=False,
        default=datetime.now
    )

    created_at = db.Column(
        db.DateTime,
        nullable=False,
        default=datetime.now
    )


class PersonnelNotificationRead(db.Model):

    __tablename__ = "personnel_notification_reads"

    id = db.Column(
        db.Integer,
        primary_key=True
    )

    personnel_id = db.Column(
        db.String(50),
        nullable=False,
        index=True
    )

    event_key = db.Column(
        db.String(180),
        nullable=False,
        index=True
    )

    read_at = db.Column(
        db.DateTime,
        nullable=False,
        default=datetime.now
    )

    __table_args__ = (
        db.UniqueConstraint(
            "personnel_id",
            "event_key",
            name="uq_personnel_notification_read"
        ),
    )


class OfficerNotificationRead(db.Model):

    __tablename__ = "officer_notification_reads"

    id = db.Column(
        db.Integer,
        primary_key=True
    )

    officer_id = db.Column(
        db.String(80),
        nullable=False,
        index=True
    )

    event_key = db.Column(
        db.String(180),
        nullable=False,
        index=True
    )

    read_at = db.Column(
        db.DateTime,
        nullable=False,
        default=datetime.now
    )

    __table_args__ = (
        db.UniqueConstraint(
            "officer_id",
            "event_key",
            name="uq_officer_notification_read"
        ),
    )


# --------------------------------------------------
# VOLUNTARY WELLNESS CHECK-IN HELPERS
# --------------------------------------------------

_CHECKIN_ALLOWED_RESPONSES = {
    "Not at all",
    "Slightly",
    "Moderately",
    "Very much"
}

_CHECKIN_NEGATIVE_FIELDS = [
    "workload_pressure",
    "duty_adjustment",
    "low_energy",
    "family_separation",
    "mental_tiredness",
    "pattern_change"
]

_CHECKIN_POSITIVE_FIELDS = [
    "recovery_time",
    "rest_quality",
    "family_connection",
    "focus"
]

_CHECKIN_RESPONSE_SCORE = {
    "Not at all": 0,
    "Slightly": 1,
    "Moderately": 2,
    "Very much": 3
}


def _score_wellness_checkin(payload):

    item_scores = []

    for field_name in _CHECKIN_NEGATIVE_FIELDS:

        item_scores.append(
            _CHECKIN_RESPONSE_SCORE[
                payload[field_name]
            ]
        )

    for field_name in _CHECKIN_POSITIVE_FIELDS:

        raw_score = (
            _CHECKIN_RESPONSE_SCORE[
                payload[field_name]
            ]
        )

        item_scores.append(
            3 - raw_score
        )

    concern_score = sum(
        item_scores
    )

    concern_count = sum(
        1
        for score in item_scores
        if score >= 2
    )

    high_concern_count = sum(
        1
        for score in item_scores
        if score >= 3
    )

    if concern_count >= 4:

        summary_label = (
            "Multiple self-reported concerns"
        )

    elif concern_count >= 2:

        summary_label = (
            "Several self-reported concerns"
        )

    else:

        summary_label = (
            "Few self-reported concerns"
        )

    return {
        "concern_score":
            concern_score,
        "concern_count":
            concern_count,
        "high_concern_count":
            high_concern_count,
        "summary_label":
            summary_label
    }


def _serialize_wellness_checkin(item):

    return {
        "id":
            item.id,

        "personnel_id":
            item.personnel_id,

        "summary_label":
            item.summary_label,

        "concern_score":
            item.concern_score,

        "concern_count":
            item.concern_count,

        "high_concern_count":
            item.high_concern_count,

        "support_request":
            item.support_request,

        "submitted_at":
            (
                item.submitted_at.isoformat()
                if item.submitted_at
                else None
            )
    }


def _recent_wellness_checkin(
    personnel_id,
    days=7
):

    cutoff = (
        datetime.now()
        - timedelta(days=days)
    )

    return (
        WellnessCheckIn.query
        .filter(
            WellnessCheckIn.personnel_id
            == personnel_id,

            WellnessCheckIn.submitted_at
            >= cutoff
        )
        .order_by(
            WellnessCheckIn.submitted_at.desc()
        )
        .first()
    )



def _recent_usable_wellness_scan(
    personnel_id,
    hours=24
):

    cutoff = (
        datetime.now()
        - timedelta(hours=hours)
    )

    return (
        WellnessScan.query
        .filter(
            WellnessScan.personnel_id
            == personnel_id,

            WellnessScan.saved_at
            >= cutoff,

            WellnessScan.confidence.in_(
                [
                    "High",
                    "Medium"
                ]
            )
        )
        .order_by(
            WellnessScan.saved_at.desc()
        )
        .first()
    )


# --------------------------------------------------
# PERSONNEL / LEAVE HELPERS
# --------------------------------------------------

def _parse_date(value):

    if not value:
        return None

    try:
        return datetime.strptime(
            str(value).strip(),
            "%Y-%m-%d"
        ).date()
    except (TypeError, ValueError):
        return None


def _clamp_score(value):

    if value is None:
        return None

    return round(
        max(0.0, min(100.0, float(value))),
        1
    )


def _average_score(values):

    clean = [
        float(value)
        for value in values
        if value is not None
    ]

    if not clean:
        return None

    return round(
        sum(clean) / len(clean),
        1
    )


def _wellness_response_points(response, reverse=False):

    response_map = {
        "Not at all": 0,
        "Slightly": 1,
        "Moderately": 2,
        "Very much": 3
    }

    if response not in response_map:
        return None

    points = response_map[response]

    if reverse:
        points = 3 - points

    return points


def _points_to_100(points):

    if points is None:
        return None

    return round(
        (float(points) / 3.0) * 100.0,
        1
    )


def _severity_from_score(score):

    if score is None:
        return "Unavailable"

    if score >= 65:
        return "High"

    if score >= 35:
        return "Attention"

    return "Low"


def _base_signal_score_from_analysis(
    item
):

    """
    Reconstruct the pre-trend score from a stored analysis row.

    This prevents the trend engine from recursively trending a score
    that already contains trend weighting.
    """

    if (
        item is None
        or item.operational_score is None
    ):

        return None

    if item.checkin_score is not None:

        return round(
            (
                float(
                    item.operational_score
                )
                * 0.65
            )
            +
            (
                float(
                    item.checkin_score
                )
                * 0.35
            ),
            1
        )

    return round(
        float(
            item.operational_score
        ),
        1
    )


def _calculate_longitudinal_trend(
    personnel_id,
    current_base_score,
    current_assessment_id,
    max_points=5
):

    """
    Calculate concern trend from stored previous VeerPulse analyses.

    Minimum:
        current score + 2 previous valid analysis points

    Direction:
        positive slope  -> increasing concern / declining wellness
        near-zero slope -> stable
        negative slope  -> improving

    Trend score:
        only worsening slope contributes concern.
        +10 score-points per assessment interval maps to 100/100.
    """

    if current_base_score is None:

        return {
            "available": False,
            "score": None,
            "direction": "Insufficient Trend Data",
            "strength": "Unavailable",
            "slope": None,
            "point_count": 0,
            "scores": []
        }

    previous_rows = (
        WelfareAnalysisResult.query
        .filter(
            WelfareAnalysisResult.personnel_id
            == personnel_id,

            WelfareAnalysisResult.assessment_id
            != current_assessment_id,

            WelfareAnalysisResult.operational_score
            .isnot(None)
        )
        .order_by(
            WelfareAnalysisResult.created_at.desc()
        )
        .limit(
            max(
                0,
                int(max_points) - 1
            )
        )
        .all()
    )

    previous_rows.reverse()

    series = []

    for item in previous_rows:

        score = (
            _base_signal_score_from_analysis(
                item
            )
        )

        if score is not None:

            series.append(
                float(score)
            )

    series.append(
        float(
            current_base_score
        )
    )

    if len(series) < 3:

        return {
            "available": False,
            "score": None,
            "direction": "Insufficient Trend Data",
            "strength": "Need 3+ Assessments",
            "slope": None,
            "point_count": len(series),
            "scores": [
                round(value, 1)
                for value in series
            ]
        }

    # Simple least-squares slope across equally spaced
    # assessment points.
    x_values = list(
        range(
            len(series)
        )
    )

    x_mean = (
        sum(x_values)
        / len(x_values)
    )

    y_mean = (
        sum(series)
        / len(series)
    )

    numerator = sum(
        (
            x_value - x_mean
        )
        *
        (
            y_value - y_mean
        )
        for x_value, y_value
        in zip(
            x_values,
            series
        )
    )

    denominator = sum(
        (
            x_value - x_mean
        )
        ** 2
        for x_value
        in x_values
    )

    slope = (
        numerator / denominator
        if denominator
        else 0.0
    )

    slope = round(
        slope,
        2
    )

    # Only a worsening trend adds concern.
    trend_score = _clamp_score(
        max(
            0.0,
            slope
        )
        / 10.0
        * 100.0
    )

    if slope >= 5.0:

        direction = "Declining"

    elif slope >= 2.0:

        direction = "Gradual Decline"

    elif slope <= -5.0:

        direction = "Improving"

    elif slope <= -2.0:

        direction = "Gradual Improvement"

    else:

        direction = "Stable"

    absolute_slope = abs(
        slope
    )

    if absolute_slope >= 5.0:

        strength = "Strong"

    elif absolute_slope >= 2.0:

        strength = "Moderate"

    else:

        strength = "Stable / Weak"

    increasing_steps = sum(
        1
        for previous, current
        in zip(
            series,
            series[1:]
        )
        if (
            current
            - previous
        ) >= 2.0
    )

    decreasing_steps = sum(
        1
        for previous, current
        in zip(
            series,
            series[1:]
        )
        if (
            previous
            - current
        ) >= 2.0
    )

    return {
        "available": True,
        "score": trend_score,
        "direction": direction,
        "strength": strength,
        "slope": slope,
        "point_count": len(series),
        "scores": [
            round(value, 1)
            for value in series
        ],
        "increasing_steps":
            increasing_steps,
        "decreasing_steps":
            decreasing_steps
    }


def _calculate_signal_agreement(
    operational_score,
    checkin_score,
    trend_score,
    trend_direction,
    scan_evidence=None
):

    """
    VeerPulse Signal Agreement Engine v1.0

    Purpose:
        Describe whether independent welfare-support signals point
        in a similar direction.

    Important:
        - NOT lie detection.
        - NOT a medical diagnosis.
        - Camera HR / breathing values are NOT interpreted as
          stress or concern in this function.
        - Camera evidence is used only as supplementary
          availability / quality context.

    Directional signals:
        Operational pattern
        Voluntary Check-In
        Longitudinal trend

    Concern threshold for operational / Check-In:
        score >= 35

    Trend:
        Declining / Gradual Decline -> Concern
        Stable / Improving / Gradual Improvement -> Stable
    """

    signals = []

    if operational_score is not None:

        signals.append({
            "name":
                "Operational Pattern",
            "state":
                (
                    "Concern"
                    if operational_score >= 35
                    else "Stable"
                ),
            "score":
                round(
                    float(
                        operational_score
                    ),
                    1
                ),
            "source":
                "OperationalRecord"
        })

    if checkin_score is not None:

        signals.append({
            "name":
                "Voluntary Check-In",
            "state":
                (
                    "Concern"
                    if checkin_score >= 35
                    else "Stable"
                ),
            "score":
                round(
                    float(
                        checkin_score
                    ),
                    1
                ),
            "source":
                "WellnessCheckIn"
        })

    if (
        trend_score is not None
        and trend_direction
        not in (
            None,
            "",
            "Insufficient Trend Data",
            "Not Calculated Yet"
        )
    ):

        trend_concern = (
            trend_direction
            in (
                "Declining",
                "Gradual Decline"
            )
        )

        signals.append({
            "name":
                "Longitudinal Trend",
            "state":
                (
                    "Concern"
                    if trend_concern
                    else "Stable"
                ),
            "score":
                round(
                    float(
                        trend_score
                    ),
                    1
                ),
            "source":
                "WelfareAnalysisResult"
        })

    scan_context = {
        "available":
            False,
        "usable":
            False,
        "signal_quality":
            None,
        "confidence":
            None,
        "interpretation":
            (
                "No usable camera context linked."
            )
    }

    if scan_evidence:

        scan_context["available"] = True

        scan_context[
            "signal_quality"
        ] = (
            scan_evidence.signal_quality
        )

        scan_context[
            "confidence"
        ] = (
            scan_evidence.confidence
        )

        scan_context["usable"] = (
            scan_evidence.confidence
            in (
                "High",
                "Medium"
            )
        )

        if scan_context["usable"]:

            scan_context[
                "interpretation"
            ] = (
                "Usable supplementary camera context is available, "
                "but HR / breathing values are not direction-scored "
                "as stress evidence."
            )

        else:

            scan_context[
                "interpretation"
            ] = (
                "Camera context exists but does not meet the "
                "usable-confidence requirement."
            )

    signal_count = len(
        signals
    )

    if signal_count < 2:

        return {
            "label":
                "Insufficient Agreement Data",
            "agreement_score":
                None,
            "direction":
                None,
            "signal_count":
                signal_count,
            "concern_count":
                sum(
                    1
                    for item in signals
                    if item["state"]
                    == "Concern"
                ),
            "stable_count":
                sum(
                    1
                    for item in signals
                    if item["state"]
                    == "Stable"
                ),
            "signals":
                signals,
            "scan_context":
                scan_context,
            "explanation":
                (
                    "At least two independent directional signals "
                    "are required before VeerPulse describes signal agreement."
                )
        }

    concern_count = sum(
        1
        for item in signals
        if item["state"]
        == "Concern"
    )

    stable_count = sum(
        1
        for item in signals
        if item["state"]
        == "Stable"
    )

    dominant_count = max(
        concern_count,
        stable_count
    )

    agreement_score = round(
        (
            dominant_count
            / signal_count
        )
        * 100.0,
        1
    )

    if concern_count == signal_count:

        label = "High Agreement"
        direction = "Concern"

    elif stable_count == signal_count:

        label = "High Agreement"
        direction = "Stable"

    elif (
        concern_count
        == stable_count
    ):

        label = "Signal Mismatch"
        direction = "Mixed"

    else:

        label = "Partial Agreement"

        direction = (
            "Concern"
            if concern_count
            > stable_count
            else "Stable"
        )

    explanation_parts = [
        (
            item["name"]
            + ": "
            + item["state"]
        )
        for item in signals
    ]

    explanation = (
        "; ".join(
            explanation_parts
        )
        + ". Agreement describes consistency between available "
          "welfare-support signals and does not determine whether "
          "a person is truthful."
    )

    return {
        "label":
            label,
        "agreement_score":
            agreement_score,
        "direction":
            direction,
        "signal_count":
            signal_count,
        "concern_count":
            concern_count,
        "stable_count":
            stable_count,
        "signals":
            signals,
        "scan_context":
            scan_context,
        "explanation":
            explanation
    }


def _build_veerpulse_analysis(
    personnel_id,
    assessment,
    evidence_snapshot,
    recent_checkin=None,
    checkin_evidence=None,
    scan_evidence=None,
    current_records=None,
    historical_records=None
):

    """
    VeerPulse Early-Warning Engine v1.0.

    Prototype welfare-support scoring only.
    Not a medical or psychological diagnosis.

    Weighting:
        65% validated operational evidence
        35% voluntary Check-In evidence

    Camera scan remains supplementary and is not numerically weighted.
    """

    current_records = current_records or []
    historical_records = historical_records or []

    duty_score = None
    rest_score = None
    operational_recovery_score = None
    operational_score = None

    duty_delta = (
        evidence_snapshot.duty_delta
        if evidence_snapshot
        else None
    )

    rest_delta = (
        evidence_snapshot.rest_delta
        if evidence_snapshot
        else None
    )

    recovery_delta = (
        evidence_snapshot.recovery_day_delta
        if evidence_snapshot
        else None
    )

    operational_evidence_ready = (
        evidence_snapshot is not None
        and evidence_snapshot.evidence_status
        in [
            "Sufficient Operational Evidence",
            "Operational + Check-In Context",
            "Operational + Scan Context",
            "Operational + Check-In + Scan Context"
        ]
        and duty_delta is not None
        and rest_delta is not None
        and recovery_delta is not None
    )

    if operational_evidence_ready:

        duty_score = _clamp_score(
            max(0.0, duty_delta)
            / 2.0
            * 100.0
        )

        rest_score = _clamp_score(
            max(0.0, -rest_delta)
            / 2.0
            * 100.0
        )

        operational_recovery_score = _clamp_score(
            max(
                0.0,
                -float(recovery_delta)
            )
            / 2.0
            * 100.0
        )

        operational_score = _average_score([
            duty_score,
            rest_score,
            operational_recovery_score
        ])

    workload_score = None
    checkin_recovery_score = None
    connection_score = None
    cognitive_score = None
    checkin_score = None

    if recent_checkin:

        workload_items = [
            _points_to_100(
                _wellness_response_points(
                    recent_checkin.workload_pressure
                )
            ),
            _points_to_100(
                _wellness_response_points(
                    recent_checkin.duty_adjustment
                )
            )
        ]

        recovery_items = [
            _points_to_100(
                _wellness_response_points(
                    recent_checkin.recovery_time,
                    reverse=True
                )
            ),
            _points_to_100(
                _wellness_response_points(
                    recent_checkin.rest_quality,
                    reverse=True
                )
            ),
            _points_to_100(
                _wellness_response_points(
                    recent_checkin.low_energy
                )
            )
        ]

        connection_items = [
            _points_to_100(
                _wellness_response_points(
                    recent_checkin.family_connection,
                    reverse=True
                )
            ),
            _points_to_100(
                _wellness_response_points(
                    recent_checkin.family_separation
                )
            )
        ]

        cognitive_items = [
            _points_to_100(
                _wellness_response_points(
                    recent_checkin.focus,
                    reverse=True
                )
            ),
            _points_to_100(
                _wellness_response_points(
                    recent_checkin.mental_tiredness
                )
            ),
            _points_to_100(
                _wellness_response_points(
                    recent_checkin.pattern_change
                )
            )
        ]

        workload_score = _average_score(
            workload_items
        )

        checkin_recovery_score = _average_score(
            recovery_items
        )

        connection_score = _average_score(
            connection_items
        )

        cognitive_score = _average_score(
            cognitive_items
        )

        checkin_score = _average_score(
            workload_items
            + recovery_items
            + connection_items
            + cognitive_items
        )

    # --------------------------------------------------
    # CURRENT SIGNAL SCORE
    #
    # This preserves Step 23's transparent fusion:
    #   Operational 65% + Check-In 35%
    # before longitudinal trend is applied.
    # --------------------------------------------------

    base_signal_score = None
    overall_score = None
    attention_level = "Insufficient Data"

    if operational_score is not None:

        if checkin_score is not None:

            base_signal_score = round(
                (
                    operational_score
                    * 0.65
                )
                +
                (
                    checkin_score
                    * 0.35
                ),
                1
            )

        else:

            base_signal_score = (
                operational_score
            )

    # --------------------------------------------------
    # LONGITUDINAL TREND ENGINE
    # --------------------------------------------------

    trend_result = (
        _calculate_longitudinal_trend(
            personnel_id=
                personnel_id,
            current_base_score=
                base_signal_score,
            current_assessment_id=
                assessment.id,
            max_points=5
        )
    )

    trend_score = (
        trend_result[
            "score"
        ]
    )

    trend_direction = (
        trend_result[
            "direction"
        ]
    )

    trend_strength = (
        trend_result[
            "strength"
        ]
    )

    # If a valid longitudinal trend exists, it contributes
    # 20% of the final Early-Warning Score.
    #
    # The current operational + Check-In signal keeps 80%.
    # A stable/improving trend contributes 0 additional concern.
    if base_signal_score is not None:

        if trend_score is not None:

            overall_score = round(
                (
                    base_signal_score
                    * 0.80
                )
                +
                (
                    trend_score
                    * 0.20
                ),
                1
            )

        else:

            overall_score = (
                base_signal_score
            )

        operational_factor_count = (
            evidence_snapshot.concerning_factor_count
            if (
                evidence_snapshot
                and evidence_snapshot.concerning_factor_count
                is not None
            )
            else 0
        )

        if overall_score >= 65:

            # Highest priority still requires operational concern.
            # Trend or self-report alone cannot create Elevated.
            if (
                operational_score >= 50
                and operational_factor_count >= 2
            ):

                attention_level = (
                    "Elevated"
                )

            else:

                attention_level = (
                    "Moderate"
                )

        elif overall_score >= 35:

            attention_level = (
                "Moderate"
            )

        else:

            attention_level = (
                "Low"
            )

    # --------------------------------------------------
    # CONFIDENCE — EVIDENCE COMPLETENESS,
    # NOT MEDICAL CERTAINTY
    # --------------------------------------------------

    confidence = None

    if operational_score is not None:

        confidence = "Medium"

        if (
            checkin_score is not None
            and evidence_snapshot
            and (
                evidence_snapshot.current_record_count
                or 0
            ) >= 4
            and (
                evidence_snapshot.baseline_record_count
                or 0
            ) >= 4
        ):

            confidence = "High"

    # --------------------------------------------------
    # SIGNAL AGREEMENT / RELIABILITY
    #
    # This does not alter the welfare-attention score.
    # It explains whether independent evidence sources point
    # in the same direction.
    # --------------------------------------------------

    agreement_result = (
        _calculate_signal_agreement(
            operational_score=
                operational_score,
            checkin_score=
                checkin_score,
            trend_score=
                trend_score,
            trend_direction=
                trend_direction,
            scan_evidence=
                scan_evidence
        )
    )

    signal_agreement = (
        agreement_result[
            "label"
        ]
    )

    # Confidence represents evidence completeness + coherence,
    # not medical certainty.
    if operational_score is not None:

        if (
            confidence == "High"
            and signal_agreement
            == "Signal Mismatch"
        ):

            confidence = "Medium"

        elif (
            confidence == "Medium"
            and signal_agreement
            == "High Agreement"
            and agreement_result[
                "signal_count"
            ] >= 2
        ):

            confidence = "Medium"

    if overall_score is None:

        analysis_reason = (
            "VeerPulse withheld the Early-Warning Score because "
            "sufficient comparable operational evidence is not available. "
            "Voluntary Check-In or camera information does not replace "
            "the operational evidence gate."
        )

        recommendation = (
            "Collect sufficient comparable operational evidence "
            "before producing a welfare-attention score."
        )

    else:

        analysis_reason = (
            "VeerPulse Early-Warning Engine v1.0 calculated an "
            f"operational strain score of {operational_score:.1f}/100"
        )

        if checkin_score is not None:

            analysis_reason += (
                " and a voluntary self-reported wellness concern "
                f"score of {checkin_score:.1f}/100. "
                "The current prototype combines these at "
                "65% operational evidence and 35% voluntary "
                "Check-In evidence."
            )

        else:

            analysis_reason += (
                ". No recent voluntary Check-In was available, "
                "so the current score uses operational evidence only."
            )

        if trend_score is not None:

            analysis_reason += (
                f" Longitudinal analysis used "
                f"{trend_result['point_count']} stored assessment points. "
                f"The trend is {trend_direction} "
                f"({trend_strength}), with a slope of "
                f"{trend_result['slope']:+.2f} score points per "
                "assessment interval and a trend concern score of "
                f"{trend_score:.1f}/100. "
                "The final score applies 80% current signal and "
                "20% longitudinal trend."
            )

        else:

            analysis_reason += (
                " Longitudinal trend was not weighted because at least "
                "three valid analysis points are required."
            )

        analysis_reason += (
            f" The resulting Early-Warning Score is "
            f"{overall_score:.1f}/100, corresponding to "
            f"{attention_level} welfare attention. "
            "This is a welfare-support indicator, not a medical diagnosis."
        )

        analysis_reason += (
            " Signal Agreement: "
            + signal_agreement
            + ". "
            + agreement_result[
                "explanation"
            ]
        )

        if agreement_result[
            "scan_context"
        ][
            "available"
        ]:

            analysis_reason += (
                " "
                + agreement_result[
                    "scan_context"
                ][
                    "interpretation"
                ]
            )

        if (
            recent_checkin
            and recent_checkin.support_request
            == "Yes"
        ):

            analysis_reason += (
                " The personnel member separately requested "
                "human welfare support."
            )

        if attention_level == "Elevated":

            recommendation = (
                "Priority human welfare review recommended. "
                "No automated disciplinary, leave, posting or "
                "deployment decision should be made from this score."
            )

        elif attention_level == "Moderate":

            recommendation = (
                "Human welfare review or continued closer monitoring "
                "may be appropriate."
            )

        else:

            recommendation = (
                "Continue normal welfare monitoring and reassess "
                "when new validated evidence becomes available."
            )

    (
        WelfareAnalysisResult.query
        .filter(
            WelfareAnalysisResult.personnel_id
            == personnel_id,
            WelfareAnalysisResult.is_current
            == True,
            WelfareAnalysisResult.assessment_id
            != assessment.id
        )
        .update({
            "is_current": False
        })
    )

    result = (
        WelfareAnalysisResult.query
        .filter_by(
            assessment_id=assessment.id
        )
        .first()
    )

    if result is None:

        result = WelfareAnalysisResult(
            personnel_id=personnel_id,
            assessment_id=assessment.id,
            assessment_version=
                assessment.assessment_version
        )

        db.session.add(result)
        db.session.flush()

    result.personnel_id = personnel_id
    result.assessment_version = (
        assessment.assessment_version
    )
    result.operational_score = operational_score
    result.checkin_score = checkin_score
    result.trend_score = trend_score
    result.overall_score = overall_score
    result.workload_score = workload_score
    result.recovery_score = (
        checkin_recovery_score
    )
    result.connection_score = (
        connection_score
    )
    result.cognitive_score = cognitive_score
    result.attention_level = attention_level
    result.trend_direction = trend_direction
    result.trend_strength = trend_strength
    result.signal_agreement = signal_agreement
    result.confidence = confidence
    result.analysis_reason = analysis_reason
    result.recommendation = recommendation
    result.rules_version = "VP-BRAIN-1.2"
    result.operational_snapshot_id = (
        evidence_snapshot.id
        if evidence_snapshot
        else None
    )
    result.checkin_evidence_id = (
        checkin_evidence.id
        if checkin_evidence
        else None
    )
    result.scan_evidence_id = (
        scan_evidence.id
        if scan_evidence
        else None
    )
    result.is_current = True
    result.updated_at = datetime.now()

    AnalysisFactor.query.filter_by(
        analysis_result_id=result.id
    ).delete()

    def average_raw(values):

        clean = [
            float(value)
            for value in values
            if value is not None
        ]

        if not clean:
            return None

        return round(
            sum(clean) / len(clean),
            1
        )

    current_worked = [
        item
        for item in current_records
        if (
            item.duty_hours is not None
            and item.duty_hours > 0
        )
    ]

    baseline_worked = [
        item
        for item in historical_records
        if (
            item.duty_hours is not None
            and item.duty_hours > 0
        )
    ]

    current_duty = average_raw([
        item.duty_hours
        for item in current_worked
    ])

    baseline_duty = average_raw([
        item.duty_hours
        for item in baseline_worked
    ])

    current_rest = average_raw([
        item.rest_hours
        for item in current_records
    ])

    baseline_rest = average_raw([
        item.rest_hours
        for item in historical_records
    ])

    current_recovery_days = (
        sum(
            1
            for item in current_records
            if (
                item.duty_hours is not None
                and float(item.duty_hours) == 0
            )
        )
        if current_records
        else None
    )

    baseline_recovery_days = (
        sum(
            1
            for item in historical_records
            if (
                item.duty_hours is not None
                and float(item.duty_hours) == 0
            )
        )
        if historical_records
        else None
    )

    factor_rows = []

    if operational_score is not None:

        factor_rows.extend([
            {
                "key": "duty_intensity",
                "label":
                    "Worked-Day Duty Intensity",
                "source":
                    "OperationalRecord",
                "current": current_duty,
                "baseline": baseline_duty,
                "deviation":
                    (
                        round(
                            (
                                (
                                    current_duty
                                    - baseline_duty
                                )
                                / baseline_duty
                            )
                            * 100,
                            1
                        )
                        if (
                            current_duty is not None
                            and baseline_duty
                            not in (None, 0)
                        )
                        else None
                    ),
                "score": duty_score,
                "severity":
                    _severity_from_score(
                        duty_score
                    ),
                "direction":
                    (
                        "Above Baseline"
                        if (
                            duty_delta is not None
                            and duty_delta > 0
                        )
                        else "Stable / Lower"
                    ),
                "explanation":
                    (
                        "Concern increases when worked-day duty "
                        "rises above the person's comparable baseline."
                    )
            },
            {
                "key": "average_rest",
                "label":
                    "Average Rest",
                "source":
                    "OperationalRecord",
                "current": current_rest,
                "baseline": baseline_rest,
                "deviation":
                    (
                        round(
                            (
                                (
                                    current_rest
                                    - baseline_rest
                                )
                                / baseline_rest
                            )
                            * 100,
                            1
                        )
                        if (
                            current_rest is not None
                            and baseline_rest
                            not in (None, 0)
                        )
                        else None
                    ),
                "score": rest_score,
                "severity":
                    _severity_from_score(
                        rest_score
                    ),
                "direction":
                    (
                        "Below Baseline"
                        if (
                            rest_delta is not None
                            and rest_delta < 0
                        )
                        else "Stable / Higher"
                    ),
                "explanation":
                    (
                        "Concern increases when average rest falls "
                        "below the person's comparable baseline."
                    )
            },
            {
                "key": "recovery_days",
                "label":
                    "Recovery-Day Availability",
                "source":
                    "OperationalRecord",
                "current":
                    current_recovery_days,
                "baseline":
                    baseline_recovery_days,
                "deviation": None,
                "score":
                    operational_recovery_score,
                "severity":
                    _severity_from_score(
                        operational_recovery_score
                    ),
                "direction":
                    (
                        "Reduced"
                        if (
                            recovery_delta is not None
                            and recovery_delta < 0
                        )
                        else "Stable / Higher"
                    ),
                "explanation":
                    (
                        "Concern increases when recovery-day "
                        "availability falls below the comparable baseline."
                    )
            }
        ])

    if checkin_score is not None:

        factor_rows.extend([
            {
                "key": "checkin_workload",
                "label":
                    "Check-In Workload / Adjustment",
                "source":
                    "WellnessCheckIn",
                "current": workload_score,
                "baseline": None,
                "deviation": None,
                "score": workload_score,
                "severity":
                    _severity_from_score(
                        workload_score
                    ),
                "direction":
                    "Current Self-Report",
                "explanation":
                    (
                        "Derived from workload pressure and duty-adjustment "
                        "responses in the voluntary Check-In."
                    )
            },
            {
                "key": "checkin_recovery",
                "label":
                    "Check-In Recovery / Energy",
                "source":
                    "WellnessCheckIn",
                "current":
                    checkin_recovery_score,
                "baseline": None,
                "deviation": None,
                "score":
                    checkin_recovery_score,
                "severity":
                    _severity_from_score(
                        checkin_recovery_score
                    ),
                "direction":
                    "Current Self-Report",
                "explanation":
                    (
                        "Derived from recovery time, rest quality and "
                        "low-energy responses."
                    )
            },
            {
                "key":
                    "checkin_connection",
                "label":
                    "Check-In Social / Family Connection",
                "source":
                    "WellnessCheckIn",
                "current":
                    connection_score,
                "baseline": None,
                "deviation": None,
                "score":
                    connection_score,
                "severity":
                    _severity_from_score(
                        connection_score
                    ),
                "direction":
                    "Current Self-Report",
                "explanation":
                    (
                        "Derived from family-connection and "
                        "family-separation responses."
                    )
            },
            {
                "key":
                    "checkin_cognitive",
                "label":
                    "Check-In Cognitive / Mental Fatigue",
                "source":
                    "WellnessCheckIn",
                "current":
                    cognitive_score,
                "baseline": None,
                "deviation": None,
                "score":
                    cognitive_score,
                "severity":
                    _severity_from_score(
                        cognitive_score
                    ),
                "direction":
                    "Current Self-Report",
                "explanation":
                    (
                        "Derived from focus, mental tiredness and "
                        "pattern-change responses."
                    )
            }
        ])

    if (
        agreement_result[
            "signal_count"
        ] >= 2
    ):

        factor_rows.append({
            "key":
                "signal_agreement",
            "label":
                "Wellness Signal Agreement",
            "source":
                "Multi-Source Analysis",
            "current":
                agreement_result[
                    "agreement_score"
                ],
            "baseline":
                None,
            "deviation":
                None,
            "score":
                agreement_result[
                    "agreement_score"
                ],
            "severity":
                (
                    "Mismatch"
                    if signal_agreement
                    == "Signal Mismatch"
                    else (
                        "High Agreement"
                        if signal_agreement
                        == "High Agreement"
                        else "Partial Agreement"
                    )
                ),
            "direction":
                agreement_result[
                    "direction"
                ],
            "explanation":
                agreement_result[
                    "explanation"
                ]
        })

    if trend_score is not None:

        factor_rows.append({
            "key":
                "longitudinal_trend",
            "label":
                "Longitudinal Early-Warning Trend",
            "source":
                "WelfareAnalysisResult",
            "current":
                trend_result[
                    "slope"
                ],
            "baseline":
                None,
            "deviation":
                None,
            "score":
                trend_score,
            "severity":
                _severity_from_score(
                    trend_score
                ),
            "direction":
                trend_direction,
            "explanation":
                (
                    "Calculated from "
                    + str(
                        trend_result[
                            "point_count"
                        ]
                    )
                    + " stored assessment points. "
                    + "Scores used: "
                    + ", ".join(
                        str(value)
                        for value
                        in trend_result[
                            "scores"
                        ]
                    )
                    + ". Positive slope means concern is increasing; "
                    + "negative slope means the pattern is improving."
                )
        })

    for item in factor_rows:

        db.session.add(
            AnalysisFactor(
                analysis_result_id=result.id,
                assessment_id=assessment.id,
                personnel_id=personnel_id,
                factor_key=item["key"],
                factor_label=item["label"],
                source_type=item["source"],
                current_value=item["current"],
                baseline_value=item["baseline"],
                deviation_pct=item["deviation"],
                factor_score=item["score"],
                severity=item["severity"],
                direction=item["direction"],
                explanation=item["explanation"],
                created_at=datetime.now()
            )
        )

    return result


def _analysis_result_to_dict(
    item,
    include_factors=True
):

    if item is None:
        return None

    data = {
        "id":
            item.id,

        "personnel_id":
            item.personnel_id,

        "assessment_id":
            item.assessment_id,

        "assessment_version":
            item.assessment_version,

        "operational_score":
            item.operational_score,

        "checkin_score":
            item.checkin_score,

        "trend_score":
            item.trend_score,

        "overall_score":
            item.overall_score,

        "workload_score":
            item.workload_score,

        "recovery_score":
            item.recovery_score,

        "connection_score":
            item.connection_score,

        "cognitive_score":
            item.cognitive_score,

        "attention_level":
            item.attention_level,

        "trend_direction":
            item.trend_direction,

        "trend_strength":
            item.trend_strength,

        "signal_agreement":
            item.signal_agreement,

        "confidence":
            item.confidence,

        "analysis_reason":
            item.analysis_reason,

        "recommendation":
            item.recommendation,

        "rules_version":
            item.rules_version,

        "operational_snapshot_id":
            item.operational_snapshot_id,

        "checkin_evidence_id":
            item.checkin_evidence_id,

        "scan_evidence_id":
            item.scan_evidence_id,

        "is_current":
            item.is_current,

        "created_at":
            (
                item.created_at.isoformat()
                if item.created_at
                else None
            ),

        "updated_at":
            (
                item.updated_at.isoformat()
                if item.updated_at
                else None
            )
    }

    if include_factors:

        factors = (
            AnalysisFactor.query
            .filter_by(
                analysis_result_id=
                    item.id
            )
            .order_by(
                AnalysisFactor.id.asc()
            )
            .all()
        )

        data["factors"] = [
            {
                "id":
                    factor.id,

                "factor_key":
                    factor.factor_key,

                "factor_label":
                    factor.factor_label,

                "source_type":
                    factor.source_type,

                "current_value":
                    factor.current_value,

                "baseline_value":
                    factor.baseline_value,

                "deviation_pct":
                    factor.deviation_pct,

                "factor_score":
                    factor.factor_score,

                "severity":
                    factor.severity,

                "direction":
                    factor.direction,

                "explanation":
                    factor.explanation
            }
            for factor in factors
        ]

    return data


def _current_welfare_analysis(
    personnel_id
):

    return (
        WelfareAnalysisResult.query
        .filter_by(
            personnel_id=
                personnel_id,
            is_current=True
        )
        .order_by(
            WelfareAnalysisResult.created_at.desc()
        )
        .first()
    )


def _personnel_to_dict(personnel):

    return {
        "id": personnel.id,
        "personnel_id": personnel.personnel_id,
        "full_name": personnel.full_name,
        "rank_designation": personnel.rank_designation,
        "unit": personnel.unit,
        "department": personnel.department,
        "posting_location": personnel.posting_location,
        "joining_date": (
            personnel.joining_date.isoformat()
            if personnel.joining_date
            else None
        ),
        "email": personnel.email,
        "phone": personnel.phone,
        "employment_status": personnel.employment_status,
        "created_by": personnel.created_by,
        "updated_by": personnel.updated_by,
        "created_at": (
            personnel.created_at.isoformat()
            if personnel.created_at
            else None
        ),
        "updated_at": (
            personnel.updated_at.isoformat()
            if personnel.updated_at
            else None
        )
    }


def _leave_to_dict(item):

    return {
        "id": item.id,
        "request_code": item.request_code,
        "personnel_id": item.personnel_id,
        "leave_type": item.leave_type,
        "start_date": item.start_date.isoformat(),
        "end_date": item.end_date.isoformat(),
        "requested_days": item.requested_days,
        "reason_category": item.reason_category,
        "reason": item.reason,
        "additional_note": item.additional_note,
        "status": item.status,
        "decision_reason": item.decision_reason,
        "decision_note": item.decision_note,
        "decided_by": item.decided_by,
        "decided_at": (
            item.decided_at.isoformat()
            if item.decided_at
            else None
        ),
        "cancelled_at": (
            item.cancelled_at.isoformat()
            if item.cancelled_at
            else None
        ),
        "created_at": item.created_at.isoformat(),
        "updated_at": item.updated_at.isoformat()
    }


def _add_audit_event(
    actor_role,
    actor_id,
    event_type,
    personnel_id,
    resource_type,
    resource_id,
    event_summary
):

    db.session.add(
        AuditEvent(
            actor_role=actor_role,
            actor_id=actor_id,
            event_type=event_type,
            personnel_id=personnel_id,
            resource_type=resource_type,
            resource_id=(
                str(resource_id)
                if resource_id is not None
                else None
            ),
            event_summary=event_summary,
            source_type="VeerPulse Application",
            created_at=datetime.now()
        )
    )


def _approved_leave_to_operational_records(leave_request):

    current_day = leave_request.start_date

    while current_day <= leave_request.end_date:

        existing = (
            OperationalRecord.query
            .filter_by(
                personnel_id=leave_request.personnel_id,
                record_date=current_day,
                is_valid=True
            )
            .order_by(
                OperationalRecord.id.desc()
            )
            .first()
        )

        source_record_id = (
            "LEAVE-"
            + str(leave_request.id)
            + "-"
            + current_day.strftime("%Y%m%d")
        )

        if existing:
            existing.duty_hours = 0.0
            existing.rest_hours = 24.0
            existing.context_type = "Approved Leave"
            existing.deployment_active = False
            existing.training_hours = 0.0
            existing.leave_status = "Approved Leave"
            existing.source_type = "Leave Management"
            existing.source_record_id = source_record_id
        else:
            db.session.add(
                OperationalRecord(
                    personnel_id=leave_request.personnel_id,
                    record_date=current_day,
                    duty_hours=0.0,
                    rest_hours=24.0,
                    context_type="Approved Leave",
                    deployment_active=False,
                    training_hours=0.0,
                    leave_status="Approved Leave",
                    source_type="Leave Management",
                    source_record_id=source_record_id,
                    is_valid=True,
                    created_at=datetime.now()
                )
            )

        current_day += timedelta(days=1)


# --------------------------------------------------
# CREATE FLASK APP
# --------------------------------------------------

def create_app():

    app = Flask(__name__)

    # Use a deployment SECRET_KEY when available.
    # A per-process random fallback avoids shipping a hard-coded secret.
    # For Render/production, set SECRET_KEY in environment variables.
    app.secret_key = (
        os.getenv("SECRET_KEY")
        or os.urandom(32)
    )


    # --------------------------------------------------
    # POSTGRESQL CONFIGURATION
    # --------------------------------------------------

    database_url = os.getenv("DATABASE_URL")

    if not database_url:

        raise RuntimeError(
            "DATABASE_URL was not found. Check the .env file."
        )

    app.config["SQLALCHEMY_DATABASE_URI"] = database_url

    app.config[
        "SQLALCHEMY_TRACK_MODIFICATIONS"
    ] = False

    db.init_app(app)

    with app.app_context():
        db.create_all()

        initial_officer_password = (
            os.getenv(
                "VEERPULSE_INITIAL_OFFICER_PASSWORD"
            )
            or ""
        )

        default_officer = (
            WelfareOfficer.query
            .filter_by(
                officer_id="officer001"
            )
            .first()
        )

        if (
            default_officer is None
            and initial_officer_password
        ):

            default_officer = WelfareOfficer(
                officer_id="officer001",
                full_name="Welfare Officer",
                designation="Authorized Welfare Officer",
                assigned_unit="All Units",
                email=None,
                password_hash=generate_password_hash(
                    initial_officer_password
                ),
                account_status="Active",
                created_at=datetime.now(),
                updated_at=datetime.now()
            )

            db.session.add(
                default_officer
            )

            db.session.flush()

        if default_officer is not None:

            default_officer_preference = (
                WelfareOfficerPreference.query
                .filter_by(
                    officer_id=
                        default_officer.officer_id
                )
                .first()
            )

            if (
                default_officer_preference
                is None
            ):

                db.session.add(
                    WelfareOfficerPreference(
                        officer_id=
                            default_officer.officer_id,
                        notify_welfare_alerts=True,
                        notify_support_requests=True,
                        notify_case_updates=True,
                        updated_at=datetime.now()
                    )
                )

        initial_personnel_password = (
            os.getenv(
                "VEERPULSE_INITIAL_PERSONNEL_PASSWORD"
            )
            or ""
        )

        default_personnel = (
            Personnel.query
            .filter_by(
                personnel_id="VP1042"
            )
            .first()
        )

        if (
            default_personnel is None
            and initial_personnel_password
        ):

            default_personnel = Personnel(
                personnel_id="VP1042",
                full_name="Prototype Personnel",
                rank_designation="Personnel",
                unit="Prototype Unit",
                department="Operations",
                posting_location=None,
                joining_date=None,
                email=None,
                phone=None,
                password_hash=generate_password_hash(
                    initial_personnel_password
                ),
                employment_status="Active",
                created_by="system",
                updated_by="system",
                created_at=datetime.now(),
                updated_at=datetime.now()
            )

            db.session.add(
                default_personnel
            )

        db.session.commit()


    # --------------------------------------------------
    # HOME
    # --------------------------------------------------

    @app.route("/")
    def home():

        return render_template(
            "home.html"
        )


    # --------------------------------------------------
    # PERSONNEL LOGIN
    # --------------------------------------------------

    @app.route(
        "/personnel-login",
        methods=["GET", "POST"]
    )
    def personnel_login():

        error = None

        if request.method == "POST":

            personnel_id = request.form.get(
                "personnel_id"
            )

            password = request.form.get(
                "password"
            )

            personnel_id = (
                personnel_id
                or ""
            ).strip()

            password = (
                password
                or ""
            )

            personnel = (
                Personnel.query
                .filter_by(
                    personnel_id=personnel_id
                )
                .first()
            )

            if (
                personnel
                and personnel.employment_status == "Active"
                and check_password_hash(
                    personnel.password_hash,
                    password
                )
            ):

                session["role"] = "personnel"

                session[
                    "personnel_id"
                ] = personnel.personnel_id

                session[
                    "personnel_name"
                ] = personnel.full_name

                return redirect(
                    url_for(
                        "personnel_dashboard"
                    )
                )

            error = (
                "Invalid Personnel ID or Password, or the account is not active."
            )

        return render_template(
            "personnel_login.html",
            error=error
        )


    # --------------------------------------------------
    # WELFARE OFFICER LOGIN
    # --------------------------------------------------

    @app.route(
        "/welfare-officer-login",
        methods=["GET", "POST"]
    )
    def welfare_officer_login():

        error = None

        if request.method == "POST":

            officer_id = (
                request.form.get(
                    "officer_id"
                )
                or ""
            ).strip()

            password = (
                request.form.get(
                    "password"
                )
                or ""
            )

            officer = (
                WelfareOfficer.query
                .filter_by(
                    officer_id=officer_id
                )
                .first()
            )

            if (
                officer
                and officer.account_status == "Active"
                and check_password_hash(
                    officer.password_hash,
                    password
                )
            ):

                session[
                    "role"
                ] = "welfare_officer"

                session[
                    "officer_id"
                ] = officer.officer_id

                session[
                    "officer_name"
                ] = officer.full_name

                officer.last_login_at = (
                    datetime.now()
                )

                db.session.add(
                    AuditEvent(
                        actor_role=
                            "welfare_officer",
                        actor_id=
                            officer.officer_id,
                        event_type=
                            "Welfare Officer Login",
                        personnel_id=
                            None,
                        resource_type=
                            "WelfareOfficer",
                        resource_id=
                            str(officer.id),
                        event_summary=
                            "Authorized Welfare Officer login succeeded.",
                        source_type=
                            "VeerPulse Application",
                        created_at=
                            datetime.now()
                    )
                )

                db.session.commit()

                return redirect(
                    url_for(
                        "welfare_officer_dashboard"
                    )
                )

            error = (
                "Invalid Officer ID or Password, "
                "or the account is not active."
            )

        return render_template(
            "welfare_officer_login.html",
            error=error
        )


    # --------------------------------------------------
    # VEERPULSE ANALYSIS — READ-ONLY DATABASE API
    #
    # Step 22 establishes the persistent analysis layer.
    # Step 23 will populate these tables using the
    # VeerPulse Early-Warning Engine.
    # --------------------------------------------------

    @app.route(
        "/api/personnel/analysis",
        methods=["GET"]
    )
    def personnel_analysis_api():

        if (
            session.get("role")
            != "personnel"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized."
            }), 401

        personnel_id = session.get(
            "personnel_id"
        )

        analysis = (
            _current_welfare_analysis(
                personnel_id
            )
        )

        if analysis is None:

            return jsonify({
                "success": True,
                "data_status":
                    "No Analysis Yet",
                "personnel_id":
                    personnel_id,
                "analysis":
                    None,
                "message":
                    (
                        "The analysis database layer is ready, "
                        "but the VeerPulse Early-Warning Engine "
                        "has not produced a stored analysis yet."
                    )
            })

        return jsonify({
            "success": True,
            "data_status":
                "Analysis Available",
            "analysis":
                _analysis_result_to_dict(
                    analysis
                )
        })


    @app.route(
        "/api/welfare-officer/personnel/"
        "<string:personnel_id>/analysis",
        methods=["GET"]
    )
    def welfare_officer_personnel_analysis_api(
        personnel_id
    ):

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized."
            }), 401

        personnel_id = (
            personnel_id
            or ""
        ).strip()

        personnel = (
            Personnel.query
            .filter_by(
                personnel_id=
                    personnel_id
            )
            .first()
        )

        if personnel is None:

            return jsonify({
                "success": False,
                "message":
                    "Personnel record was not found."
            }), 404

        analysis = (
            _current_welfare_analysis(
                personnel_id
            )
        )

        if analysis is None:

            return jsonify({
                "success": True,
                "data_status":
                    "No Analysis Yet",
                "personnel_id":
                    personnel_id,
                "analysis":
                    None
            })

        return jsonify({
            "success": True,
            "data_status":
                "Analysis Available",
            "analysis":
                _analysis_result_to_dict(
                    analysis
                )
        })


    # --------------------------------------------------
    # VEERPULSE ANALYSIS SCREEN
    # --------------------------------------------------

    @app.route(
        "/personnel-analysis"
    )
    def personnel_analysis_page():

        if (
            session.get("role")
            != "personnel"
        ):

            return redirect(
                url_for(
                    "personnel_login"
                )
            )

        personnel_id = session.get(
            "personnel_id"
        )

        personnel = (
            Personnel.query
            .filter_by(
                personnel_id=
                    personnel_id
            )
            .first()
        )

        analysis = (
            _current_welfare_analysis(
                personnel_id
            )
        )

        current_assessment = (
            WelfareAssessment.query
            .filter_by(
                personnel_id=
                    personnel_id,
                is_current=True
            )
            .order_by(
                WelfareAssessment.assessment_time.desc()
            )
            .first()
        )

        factors = []

        signal_agreement_factor = None
        trend_factor = None

        if analysis:

            factors = (
                AnalysisFactor.query
                .filter_by(
                    analysis_result_id=
                        analysis.id
                )
                .order_by(
                    AnalysisFactor.id.asc()
                )
                .all()
            )

            signal_agreement_factor = next(
                (
                    factor
                    for factor in factors
                    if factor.factor_key
                    == "signal_agreement"
                ),
                None
            )

            trend_factor = next(
                (
                    factor
                    for factor in factors
                    if factor.factor_key
                    == "longitudinal_trend"
                ),
                None
            )

        return render_template(
            "veerpulse_analysis.html",
            viewer_role=
                "personnel",
            personnel=
                personnel,
            analysis=
                analysis,
            current_assessment=
                current_assessment,
            factors=
                factors,
            signal_agreement_factor=
                signal_agreement_factor,
            trend_factor=
                trend_factor
        )


    @app.route(
        "/welfare-officer-analysis/"
        "<string:personnel_id>"
    )
    def welfare_officer_analysis_page(
        personnel_id
    ):

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return redirect(
                url_for(
                    "welfare_officer_login"
                )
            )

        personnel_id = (
            personnel_id
            or ""
        ).strip()

        personnel = (
            Personnel.query
            .filter_by(
                personnel_id=
                    personnel_id
            )
            .first()
        )

        if personnel is None:

            return redirect(
                url_for(
                    "welfare_officer_priority_review"
                )
            )

        analysis = (
            _current_welfare_analysis(
                personnel_id
            )
        )

        current_assessment = (
            WelfareAssessment.query
            .filter_by(
                personnel_id=
                    personnel_id,
                is_current=True
            )
            .order_by(
                WelfareAssessment.assessment_time.desc()
            )
            .first()
        )

        factors = []

        signal_agreement_factor = None
        trend_factor = None

        if analysis:

            factors = (
                AnalysisFactor.query
                .filter_by(
                    analysis_result_id=
                        analysis.id
                )
                .order_by(
                    AnalysisFactor.id.asc()
                )
                .all()
            )

            signal_agreement_factor = next(
                (
                    factor
                    for factor in factors
                    if factor.factor_key
                    == "signal_agreement"
                ),
                None
            )

            trend_factor = next(
                (
                    factor
                    for factor in factors
                    if factor.factor_key
                    == "longitudinal_trend"
                ),
                None
            )

        return render_template(
            "veerpulse_analysis.html",
            viewer_role=
                "welfare_officer",
            personnel=
                personnel,
            analysis=
                analysis,
            current_assessment=
                current_assessment,
            factors=
                factors,
            signal_agreement_factor=
                signal_agreement_factor,
            trend_factor=
                trend_factor
        )


    # --------------------------------------------------
    # PERSONNEL DASHBOARD
    # --------------------------------------------------

    @app.route(
        "/personnel-dashboard"
    )
    def personnel_dashboard():

        if (
            session.get("role")
            != "personnel"
        ):

            return redirect(
                url_for(
                    "personnel_login"
                )
            )

        personnel_id = session.get(
            "personnel_id"
        )

        today = datetime.now().date()

        period_start = (
            today
            - timedelta(days=6)
        )

        personnel = (
            Personnel.query
            .filter_by(
                personnel_id=personnel_id
            )
            .first()
        )

        current_assessment = (
            WelfareAssessment.query
            .filter_by(
                personnel_id=personnel_id,
                is_current=True
            )
            .order_by(
                WelfareAssessment.assessment_time.desc()
            )
            .first()
        )

        current_records = (
            OperationalRecord.query
            .filter(
                OperationalRecord.personnel_id
                == personnel_id,

                OperationalRecord.is_valid
                == True,

                OperationalRecord.record_date
                >= period_start,

                OperationalRecord.record_date
                <= today
            )
            .order_by(
                OperationalRecord.record_date.asc()
            )
            .all()
        )

        def _dashboard_total(values):

            clean = [
                value
                for value in values
                if value is not None
            ]

            if not clean:
                return None

            return round(
                sum(clean),
                1
            )

        def _dashboard_average(values):

            clean = [
                value
                for value in values
                if value is not None
            ]

            if not clean:
                return None

            return round(
                sum(clean) / len(clean),
                1
            )

        current_duty_total = (
            _dashboard_total([
                item.duty_hours
                for item in current_records
            ])
        )

        current_rest_average = (
            _dashboard_average([
                item.rest_hours
                for item in current_records
            ])
        )

        current_recovery_days = sum(
            1
            for item in current_records
            if (
                item.duty_hours is not None
                and item.duty_hours == 0
            )
        )

        current_training_total = (
            _dashboard_total([
                item.training_hours
                for item in current_records
            ])
        )

        latest_operational = (
            current_records[-1]
            if current_records
            else None
        )

        current_context = (
            latest_operational.context_type
            if latest_operational
            and latest_operational.context_type
            else "No Recent Context"
        )

        deployment_days = sum(
            1
            for item in current_records
            if item.deployment_active
        )

        leave_contexts = [
            item.leave_status
            for item in current_records
            if item.leave_status
        ]

        latest_leave_context = (
            leave_contexts[-1]
            if leave_contexts
            else "No recent leave context"
        )

        # ----------------------------------------------
        # Context-aware equal-length personal baseline
        # ----------------------------------------------

        baseline_status = (
            "Insufficient Data"
        )

        baseline_reason = (
            "At least 4 valid current-period records are required."
        )

        baseline_summary = {
            "status":
                baseline_status,
            "reason":
                baseline_reason,
            "context":
                current_context,
            "current_record_count":
                len(current_records),
            "baseline_record_count":
                0,
            "current_duty_total":
                current_duty_total,
            "baseline_duty_total":
                None,
            "current_rest_average":
                current_rest_average,
            "baseline_rest_average":
                None,
            "current_recovery_days":
                current_recovery_days,
            "baseline_recovery_days":
                None,
            "current_training_total":
                current_training_total,
            "baseline_training_total":
                None
        }

        if (
            len(current_records) >= 4
            and latest_operational
        ):

            comparison_limit = (
                len(current_records)
            )

            # Personal operational baseline:
            # use the immediately preceding validated records,
            # including recovery-day records, rather than forcing
            # every historical row to share the latest context label.
            historical_records = (
                OperationalRecord.query
                .filter(
                    OperationalRecord.personnel_id
                    == personnel_id,

                    OperationalRecord.is_valid
                    == True,

                    OperationalRecord.record_date
                    < period_start
                )
                .order_by(
                    OperationalRecord.record_date.desc()
                )
                .limit(
                    comparison_limit
                )
                .all()
            )

            historical_records.reverse()

            if len(historical_records) >= 4:

                baseline_status = (
                    "Available"
                )

                baseline_reason = (
                    "Compared with the same person's immediately preceding "
                    "equal-length validated operational records."
                )

                baseline_duty_total = (
                    _dashboard_total([
                        item.duty_hours
                        for item in historical_records
                    ])
                )

                baseline_rest_average = (
                    _dashboard_average([
                        item.rest_hours
                        for item in historical_records
                    ])
                )

                baseline_recovery_days = sum(
                    1
                    for item in historical_records
                    if (
                        item.duty_hours is not None
                        and item.duty_hours == 0
                    )
                )

                baseline_training_total = (
                    _dashboard_total([
                        item.training_hours
                        for item in historical_records
                    ])
                )

                baseline_summary.update({
                    "status":
                        baseline_status,
                    "reason":
                        baseline_reason,
                    "baseline_record_count":
                        len(historical_records),
                    "baseline_duty_total":
                        baseline_duty_total,
                    "baseline_rest_average":
                        baseline_rest_average,
                    "baseline_recovery_days":
                        baseline_recovery_days,
                    "baseline_training_total":
                        baseline_training_total
                })

            else:

                baseline_summary.update({
                    "reason":
                        "Not enough validated historical records exist "
                        "for the personal operational baseline.",
                    "baseline_record_count":
                        len(historical_records)
                })

        latest_checkin = (
            WellnessCheckIn.query
            .filter_by(
                personnel_id=personnel_id
            )
            .order_by(
                WellnessCheckIn.submitted_at.desc()
            )
            .first()
        )

        latest_scan = (
            WellnessScan.query
            .filter_by(
                personnel_id=personnel_id
            )
            .order_by(
                WellnessScan.saved_at.desc()
            )
            .first()
        )

        latest_support_request = (
            SupportRequest.query
            .filter_by(
                personnel_id=personnel_id
            )
            .order_by(
                SupportRequest.created_at.desc()
            )
            .first()
        )

        latest_leave_request = (
            LeaveRequest.query
            .filter_by(
                personnel_id=personnel_id
            )
            .order_by(
                LeaveRequest.created_at.desc()
            )
            .first()
        )

        dashboard_summary = {
            "period_start":
                period_start,
            "period_end":
                today,
            "record_count":
                len(current_records),
            "duty_total":
                current_duty_total,
            "rest_average":
                current_rest_average,
            "recovery_days":
                current_recovery_days,
            "training_total":
                current_training_total,
            "current_context":
                current_context,
            "deployment_days":
                deployment_days,
            "latest_leave_context":
                latest_leave_context,
            "attention_level":
                (
                    current_assessment.attention_level
                    if current_assessment
                    else "Insufficient Data"
                ),
            "assessment_version":
                (
                    current_assessment.assessment_version
                    if current_assessment
                    else None
                ),
            "evidence_status":
                (
                    current_assessment.evidence_status
                    if current_assessment
                    else "No Current Assessment"
                )
        }

        operational_factors = []

        if current_records:

            operational_factors.append({
                "title":
                    "Current Duty Window",
                "detail":
                    (
                        (
                            str(current_duty_total)
                            + " duty hours"
                        )
                        if current_duty_total is not None
                        else "Duty hours unavailable"
                    )
                    + " across "
                    + str(len(current_records))
                    + " stored record(s)."
            })

            operational_factors.append({
                "title":
                    "Recovery Pattern",
                "detail":
                    (
                        str(current_recovery_days)
                        + " recovery day(s) in the current seven-day window."
                    )
            })

            operational_factors.append({
                "title":
                    "Operational Context",
                "detail":
                    current_context
                    + (
                        " with active deployment recorded."
                        if latest_operational
                        and latest_operational.deployment_active
                        else "."
                    )
            })

            if (
                current_training_total is not None
                and current_training_total > 0
            ):

                operational_factors.append({
                    "title":
                        "Training Load",
                    "detail":
                        (
                            str(current_training_total)
                            + " training hour(s) are stored in the current window."
                        )
                })

            if leave_contexts:

                operational_factors.append({
                    "title":
                        "Leave Context",
                    "detail":
                        (
                            latest_leave_context
                            + " is present as supporting operational context."
                        )
                })

        latest_wellness_context = {
            "checkin":
                latest_checkin,
            "scan":
                latest_scan,
            "support":
                latest_support_request,
            "leave":
                latest_leave_request
        }

        current_analysis = (
            _current_welfare_analysis(
                personnel_id
            )
        )

        return render_template(
            "personnel_dashboard.html",
            personnel=personnel,
            current_assessment=
                current_assessment,
            current_analysis=
                current_analysis,
            current_records=
                current_records,
            dashboard_summary=
                dashboard_summary,
            baseline_summary=
                baseline_summary,
            operational_factors=
                operational_factors,
            latest_wellness_context=
                latest_wellness_context
        )


    # --------------------------------------------------
    # CURRENT PERSONNEL WELFARE ASSESSMENT API
    # --------------------------------------------------

    @app.route(
        "/api/personnel/current-assessment"
    )
    def personnel_current_assessment():

        if (
            session.get("role")
            != "personnel"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        personnel_id = session.get(
            "personnel_id"
        )

        current_assessment = (
            WelfareAssessment.query
            .filter_by(
                personnel_id=personnel_id,
                is_current=True
            )
            .order_by(
                WelfareAssessment.assessment_time.desc()
            )
            .first()
        )

        if current_assessment is None:

            current_assessment = WelfareAssessment(

                personnel_id=personnel_id,

                assessment_version=1,

                attention_level="Insufficient Data",

                evidence_status="Awaiting Sufficient Evidence",

                confidence=None,

                baseline_context=None,

                baseline_version=None,

                rules_version="VP-RULES-1.0",

                explanation_summary=(
                    "A welfare-attention category has not been produced "
                    "because sufficient validated operational evidence "
                    "has not yet been processed."
                ),

                is_current=True,

                assessment_time=datetime.now(),

                created_at=datetime.now()
            )

            try:

                db.session.add(
                    current_assessment
                )

                db.session.commit()

            except Exception:

                db.session.rollback()

                return jsonify({
                    "success": False,
                    "message":
                        "Could not create the current assessment record."
                }), 500

        return jsonify({

            "success": True,

            "assessment": {

                "id":
                    current_assessment.id,

                "personnel_id":
                    current_assessment.personnel_id,

                "assessment_version":
                    current_assessment.assessment_version,

                "attention_level":
                    current_assessment.attention_level,

                "evidence_status":
                    current_assessment.evidence_status,

                "confidence":
                    current_assessment.confidence,

                "baseline_context":
                    current_assessment.baseline_context,

                "baseline_version":
                    current_assessment.baseline_version,

                "rules_version":
                    current_assessment.rules_version,

                "explanation_summary":
                    current_assessment.explanation_summary,

                "assessment_time":
                    current_assessment.assessment_time.isoformat()
                    if current_assessment.assessment_time
                    else None
            }
        })


    # --------------------------------------------------
    # PROTOTYPE OPERATIONAL DATA SEED
    # --------------------------------------------------

    @app.route(
        "/api/personnel/seed-operational-history",
        methods=["POST"]
    )
    def seed_operational_history():

        if (
            session.get("role")
            != "personnel"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        personnel_id = session.get(
            "personnel_id"
        )

        existing_count = (
            OperationalRecord.query
            .filter_by(
                personnel_id=personnel_id
            )
            .count()
        )

        if existing_count > 0:

            return jsonify({
                "success": True,
                "message":
                    "Operational history already exists. No duplicate records were created.",
                "records_created": 0,
                "existing_records": existing_count
            })

        today = datetime.now().date()

        prototype_rows = [

            {
                "days_ago": 41,
                "duty_hours": 8.0,
                "rest_hours": 16.0,
                "context_type": "Station Duty",
                "deployment_active": False,
                "training_hours": 0.0,
                "leave_status": "None"
            },

            {
                "days_ago": 40,
                "duty_hours": 8.0,
                "rest_hours": 16.0,
                "context_type": "Station Duty",
                "deployment_active": False,
                "training_hours": 0.0,
                "leave_status": "None"
            },

            {
                "days_ago": 39,
                "duty_hours": 7.5,
                "rest_hours": 16.5,
                "context_type": "Station Duty",
                "deployment_active": False,
                "training_hours": 0.0,
                "leave_status": "None"
            },

            {
                "days_ago": 38,
                "duty_hours": 8.5,
                "rest_hours": 15.5,
                "context_type": "Station Duty",
                "deployment_active": False,
                "training_hours": 0.0,
                "leave_status": "None"
            },

            {
                "days_ago": 37,
                "duty_hours": 8.0,
                "rest_hours": 16.0,
                "context_type": "Station Duty",
                "deployment_active": False,
                "training_hours": 0.0,
                "leave_status": "None"
            },

            {
                "days_ago": 36,
                "duty_hours": 0.0,
                "rest_hours": 24.0,
                "context_type": "Station Duty",
                "deployment_active": False,
                "training_hours": 0.0,
                "leave_status": "Rest Day"
            },

            {
                "days_ago": 35,
                "duty_hours": 0.0,
                "rest_hours": 24.0,
                "context_type": "Station Duty",
                "deployment_active": False,
                "training_hours": 0.0,
                "leave_status": "Rest Day"
            },

            {
                "days_ago": 13,
                "duty_hours": 10.0,
                "rest_hours": 14.0,
                "context_type": "Deployment",
                "deployment_active": True,
                "training_hours": 0.0,
                "leave_status": "None"
            },

            {
                "days_ago": 12,
                "duty_hours": 10.5,
                "rest_hours": 13.5,
                "context_type": "Deployment",
                "deployment_active": True,
                "training_hours": 0.0,
                "leave_status": "None"
            },

            {
                "days_ago": 11,
                "duty_hours": 9.5,
                "rest_hours": 14.5,
                "context_type": "Deployment",
                "deployment_active": True,
                "training_hours": 0.0,
                "leave_status": "None"
            },

            {
                "days_ago": 10,
                "duty_hours": 10.0,
                "rest_hours": 14.0,
                "context_type": "Deployment",
                "deployment_active": True,
                "training_hours": 0.0,
                "leave_status": "None"
            },

            {
                "days_ago": 9,
                "duty_hours": 10.5,
                "rest_hours": 13.5,
                "context_type": "Deployment",
                "deployment_active": True,
                "training_hours": 0.0,
                "leave_status": "None"
            },

            {
                "days_ago": 8,
                "duty_hours": 0.0,
                "rest_hours": 24.0,
                "context_type": "Deployment",
                "deployment_active": True,
                "training_hours": 0.0,
                "leave_status": "Rest Day"
            },

            {
                "days_ago": 6,
                "duty_hours": 10.5,
                "rest_hours": 13.5,
                "context_type": "Deployment",
                "deployment_active": True,
                "training_hours": 0.0,
                "leave_status": "None"
            },

            {
                "days_ago": 5,
                "duty_hours": 11.0,
                "rest_hours": 13.0,
                "context_type": "Deployment",
                "deployment_active": True,
                "training_hours": 0.0,
                "leave_status": "None"
            },

            {
                "days_ago": 4,
                "duty_hours": 11.0,
                "rest_hours": 13.0,
                "context_type": "Deployment",
                "deployment_active": True,
                "training_hours": 0.0,
                "leave_status": "None"
            },

            {
                "days_ago": 3,
                "duty_hours": 10.5,
                "rest_hours": 13.5,
                "context_type": "Deployment",
                "deployment_active": True,
                "training_hours": 2.0,
                "leave_status": "None"
            },

            {
                "days_ago": 2,
                "duty_hours": 11.0,
                "rest_hours": 13.0,
                "context_type": "Deployment",
                "deployment_active": True,
                "training_hours": 2.0,
                "leave_status": "Leave Cancelled"
            },

            {
                "days_ago": 1,
                "duty_hours": 0.0,
                "rest_hours": 24.0,
                "context_type": "Deployment",
                "deployment_active": True,
                "training_hours": 0.0,
                "leave_status": "Rest Day"
            },

            {
                "days_ago": 0,
                "duty_hours": 0.0,
                "rest_hours": 24.0,
                "context_type": "Deployment",
                "deployment_active": True,
                "training_hours": 0.0,
                "leave_status": "Current Day"
            }
        ]

        created_records = []

        try:

            for index, row in enumerate(
                prototype_rows,
                start=1
            ):

                record = OperationalRecord(

                    personnel_id=personnel_id,

                    record_date=(
                        today -
                        timedelta(
                            days=row["days_ago"]
                        )
                    ),

                    duty_hours=row["duty_hours"],

                    rest_hours=row["rest_hours"],

                    context_type=row["context_type"],

                    deployment_active=row[
                        "deployment_active"
                    ],

                    training_hours=row[
                        "training_hours"
                    ],

                    leave_status=row[
                        "leave_status"
                    ],

                    source_type=
                        "Prototype Demonstration Data",

                    source_record_id=(
                        f"DEMO-{personnel_id}-{index:03d}"
                    ),

                    is_valid=True,

                    created_at=datetime.now()
                )

                db.session.add(
                    record
                )

                created_records.append(
                    record
                )

            db.session.commit()

        except Exception:

            db.session.rollback()

            return jsonify({
                "success": False,
                "message":
                    "Could not create prototype operational history."
            }), 500

        return jsonify({

            "success": True,

            "message":
                "Prototype operational history created successfully.",

            "records_created":
                len(created_records),

            "source_type":
                "Prototype Demonstration Data"
        })


    # --------------------------------------------------
    # PERSONNEL OPERATIONAL SUMMARY API
    # --------------------------------------------------

    @app.route(
        "/api/personnel/operational-summary"
    )
    def personnel_operational_summary():

        if (
            session.get("role")
            != "personnel"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        personnel_id = session.get(
            "personnel_id"
        )

        today = datetime.now().date()

        period_start = (
            today -
            timedelta(days=6)
        )

        records = (
            OperationalRecord.query
            .filter(
                OperationalRecord.personnel_id
                == personnel_id,

                OperationalRecord.is_valid
                == True,

                OperationalRecord.record_date
                >= period_start,

                OperationalRecord.record_date
                <= today
            )
            .order_by(
                OperationalRecord.record_date.asc()
            )
            .all()
        )

        if not records:

            return jsonify({

                "success": True,

                "summary": {

                    "personnel_id":
                        personnel_id,

                    "data_status":
                        "Insufficient Data",

                    "period_start":
                        period_start.isoformat(),

                    "period_end":
                        today.isoformat(),

                    "record_count":
                        0,

                    "total_duty_hours":
                        None,

                    "average_rest_hours":
                        None,

                    "total_training_hours":
                        None,

                    "leave_cancellations":
                        None,

                    "current_context":
                        None,

                    "deployment_active":
                        None,

                    "source_types":
                        []
                }
            })

        duty_values = [
            record.duty_hours
            for record in records
            if record.duty_hours is not None
        ]

        rest_values = [
            record.rest_hours
            for record in records
            if record.rest_hours is not None
        ]

        training_values = [
            record.training_hours
            for record in records
            if record.training_hours is not None
        ]

        total_duty_hours = round(
            sum(duty_values),
            1
        )

        average_rest_hours = (
            round(
                sum(rest_values)
                / len(rest_values),
                1
            )
            if rest_values
            else None
        )

        total_training_hours = round(
            sum(training_values),
            1
        )

        leave_cancellations = sum(
            1
            for record in records
            if (
                record.leave_status
                == "Leave Cancelled"
            )
        )

        latest_record = records[-1]

        source_types = sorted(
            list({
                record.source_type
                for record in records
                if record.source_type
            })
        )

        return jsonify({

            "success": True,

            "summary": {

                "personnel_id":
                    personnel_id,

                "data_status":
                    "Available",

                "period_start":
                    period_start.isoformat(),

                "period_end":
                    today.isoformat(),

                "record_count":
                    len(records),

                "total_duty_hours":
                    total_duty_hours,

                "average_rest_hours":
                    average_rest_hours,

                "total_training_hours":
                    total_training_hours,

                "leave_cancellations":
                    leave_cancellations,

                "current_context":
                    latest_record.context_type,

                "deployment_active":
                    latest_record.deployment_active,

                "source_types":
                    source_types
            }
        })


    # --------------------------------------------------
    # CONTEXT-AWARE PERSONAL BASELINE API
    # --------------------------------------------------

    @app.route(
        "/api/personnel/personal-baseline"
    )
    def personnel_personal_baseline():

        if (
            session.get("role")
            != "personnel"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        personnel_id = session.get(
            "personnel_id"
        )

        today = datetime.now().date()

        current_start = (
            today -
            timedelta(days=6)
        )

        current_records = (
            OperationalRecord.query
            .filter(
                OperationalRecord.personnel_id
                == personnel_id,

                OperationalRecord.is_valid
                == True,

                OperationalRecord.record_date
                >= current_start,

                OperationalRecord.record_date
                <= today
            )
            .order_by(
                OperationalRecord.record_date.asc()
            )
            .all()
        )

        if len(current_records) < 4:

            return jsonify({
                "success": True,
                "baseline": {
                    "status": "Insufficient Data",
                    "reason":
                        "At least 4 valid current-period records are required.",
                    "current_record_count":
                        len(current_records),
                    "baseline_record_count":
                        0
                }
            })

        current_context = (
            current_records[-1].context_type
            or "Unknown"
        )

        comparison_record_limit = (
            len(current_records)
        )

        historical_records = (
            OperationalRecord.query
            .filter(
                OperationalRecord.personnel_id
                == personnel_id,

                OperationalRecord.is_valid
                == True,

                OperationalRecord.record_date
                < current_start,

                OperationalRecord.context_type
                == current_context
            )
            .order_by(
                OperationalRecord.record_date.desc()
            )
            .limit(
                comparison_record_limit
            )
            .all()
        )

        historical_records.reverse()

        if len(historical_records) < 4:

            return jsonify({
                "success": True,
                "baseline": {
                    "status": "Insufficient Data",
                    "reason":
                        "Not enough comparable historical records exist for the current context.",
                    "context":
                        current_context,
                    "current_record_count":
                        len(current_records),
                    "baseline_record_count":
                        len(historical_records)
                }
            })

        def average(values):

            clean_values = [
                value
                for value in values
                if value is not None
            ]

            if not clean_values:
                return None

            return round(
                sum(clean_values)
                / len(clean_values),
                1
            )

        def total(values):

            clean_values = [
                value
                for value in values
                if value is not None
            ]

            if not clean_values:
                return None

            return round(
                sum(clean_values),
                1
            )

        current_worked_records = [
            record
            for record in current_records
            if (
                record.duty_hours is not None
                and record.duty_hours > 0
            )
        ]

        baseline_worked_records = [
            record
            for record in historical_records
            if (
                record.duty_hours is not None
                and record.duty_hours > 0
            )
        ]

        current_duty_avg = average([
            record.duty_hours
            for record in current_worked_records
        ])

        baseline_duty_avg = average([
            record.duty_hours
            for record in baseline_worked_records
        ])

        current_duty_total = total([
            record.duty_hours
            for record in current_records
        ])

        baseline_duty_total = total([
            record.duty_hours
            for record in historical_records
        ])

        current_rest_avg = average([
            record.rest_hours
            for record in current_records
        ])

        baseline_rest_avg = average([
            record.rest_hours
            for record in historical_records
        ])

        current_recovery_days = sum(
            1
            for record in current_records
            if (
                record.duty_hours is not None
                and record.duty_hours == 0
            )
        )

        baseline_recovery_days = sum(
            1
            for record in historical_records
            if (
                record.duty_hours is not None
                and record.duty_hours == 0
            )
        )

        current_training_total = total([
            record.training_hours
            for record in current_records
        ])

        baseline_training_avg = average([
            record.training_hours
            for record in historical_records
        ])

        duty_avg_delta = (
            round(
                current_duty_avg
                - baseline_duty_avg,
                1
            )
            if (
                current_duty_avg is not None
                and baseline_duty_avg is not None
            )
            else None
        )

        duty_total_delta = (
            round(
                current_duty_total
                - baseline_duty_total,
                1
            )
            if (
                current_duty_total is not None
                and baseline_duty_total is not None
            )
            else None
        )

        rest_delta = (
            round(
                current_rest_avg
                - baseline_rest_avg,
                1
            )
            if (
                current_rest_avg is not None
                and baseline_rest_avg is not None
            )
            else None
        )

        recovery_day_delta = (
            current_recovery_days
            - baseline_recovery_days
        )

        return jsonify({
            "success": True,
            "baseline": {
                "status": "Available",
                "personnel_id":
                    personnel_id,
                "context":
                    current_context,
                "baseline_version":
                    "VP-BASELINE-1.0",
                "comparison_method":
                    "Own comparable historical context using an equal-length record window",
                "current_period_start":
                    current_start.isoformat(),
                "current_period_end":
                    today.isoformat(),
                "current_record_count":
                    len(current_records),
                "baseline_record_count":
                    len(historical_records),
                "current": {
                    "average_duty_hours_on_worked_days":
                        current_duty_avg,
                    "total_duty_hours_in_window":
                        current_duty_total,
                    "worked_day_count":
                        len(current_worked_records),
                    "recovery_day_count":
                        current_recovery_days,
                    "average_rest_hours_per_record":
                        current_rest_avg,
                    "training_hours_total":
                        current_training_total
                },
                "historical_baseline": {
                    "average_duty_hours_on_worked_days":
                        baseline_duty_avg,
                    "total_duty_hours_in_comparison_set":
                        baseline_duty_total,
                    "worked_day_count":
                        len(baseline_worked_records),
                    "recovery_day_count":
                        baseline_recovery_days,
                    "average_rest_hours_per_record":
                        baseline_rest_avg,
                    "average_training_hours_per_record":
                        baseline_training_avg
                },
                "deviation": {
                    "average_duty_hours_on_worked_days":
                        duty_avg_delta,
                    "total_duty_hours":
                        duty_total_delta,
                    "rest_hours":
                        rest_delta,
                    "recovery_days":
                        recovery_day_delta
                },
                "guardrail_note":
                    "Context comparison does not excuse sustained recovery shortfalls. Rest days are kept separate from worked-day duty intensity, and current versus historical totals use equal-length record windows so unequal sample sizes do not distort the comparison."
            }
        })


    # --------------------------------------------------
    # EVIDENCE-GATED WELFARE ASSESSMENT ENGINE
    # --------------------------------------------------

    @app.route(
        "/api/personnel/recalculate-assessment",
        methods=["POST"]
    )
    def personnel_recalculate_assessment():

        if (
            session.get("role")
            != "personnel"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        personnel_id = session.get(
            "personnel_id"
        )

        today = datetime.now().date()

        current_start = (
            today -
            timedelta(days=6)
        )

        current_records = (
            OperationalRecord.query
            .filter(
                OperationalRecord.personnel_id
                == personnel_id,

                OperationalRecord.is_valid
                == True,

                OperationalRecord.record_date
                >= current_start,

                OperationalRecord.record_date
                <= today
            )
            .order_by(
                OperationalRecord.record_date.asc()
            )
            .all()
        )

        historical_records = []

        # ----------------------------------------------
        # DECISION REPLAY SNAPSHOT FIELDS
        # ----------------------------------------------

        snapshot_current_record_count = (
            len(current_records)
        )

        snapshot_baseline_record_count = None
        snapshot_duty_delta = None
        snapshot_rest_delta = None
        snapshot_recovery_day_delta = None
        snapshot_factor_count = None
        snapshot_severity_score = None

        # ----------------------------------------------
        # EVIDENCE GATE 1:
        # enough recent validated operational records
        # ----------------------------------------------

        if len(current_records) < 4:

            attention_level = "Insufficient Data"

            evidence_status = (
                "Awaiting Sufficient Evidence"
            )

            confidence = None

            explanation_summary = (
                "Assessment withheld because fewer than four "
                "validated current-period operational records "
                "are available."
            )

            baseline_context = None

            baseline_version = None

        else:

            current_context = (
                current_records[-1].context_type
                or "Unknown"
            )

            comparison_record_limit = (
                len(current_records)
            )

            # Personal operational baseline:
            # use the immediately preceding equal-length block of
            # validated records. Recovery Day records are deliberately
            # retained so recovery-day availability can be compared.
            historical_records = (
                OperationalRecord.query
                .filter(
                    OperationalRecord.personnel_id
                    == personnel_id,

                    OperationalRecord.is_valid
                    == True,

                    OperationalRecord.record_date
                    < current_start
                )
                .order_by(
                    OperationalRecord.record_date.desc()
                )
                .limit(
                    comparison_record_limit
                )
                .all()
            )

            historical_records.reverse()

            snapshot_baseline_record_count = (
                len(historical_records)
            )

            current_worked_records = [
                record
                for record in current_records
                if (
                    record.duty_hours is not None
                    and record.duty_hours > 0
                )
            ]

            baseline_worked_records = [
                record
                for record in historical_records
                if (
                    record.duty_hours is not None
                    and record.duty_hours > 0
                )
            ]

            # ------------------------------------------
            # EVIDENCE GATE 2:
            # equal-length comparable context + enough
            # worked-day observations
            # ------------------------------------------

            evidence_is_sufficient = (
                len(historical_records)
                == comparison_record_limit

                and len(historical_records) >= 4

                and len(current_worked_records) >= 3

                and len(baseline_worked_records) >= 3
            )

            if not evidence_is_sufficient:

                attention_level = (
                    "Insufficient Data"
                )

                evidence_status = (
                    "Awaiting Comparable Evidence"
                )

                confidence = None

                baseline_context = (
                    "Personal Operational Baseline"
                )

                baseline_version = None

                explanation_summary = (
                    "Assessment withheld because there are not "
                    "enough equal-length validated historical "
                    "records for the personal operational baseline."
                )

            else:

                def average(values):

                    clean_values = [
                        value
                        for value in values
                        if value is not None
                    ]

                    if not clean_values:
                        return None

                    return round(
                        sum(clean_values)
                        / len(clean_values),
                        1
                    )

                current_duty_avg = average([
                    record.duty_hours
                    for record
                    in current_worked_records
                ])

                baseline_duty_avg = average([
                    record.duty_hours
                    for record
                    in baseline_worked_records
                ])

                current_rest_avg = average([
                    record.rest_hours
                    for record
                    in current_records
                ])

                baseline_rest_avg = average([
                    record.rest_hours
                    for record
                    in historical_records
                ])

                current_recovery_days = sum(
                    1
                    for record in current_records
                    if (
                        record.duty_hours is not None
                        and record.duty_hours == 0
                    )
                )

                baseline_recovery_days = sum(
                    1
                    for record in historical_records
                    if (
                        record.duty_hours is not None
                        and record.duty_hours == 0
                    )
                )

                _t_vals = [r.training_hours for r in current_records if r.training_hours is not None]
                current_training_hours = round(sum(_t_vals), 1) if _t_vals else None

                leave_cancellations = sum(
                    1
                    for record in current_records
                    if (
                        record.leave_status
                        == "Leave Cancelled"
                    )
                )

                duty_delta = round(
                    current_duty_avg
                    - baseline_duty_avg,
                    1
                )

                rest_delta = round(
                    current_rest_avg
                    - baseline_rest_avg,
                    1
                )

                recovery_delta = (
                    current_recovery_days
                    - baseline_recovery_days
                )

                snapshot_duty_delta = duty_delta
                snapshot_rest_delta = rest_delta
                snapshot_recovery_day_delta = (
                    recovery_delta
                )

                # --------------------------------------
                # TRANSPARENT PROTOTYPE RULES
                #
                # Important:
                # - no single factor can elevate attention
                # - transfer/deployment/training/leave are
                #   context and never standalone triggers
                # --------------------------------------

                concerning_factors = []

                severity_score = 0

                if duty_delta >= 2.0:

                    concerning_factors.append(
                        "Worked-day duty intensity is substantially above the comparable personal baseline."
                    )

                    severity_score += 2

                elif duty_delta >= 1.0:

                    concerning_factors.append(
                        "Worked-day duty intensity is moderately above the comparable personal baseline."
                    )

                    severity_score += 1

                if rest_delta <= -2.0:

                    concerning_factors.append(
                        "Average rest is substantially below the comparable personal baseline."
                    )

                    severity_score += 2

                elif rest_delta <= -1.0:

                    concerning_factors.append(
                        "Average rest is moderately below the comparable personal baseline."
                    )

                    severity_score += 1

                if recovery_delta <= -2:

                    concerning_factors.append(
                        "Recovery-day availability is substantially below the comparable personal baseline."
                    )

                    severity_score += 2

                elif recovery_delta <= -1:

                    concerning_factors.append(
                        "Recovery-day availability is below the comparable personal baseline."
                    )

                    severity_score += 1

                # --------------------------------------
                # MULTI-FACTOR GUARDRAIL
                # --------------------------------------

                factor_count = len(
                    concerning_factors
                )

                snapshot_factor_count = (
                    factor_count
                )

                snapshot_severity_score = (
                    severity_score
                )

                if (
                    factor_count >= 3
                    and severity_score >= 4
                ):

                    attention_level = (
                        "Elevated"
                    )

                elif factor_count >= 2:

                    attention_level = (
                        "Moderate"
                    )

                else:

                    attention_level = (
                        "Low"
                    )

                evidence_status = (
                    "Sufficient Operational Evidence"
                )

                confidence = (
                    "Medium"
                )

                baseline_context = (
                    "Personal Operational Baseline"
                )

                baseline_version = (
                    "VP-BASELINE-1.1"
                )

                # --------------------------------------
                # EXPLAIN WHY
                # --------------------------------------

                if factor_count == 0:

                    explanation_summary = (
                        "Current validated operational evidence does "
                        "not show a sustained multi-factor deterioration "
                        "relative to the person's comparable "
                        "personal operational baseline. "
                        f"Worked-day duty changed by {duty_delta:+.1f} hours, "
                        f"average rest changed by {rest_delta:+.1f} hours, "
                        f"and recovery-day availability changed by "
                        f"{recovery_delta:+d} day(s). "
                        "Training and leave records are retained as "
                        "supporting operational context and do not "
                        "independently increase welfare attention."
                    )

                elif factor_count == 1:

                    explanation_summary = (
                        "One operational deviation is present, but "
                        "VeerPulse does not elevate welfare attention "
                        "from a single factor alone. "
                        + concerning_factors[0]
                    )

                else:

                    explanation_summary = (
                        "Multiple sustained operational deviations are "
                        "present relative to the person's comparable "
                        "personal operational baseline: "
                        + " ".join(concerning_factors)
                    )

                if current_training_hours > 0:

                    explanation_summary += (
                        f" {current_training_hours:.1f} training hour(s) "
                        "are recorded in the current window as supporting "
                        "context."
                    )

                if leave_cancellations > 0:

                    explanation_summary += (
                        f" {leave_cancellations} leave cancellation(s) "
                        "are recorded as supporting context."
                    )


        # ----------------------------------------------
        # OPTIONAL VOLUNTARY CHECK-IN CONTEXT
        #
        # Guardrail:
        # - the check-in contributes only when the operational
        #   evidence gate is satisfied
        # - self-report alone cannot produce Elevated attention
        # - confidential free-text is not copied into the
        #   assessment explanation
        # - employment / leave / disciplinary decisions
        #   remain outside this assessment engine
        # ----------------------------------------------

        recent_checkin = (
            _recent_wellness_checkin(
                personnel_id,
                days=7
            )
        )

        if recent_checkin:

            if (
                evidence_status
                == "Sufficient Operational Evidence"
            ):

                evidence_status = (
                    "Operational + Check-In Context"
                )

            explanation_summary += (
                " A voluntary Wellness Check-In submitted within "
                "the last seven days is linked as supporting context. "
                + recent_checkin.summary_label
                + " were recorded. "
                "Self-report is weighted as supporting evidence and "
                "cannot produce Elevated attention without "
                "meaningful operational concern."
            )

            if (
                recent_checkin.support_request
                == "Yes"
            ):

                explanation_summary += (
                    " The personnel member also requested confidential "
                    "human welfare support; that request is routed separately "
                    "and does not change this assessment level."
                )


        # ----------------------------------------------
        # OPTIONAL CAMERA WELLNESS SCAN CONTEXT
        #
        # Guardrail:
        # - only a recently saved usable scan is linked
        # - HR / breathing are experimental wellness-support
        #   measurements, not clinical measurements
        # - the scan never independently sets or raises
        #   Low / Moderate / Elevated attention
        # ----------------------------------------------

        recent_scan = (
            _recent_usable_wellness_scan(
                personnel_id,
                hours=24
            )
        )

        if recent_scan:

            if (
                evidence_status
                == "Sufficient Operational Evidence"
            ):

                evidence_status = (
                    "Operational + Scan Context"
                )

            elif (
                evidence_status
                == "Operational + Check-In Context"
            ):

                evidence_status = (
                    "Operational + Check-In + Scan Context"
                )

            explanation_summary += (
                " A usable voluntary camera Wellness Scan saved within "
                "the last 24 hours is linked as optional supporting context. "
                "Camera-derived pulse and breathing estimates did not "
                "independently change the welfare-attention level."
            )

        # ----------------------------------------------
        # VERSIONED SINGLE SOURCE OF TRUTH
        # ----------------------------------------------

        latest_assessment = (
            WelfareAssessment.query
            .filter_by(
                personnel_id=personnel_id
            )
            .order_by(
                WelfareAssessment.assessment_version.desc()
            )
            .first()
        )

        next_version = (
            latest_assessment.assessment_version + 1
            if latest_assessment
            else 1
        )

        try:

            (
                WelfareAssessment.query
                .filter_by(
                    personnel_id=personnel_id,
                    is_current=True
                )
                .update({
                    "is_current": False
                })
            )

            new_assessment = WelfareAssessment(

                personnel_id=personnel_id,

                assessment_version=
                    next_version,

                attention_level=
                    attention_level,

                evidence_status=
                    evidence_status,

                confidence=
                    confidence,

                baseline_context=
                    baseline_context,

                baseline_version=
                    baseline_version,

                rules_version=
                    "VP-RULES-1.3",

                explanation_summary=
                    explanation_summary,

                is_current=True,

                assessment_time=
                    datetime.now(),

                created_at=
                    datetime.now()
            )

            db.session.add(
                new_assessment
            )

            # Flush first so the versioned assessment receives
            # its database ID before the replay snapshot is saved.
            db.session.flush()

            evidence_snapshot = (
                AssessmentEvidenceSnapshot(

                    assessment_id=
                        new_assessment.id,

                    personnel_id=
                        personnel_id,

                    assessment_version=
                        new_assessment.assessment_version,

                    rules_version=
                        new_assessment.rules_version,

                    baseline_version=
                        new_assessment.baseline_version,

                    baseline_context=
                        new_assessment.baseline_context,

                    evidence_status=
                        new_assessment.evidence_status,

                    current_record_count=
                        snapshot_current_record_count,

                    baseline_record_count=
                        snapshot_baseline_record_count,

                    duty_delta=
                        snapshot_duty_delta,

                    rest_delta=
                        snapshot_rest_delta,

                    recovery_day_delta=
                        snapshot_recovery_day_delta,

                    concerning_factor_count=
                        snapshot_factor_count,

                    severity_score=
                        snapshot_severity_score,

                    resulting_attention_level=
                        new_assessment.attention_level,

                    explanation_summary=
                        new_assessment.explanation_summary,

                    created_at=
                        datetime.now()
                )
            )

            db.session.add(
                evidence_snapshot
            )

            checkin_evidence = None

            if recent_checkin:

                checkin_evidence = (
                    AssessmentCheckInEvidence(

                        assessment_id=
                            new_assessment.id,

                        personnel_id=
                            personnel_id,

                        assessment_version=
                            new_assessment.assessment_version,

                        checkin_id=
                            recent_checkin.id,

                        checkin_submitted_at=
                            recent_checkin.submitted_at,

                        concern_score=
                            recent_checkin.concern_score,

                        concern_count=
                            recent_checkin.concern_count,

                        high_concern_count=
                            recent_checkin.high_concern_count,

                        support_requested=
                            (
                                recent_checkin.support_request
                                == "Yes"
                            ),

                        summary_label=
                            recent_checkin.summary_label,

                        created_at=
                            datetime.now()
                    )
                )

                db.session.add(
                    checkin_evidence
                )

            scan_evidence = None

            if recent_scan:

                scan_evidence = (
                    AssessmentScanEvidence(

                        assessment_id=
                            new_assessment.id,

                        personnel_id=
                            personnel_id,

                        assessment_version=
                            new_assessment.assessment_version,

                        scan_id=
                            recent_scan.id,

                        scan_saved_at=
                            recent_scan.saved_at,

                        heart_rate=
                            recent_scan.heart_rate,

                        breathing_rate=
                            recent_scan.breathing_rate,

                        signal_quality=
                            recent_scan.signal_quality,

                        confidence=
                            recent_scan.confidence,

                        created_at=
                            datetime.now()
                    )
                )

                db.session.add(
                    scan_evidence
                )

            # ------------------------------------------
            # VEERPULSE EARLY-WARNING ENGINE
            # ------------------------------------------

            analysis_result = (
                _build_veerpulse_analysis(
                    personnel_id=personnel_id,
                    assessment=new_assessment,
                    evidence_snapshot=evidence_snapshot,
                    recent_checkin=recent_checkin,
                    checkin_evidence=checkin_evidence,
                    scan_evidence=scan_evidence,
                    current_records=current_records,
                    historical_records=historical_records
                )
            )

            new_assessment.attention_level = (
                analysis_result.attention_level
            )

            new_assessment.confidence = (
                analysis_result.confidence
            )

            new_assessment.rules_version = (
                analysis_result.rules_version
            )

            new_assessment.explanation_summary = (
                analysis_result.analysis_reason
            )

            evidence_snapshot.rules_version = (
                analysis_result.rules_version
            )

            evidence_snapshot.resulting_attention_level = (
                analysis_result.attention_level
            )

            evidence_snapshot.explanation_summary = (
                analysis_result.analysis_reason
            )

            audit_event = AuditEvent(
                actor_role=
                    session.get("role")
                    or "personnel",

                actor_id=
                    session.get("personnel_id")
                    or session.get("user_id")
                    or personnel_id,

                event_type=
                    "Welfare Assessment Recalculated",

                personnel_id=
                    personnel_id,

                resource_type=
                    "WelfareAssessment",

                resource_id=
                    str(new_assessment.id),

                event_summary=
                    (
                        "Assessment version "
                        + str(
                            new_assessment.assessment_version
                        )
                        + " recalculated with result "
                        + str(
                            new_assessment.attention_level
                        )
                        + " using rules "
                        + str(
                            new_assessment.rules_version
                        )
                        + "."
                    ),

                source_type=
                    "VeerPulse Application",

                created_at=
                    datetime.now()
            )

            db.session.add(
                audit_event
            )

            db.session.commit()

        except Exception:

            db.session.rollback()

            return jsonify({
                "success": False,
                "message":
                    "Assessment could not be recalculated."
            }), 500

        return jsonify({

            "success": True,

            "message":
                "Assessment recalculated successfully.",

            "assessment": {

                "id":
                    new_assessment.id,

                "personnel_id":
                    new_assessment.personnel_id,

                "assessment_version":
                    new_assessment.assessment_version,

                "attention_level":
                    new_assessment.attention_level,

                "evidence_status":
                    new_assessment.evidence_status,

                "confidence":
                    new_assessment.confidence,

                "baseline_context":
                    new_assessment.baseline_context,

                "baseline_version":
                    new_assessment.baseline_version,

                "rules_version":
                    new_assessment.rules_version,

                "evidence_snapshot_id":
                    evidence_snapshot.id,

                "checkin_evidence_id":
                    (
                        checkin_evidence.id
                        if checkin_evidence
                        else None
                    ),

                "recent_checkin_id":
                    (
                        recent_checkin.id
                        if recent_checkin
                        else None
                    ),

                "scan_evidence_id":
                    (
                        scan_evidence.id
                        if scan_evidence
                        else None
                    ),

                "recent_scan_id":
                    (
                        recent_scan.id
                        if recent_scan
                        else None
                    ),

                "explanation_summary":
                    new_assessment.explanation_summary,

                "analysis_result_id":
                    analysis_result.id,

                "early_warning_score":
                    analysis_result.overall_score,

                "operational_score":
                    analysis_result.operational_score,

                "checkin_score":
                    analysis_result.checkin_score,

                "trend_score":
                    analysis_result.trend_score,

                "trend_direction":
                    analysis_result.trend_direction,

                "trend_strength":
                    analysis_result.trend_strength,

                "signal_agreement":
                    analysis_result.signal_agreement,

                "signal_agreement_score":
                    (
                        agreement_result[
                            "agreement_score"
                        ]
                    ),

                "signal_agreement_direction":
                    (
                        agreement_result[
                            "direction"
                        ]
                    ),

                "signal_count":
                    (
                        agreement_result[
                            "signal_count"
                        ]
                    ),

                "analysis_confidence":
                    analysis_result.confidence,

                "recommendation":
                    analysis_result.recommendation,

                "assessment_time":
                    new_assessment.assessment_time.isoformat()
            }
        })


    # --------------------------------------------------
    # PERSONNEL PROFILE
    # --------------------------------------------------

    @app.route(
        "/personnel-profile"
    )
    def personnel_profile():

        if (
            session.get("role")
            != "personnel"
        ):

            return redirect(
                url_for(
                    "personnel_login"
                )
            )

        personnel = (
            Personnel.query
            .filter_by(
                personnel_id=session.get(
                    "personnel_id"
                )
            )
            .first()
        )

        return render_template(
            "personnel_profile.html",
            personnel=personnel
        )


    # --------------------------------------------------
    # DUTY & LEAVE
    # --------------------------------------------------

    @app.route(
        "/personnel-duty-leave"
    )
    def personnel_duty_leave():

        if (
            session.get("role")
            != "personnel"
        ):

            return redirect(
                url_for(
                    "personnel_login"
                )
            )

        personnel_id = (
            session.get(
                "personnel_id"
            )
            or ""
        ).strip()

        if not personnel_id:

            session.clear()

            return redirect(
                url_for(
                    "personnel_login"
                )
            )

        personnel = (
            Personnel.query
            .filter_by(
                personnel_id=personnel_id
            )
            .first()
        )

        if personnel is None:

            session.clear()

            return redirect(
                url_for(
                    "personnel_login"
                )
            )

        return render_template(
            "personnel_duty_leave.html",
            personnel=personnel
        )


    # --------------------------------------------------
    # PERSONNEL LEAVE REQUESTS API
    # --------------------------------------------------

    @app.route(
        "/api/personnel/leave-requests",
        methods=["GET", "POST"]
    )
    def personnel_leave_requests_api():

        if (
            session.get("role")
            != "personnel"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        personnel_id = session.get(
            "personnel_id"
        )

        personnel = (
            Personnel.query
            .filter_by(
                personnel_id=personnel_id
            )
            .first()
        )

        if (
            not personnel
            or personnel.employment_status
            != "Active"
        ):

            return jsonify({
                "success": False,
                "message":
                    "Personnel account is not active."
            }), 403

        if request.method == "GET":

            leave_items = (
                LeaveRequest.query
                .filter_by(
                    personnel_id=personnel_id
                )
                .order_by(
                    LeaveRequest.created_at.desc()
                )
                .all()
            )

            return jsonify({
                "success": True,
                "leave_requests": [
                    _leave_to_dict(item)
                    for item in leave_items
                ]
            })

        data = request.get_json(
            silent=True
        ) or {}

        leave_type = (
            data.get("leave_type")
            or ""
        ).strip()

        start_date = _parse_date(
            data.get("start_date")
        )

        end_date = _parse_date(
            data.get("end_date")
        )

        reason_category = (
            data.get("reason_category")
            or ""
        ).strip()

        reason = (
            data.get("reason")
            or ""
        ).strip()

        additional_note = (
            data.get("additional_note")
            or ""
        ).strip() or None

        allowed_leave_types = [
            "Annual Leave",
            "Casual Leave",
            "Medical Leave",
            "Emergency Leave",
            "Compensatory Leave",
            "Other"
        ]

        allowed_reason_categories = [
            "Personal",
            "Family",
            "Medical",
            "Emergency",
            "Official",
            "Other"
        ]

        if leave_type not in allowed_leave_types:

            return jsonify({
                "success": False,
                "message":
                    "Please select a valid leave type."
            }), 400

        if reason_category not in allowed_reason_categories:

            return jsonify({
                "success": False,
                "message":
                    "Please select a valid reason category."
            }), 400

        if not start_date or not end_date:

            return jsonify({
                "success": False,
                "message":
                    "Start date and end date must use YYYY-MM-DD."
            }), 400

        if end_date < start_date:

            return jsonify({
                "success": False,
                "message":
                    "End date cannot be before start date."
            }), 400

        if start_date < datetime.now().date():

            return jsonify({
                "success": False,
                "message":
                    "A new leave request cannot start in the past."
            }), 400

        requested_days = (
            end_date - start_date
        ).days + 1

        if requested_days > 90:

            return jsonify({
                "success": False,
                "message":
                    "A single leave request cannot exceed 90 days in this prototype."
            }), 400

        if len(reason) < 3:

            return jsonify({
                "success": False,
                "message":
                    "Please provide a short reason for the leave request."
            }), 400

        overlapping = (
            LeaveRequest.query
            .filter(
                LeaveRequest.personnel_id
                == personnel_id,

                LeaveRequest.status.in_([
                    "Pending",
                    "Approved"
                ]),

                LeaveRequest.start_date
                <= end_date,

                LeaveRequest.end_date
                >= start_date
            )
            .first()
        )

        if overlapping:

            return jsonify({
                "success": False,
                "message":
                    "Another pending or approved leave request overlaps these dates."
            }), 409

        leave_item = LeaveRequest(
            personnel_id=personnel_id,
            leave_type=leave_type,
            start_date=start_date,
            end_date=end_date,
            requested_days=requested_days,
            reason_category=reason_category,
            reason=reason,
            additional_note=additional_note,
            status="Pending",
            created_at=datetime.now(),
            updated_at=datetime.now()
        )

        try:

            db.session.add(
                leave_item
            )

            db.session.flush()

            leave_item.request_code = (
                "LR-"
                + str(leave_item.id).zfill(6)
            )

            _add_audit_event(
                actor_role="personnel",
                actor_id=personnel_id,
                event_type="Leave Request Submitted",
                personnel_id=personnel_id,
                resource_type="LeaveRequest",
                resource_id=leave_item.id,
                event_summary=(
                    leave_item.request_code
                    + " submitted for "
                    + str(requested_days)
                    + " day(s)."
                )
            )

            db.session.commit()

        except Exception:

            db.session.rollback()

            return jsonify({
                "success": False,
                "message":
                    "The leave request could not be saved."
            }), 500

        return jsonify({
            "success": True,
            "message":
                "Leave request submitted successfully.",
            "leave_request":
                _leave_to_dict(leave_item)
        }), 201


    @app.route(
        "/api/personnel/leave-requests/<int:leave_id>/cancel",
        methods=["POST"]
    )
    def cancel_personnel_leave_request(
        leave_id
    ):

        if (
            session.get("role")
            != "personnel"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        personnel_id = session.get(
            "personnel_id"
        )

        leave_item = (
            LeaveRequest.query
            .filter_by(
                id=leave_id,
                personnel_id=personnel_id
            )
            .first()
        )

        if not leave_item:

            return jsonify({
                "success": False,
                "message":
                    "Leave request not found."
            }), 404

        if leave_item.status != "Pending":

            return jsonify({
                "success": False,
                "message":
                    "Only a pending leave request can be cancelled directly."
            }), 409

        leave_item.status = "Cancelled"
        leave_item.cancelled_at = datetime.now()
        leave_item.updated_at = datetime.now()

        _add_audit_event(
            actor_role="personnel",
            actor_id=personnel_id,
            event_type="Leave Request Cancelled",
            personnel_id=personnel_id,
            resource_type="LeaveRequest",
            resource_id=leave_item.id,
            event_summary=(
                leave_item.request_code
                + " cancelled by personnel before officer decision."
            )
        )

        try:
            db.session.commit()
        except Exception:
            db.session.rollback()
            return jsonify({
                "success": False,
                "message":
                    "The leave request could not be cancelled."
            }), 500

        return jsonify({
            "success": True,
            "message":
                "Leave request cancelled successfully.",
            "leave_request":
                _leave_to_dict(leave_item)
        })


    @app.route(
        "/api/personnel/leave-metrics",
        methods=["GET"]
    )
    def personnel_leave_metrics():

        if (
            session.get("role")
            != "personnel"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        personnel_id = session.get(
            "personnel_id"
        )

        try:
            window_days = int(
                request.args.get(
                    "days",
                    90
                )
            )
        except ValueError:
            window_days = 90

        window_days = max(
            30,
            min(window_days, 365)
        )

        window_start = (
            datetime.now()
            - timedelta(days=window_days)
        )

        items = (
            LeaveRequest.query
            .filter(
                LeaveRequest.personnel_id
                == personnel_id,
                LeaveRequest.created_at
                >= window_start
            )
            .order_by(
                LeaveRequest.created_at.asc()
            )
            .all()
        )

        counts = {
            "Pending": 0,
            "Approved": 0,
            "Denied": 0,
            "Cancelled": 0
        }

        for item in items:
            if item.status in counts:
                counts[item.status] += 1

        approved_days = sum(
            item.requested_days
            for item in items
            if item.status == "Approved"
        )

        gaps = []

        for index in range(1, len(items)):
            gaps.append(
                (
                    items[index].created_at
                    - items[index - 1].created_at
                ).total_seconds()
                / 86400
            )

        average_gap = (
            round(
                sum(gaps) / len(gaps),
                1
            )
            if gaps
            else None
        )

        total = len(items)

        return jsonify({
            "success": True,
            "metrics": {
                "personnel_id": personnel_id,
                "window_days": window_days,
                "applications": total,
                "pending": counts["Pending"],
                "approved": counts["Approved"],
                "denied": counts["Denied"],
                "cancelled": counts["Cancelled"],
                "approved_leave_days": approved_days,
                "applications_per_30_days": (
                    round(
                        total * 30 / window_days,
                        2
                    )
                    if window_days
                    else 0
                ),
                "average_days_between_requests":
                    average_gap,
                "denial_rate_percent": (
                    round(
                        counts["Denied"]
                        * 100
                        / total,
                        1
                    )
                    if total
                    else 0
                ),
                "cancellation_rate_percent": (
                    round(
                        counts["Cancelled"]
                        * 100
                        / total,
                        1
                    )
                    if total
                    else 0
                ),
                "interpretation_note": (
                    "Leave activity is reported as operational context only. "
                    "It must not independently determine welfare attention, "
                    "discipline, or future leave approval."
                )
            }
        })


    # --------------------------------------------------
    # MY TRENDS
    # --------------------------------------------------

    @app.route(
        "/personnel-trends"
    )
    def personnel_trends():

        if (
            session.get("role")
            != "personnel"
        ):

            return redirect(
                url_for(
                    "personnel_login"
                )
            )

        personnel_id = session.get(
            "personnel_id"
        )

        personnel = (
            Personnel.query
            .filter_by(
                personnel_id=personnel_id
            )
            .first()
        )

        current_assessment = (
            WelfareAssessment.query
            .filter_by(
                personnel_id=personnel_id,
                is_current=True
            )
            .order_by(
                WelfareAssessment.assessment_time.desc()
            )
            .first()
        )

        wellness_checkin_history = (
            WellnessCheckIn.query
            .filter_by(
                personnel_id=personnel_id
            )
            .order_by(
                WellnessCheckIn.submitted_at.desc()
            )
            .limit(10)
            .all()
        )

        latest_checkin = (
            wellness_checkin_history[0]
            if wellness_checkin_history
            else None
        )

        wellness_scan_history = (
            WellnessScan.query
            .filter_by(
                personnel_id=personnel_id
            )
            .order_by(
                WellnessScan.saved_at.desc()
            )
            .limit(10)
            .all()
        )

        current_analysis = (
            _current_welfare_analysis(
                personnel_id
            )
        )

        analysis_factors = []

        if current_analysis:

            analysis_factors = (
                AnalysisFactor.query
                .filter_by(
                    analysis_result_id=
                        current_analysis.id
                )
                .order_by(
                    AnalysisFactor.id.asc()
                )
                .all()
            )

        analysis_history = (
            WelfareAnalysisResult.query
            .filter_by(
                personnel_id=
                    personnel_id
            )
            .order_by(
                WelfareAnalysisResult.created_at.desc()
            )
            .limit(5)
            .all()
        )

        analysis_history.reverse()

        return render_template(
            "personnel_trends.html",
            personnel=personnel,
            current_assessment=current_assessment,
            current_analysis=current_analysis,
            analysis_factors=analysis_factors,
            analysis_history=analysis_history,
            latest_checkin=latest_checkin,
            wellness_checkin_history=
                wellness_checkin_history,
            wellness_scan_history=
                wellness_scan_history
        )


    # --------------------------------------------------
    # WELLNESS CHECK-IN
    # --------------------------------------------------

    @app.route(
        "/personnel-wellness-checkin"
    )
    def personnel_wellness_checkin():

        if (
            session.get("role")
            != "personnel"
        ):

            return redirect(
                url_for(
                    "personnel_login"
                )
            )

        personnel_id = session.get(
            "personnel_id"
        )

        personnel = (
            Personnel.query
            .filter_by(
                personnel_id=personnel_id
            )
            .first()
        )

        recent_checkins = (
            WellnessCheckIn.query
            .filter_by(
                personnel_id=personnel_id
            )
            .order_by(
                WellnessCheckIn.submitted_at.desc()
            )
            .limit(5)
            .all()
        )

        return render_template(
            "personnel_wellness_checkin.html",
            personnel=personnel,
            recent_checkins=recent_checkins
        )


    # --------------------------------------------------
    # SAVE VOLUNTARY WELLNESS CHECK-IN
    # --------------------------------------------------

    @app.route(
        "/api/personnel/wellness-checkins",
        methods=["POST"]
    )
    def save_personnel_wellness_checkin():

        if (
            session.get("role")
            != "personnel"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        data = request.get_json(
            silent=True
        )

        if not data:

            return jsonify({
                "success": False,
                "message":
                    "No wellness check-in data was received."
            }), 400

        required_fields = (
            _CHECKIN_NEGATIVE_FIELDS
            + _CHECKIN_POSITIVE_FIELDS
        )

        cleaned = {}

        for field_name in required_fields:

            value = (
                str(
                    data.get(
                        field_name,
                        ""
                    )
                )
                .strip()
            )

            if (
                value
                not in _CHECKIN_ALLOWED_RESPONSES
            ):

                return jsonify({
                    "success": False,
                    "message":
                        "Please answer every wellness question before submitting."
                }), 400

            cleaned[
                field_name
            ] = value

        support_request = (
            str(
                data.get(
                    "support_request",
                    "Not Answered"
                )
            )
            .strip()
            or "Not Answered"
        )

        if support_request not in {
            "No",
            "Maybe Later",
            "Yes",
            "Not Answered"
        }:

            return jsonify({
                "success": False,
                "message":
                    "Invalid support-request selection."
            }), 400

        optional_note = (
            str(
                data.get(
                    "optional_note",
                    ""
                )
            )
            .strip()
        )

        if len(optional_note) > 1000:

            return jsonify({
                "success": False,
                "message":
                    "The optional note must be 1000 characters or fewer."
            }), 400

        score_summary = (
            _score_wellness_checkin(
                cleaned
            )
        )

        personnel_id = session.get(
            "personnel_id"
        )

        checkin = WellnessCheckIn(

            personnel_id=
                personnel_id,

            workload_pressure=
                cleaned[
                    "workload_pressure"
                ],

            recovery_time=
                cleaned[
                    "recovery_time"
                ],

            duty_adjustment=
                cleaned[
                    "duty_adjustment"
                ],

            rest_quality=
                cleaned[
                    "rest_quality"
                ],

            low_energy=
                cleaned[
                    "low_energy"
                ],

            family_connection=
                cleaned[
                    "family_connection"
                ],

            family_separation=
                cleaned[
                    "family_separation"
                ],

            focus=
                cleaned[
                    "focus"
                ],

            mental_tiredness=
                cleaned[
                    "mental_tiredness"
                ],

            pattern_change=
                cleaned[
                    "pattern_change"
                ],

            support_request=
                support_request,

            optional_note=
                (
                    optional_note
                    if optional_note
                    else None
                ),

            concern_score=
                score_summary[
                    "concern_score"
                ],

            concern_count=
                score_summary[
                    "concern_count"
                ],

            high_concern_count=
                score_summary[
                    "high_concern_count"
                ],

            summary_label=
                score_summary[
                    "summary_label"
                ],

            submitted_at=
                datetime.now()
        )

        support_request_created = False

        try:

            db.session.add(
                checkin
            )

            db.session.flush()

            # A "Yes" response routes a generic human-support request.
            # The confidential free-text note is NOT copied into that request.
            # It also does not automatically change the welfare attention level.
            if support_request == "Yes":

                existing_support = (
                    SupportRequest.query
                    .filter(
                        SupportRequest.personnel_id
                        == personnel_id,

                        SupportRequest.support_category
                        == "Personal Wellness Check",

                        SupportRequest.status.in_(
                            [
                                "Open",
                                "In Review",
                                "In Progress"
                            ]
                        )
                    )
                    .order_by(
                        SupportRequest.created_at.desc()
                    )
                    .first()
                )

                if not existing_support:

                    support_item = SupportRequest(

                        personnel_id=
                            personnel_id,

                        support_category=
                            "Personal Wellness Check",

                        contact_method=
                            "Secure Message",

                        priority=
                            "Soon",

                        message=
                            (
                                "Personnel voluntarily requested confidential "
                                "human welfare support through Wellness Check-In."
                            ),

                        status=
                            "Open",

                        created_at=
                            datetime.now(),

                        updated_at=
                            datetime.now()
                    )

                    db.session.add(
                        support_item
                    )

                    support_request_created = (
                        True
                    )

            _add_audit_event(
                actor_role=
                    "personnel",

                actor_id=
                    personnel_id,

                event_type=
                    "Wellness Check-In Submitted",

                personnel_id=
                    personnel_id,

                resource_type=
                    "WellnessCheckIn",

                resource_id=
                    checkin.id,

                event_summary=
                    (
                        "A voluntary wellness check-in was submitted. "
                        "Detailed questionnaire responses were not copied "
                        "into the audit event."
                    )
            )

            db.session.commit()

        except Exception:

            db.session.rollback()

            return jsonify({
                "success": False,
                "message":
                    "The wellness check-in could not be saved."
            }), 500

        return jsonify({
            "success": True,
            "message":
                "Wellness check-in saved successfully.",
            "checkin":
                _serialize_wellness_checkin(
                    checkin
                ),
            "support_request_created":
                support_request_created,
            "assessment_note":
                (
                    "The check-in is supporting welfare context only. "
                    "It does not independently create a Low, Moderate or Elevated "
                    "attention level and does not trigger employment, leave, "
                    "disciplinary or medical decisions."
                )
        })


    # --------------------------------------------------
    # CAMERA WELLNESS SCAN
    # --------------------------------------------------

    @app.route(
        "/personnel-wellness-scan"
    )
    def personnel_wellness_scan():

        if (
            session.get("role")
            != "personnel"
        ):

            return redirect(
                url_for(
                    "personnel_login"
                )
            )

        personnel_id = session.get(
            "personnel_id"
        )

        personnel = (
            Personnel.query
            .filter_by(
                personnel_id=personnel_id
            )
            .first()
        )

        return render_template(
            "personnel_wellness_scan.html",
            personnel=personnel
        )


    # --------------------------------------------------
    # PERSONNEL SUPPORT
    # --------------------------------------------------

    @app.route(
        "/personnel-support"
    )
    def personnel_support():

        if (
            session.get("role")
            != "personnel"
        ):

            return redirect(
                url_for(
                    "personnel_login"
                )
            )

        personnel_id = session.get(
            "personnel_id"
        )

        personnel = (
            Personnel.query
            .filter_by(
                personnel_id=personnel_id
            )
            .first()
        )

        if personnel is None:
            session.clear()
            return redirect(
                url_for(
                    "personnel_login"
                )
            )

        support_requests = (
            SupportRequest.query
            .filter_by(
                personnel_id=personnel_id
            )
            .order_by(
                SupportRequest.created_at.desc()
            )
            .limit(5)
            .all()
        )

        return render_template(
            "personnel_support.html",
            personnel=personnel,
            support_requests=support_requests
        )


    # --------------------------------------------------
    # SUBMIT PERSONNEL SUPPORT REQUEST
    # --------------------------------------------------

    @app.route(
        "/submit-support-request",
        methods=["POST"]
    )
    def submit_support_request():

        if (
            session.get("role")
            != "personnel"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        data = request.get_json(
            silent=True
        )

        if not data:

            return jsonify({
                "success": False,
                "message":
                    "No support request data received."
            }), 400

        support_category = (
            data.get("support_category")
            or ""
        ).strip()

        contact_method = (
            data.get("contact_method")
            or ""
        ).strip()

        priority = (
            data.get("priority")
            or ""
        ).strip()

        message = (
            data.get("message")
            or ""
        ).strip()

        allowed_categories = [
            "Workload Support",
            "Duty / Rest Concern",
            "Personal Wellness Check",
            "General Welfare Guidance",
            "Other"
        ]

        allowed_contact_methods = [
            "In-Person Meeting",
            "Phone Call",
            "Secure Message"
        ]

        allowed_priorities = [
            "Routine",
            "Soon",
            "Prompt Follow-Up"
        ]

        if support_category not in allowed_categories:

            return jsonify({
                "success": False,
                "message":
                    "Please select a valid support category."
            }), 400

        if contact_method not in allowed_contact_methods:

            return jsonify({
                "success": False,
                "message":
                    "Please select a valid contact method."
            }), 400

        if priority not in allowed_priorities:

            return jsonify({
                "success": False,
                "message":
                    "Please select a valid request priority."
            }), 400

        new_request = SupportRequest(

            personnel_id=session.get(
                "personnel_id"
            ),

            support_category=support_category,

            contact_method=contact_method,

            priority=priority,

            message=message or None,

            status="Open",

            created_at=datetime.now(),

            updated_at=datetime.now()
        )

        try:

            db.session.add(
                new_request
            )

            db.session.commit()

        except Exception:

            db.session.rollback()

            return jsonify({
                "success": False,
                "message":
                    "Could not save the support request."
            }), 500

        return jsonify({

            "success": True,

            "message":
                "Support request submitted successfully.",

            "request": {

                "id":
                    new_request.id,

                "personnel_id":
                    new_request.personnel_id,

                "support_category":
                    new_request.support_category,

                "contact_method":
                    new_request.contact_method,

                "priority":
                    new_request.priority,

                "status":
                    new_request.status,

                "created_at":
                    new_request.created_at.strftime(
                        "%d %b %Y, %I:%M %p"
                    )
            }
        })


    # --------------------------------------------------
    # PERSONNEL NOTIFICATIONS
    # --------------------------------------------------

    @app.route(
        "/personnel-notifications"
    )
    def personnel_notifications():

        if (
            session.get("role")
            != "personnel"
        ):

            return redirect(
                url_for(
                    "personnel_login"
                )
            )

        personnel_id = session.get(
            "personnel_id"
        )

        personnel = (
            Personnel.query
            .filter_by(
                personnel_id=personnel_id
            )
            .first()
        )

        now = datetime.now()

        def relative_time(value):

            if value is None:
                return "Time unavailable"

            delta = now - value

            seconds = max(
                0,
                int(
                    delta.total_seconds()
                )
            )

            if seconds < 60:
                return "Just now"

            minutes = seconds // 60

            if minutes < 60:
                return (
                    str(minutes)
                    + (
                        " minute ago"
                        if minutes == 1
                        else " minutes ago"
                    )
                )

            hours = minutes // 60

            if hours < 24:
                return (
                    str(hours)
                    + (
                        " hour ago"
                        if hours == 1
                        else " hours ago"
                    )
                )

            days = hours // 24

            if days == 1:
                return "Yesterday"

            if days < 7:
                return (
                    str(days)
                    + " days ago"
                )

            return value.strftime(
                "%d %b %Y"
            )

        notification_items = []

        # ----------------------------------------------
        # ASSESSMENT UPDATES FOR THIS PERSON ONLY
        # ----------------------------------------------

        assessments = (
            WelfareAssessment.query
            .filter_by(
                personnel_id=personnel_id
            )
            .order_by(
                WelfareAssessment.assessment_time.desc()
            )
            .limit(25)
            .all()
        )

        for assessment in assessments:

            notification_items.append({
                "event_key":
                    (
                        "assessment:"
                        + str(assessment.id)
                    ),

                "event_time":
                    assessment.assessment_time,

                "time_label":
                    relative_time(
                        assessment.assessment_time
                    ),

                "category":
                    "assessment",

                "title":
                    (
                        "Welfare Assessment Updated"
                    ),

                "message":
                    (
                        "Assessment v"
                        + str(
                            assessment.assessment_version
                        )
                        + " is available. Current welfare-attention level: "
                        + assessment.attention_level
                        + "."
                    ),

                "tag":
                    "Assessment",

                "action_label":
                    "View My Trends",

                "action_url":
                    url_for(
                        "personnel_trends"
                    )
            })

        # ----------------------------------------------
        # SUPPORT REQUEST STATUS UPDATES
        # ----------------------------------------------

        support_requests = (
            SupportRequest.query
            .filter_by(
                personnel_id=personnel_id
            )
            .order_by(
                SupportRequest.updated_at.desc()
            )
            .limit(25)
            .all()
        )

        for support_request in support_requests:

            notification_items.append({
                "event_key":
                    (
                        "support:"
                        + str(
                            support_request.id
                        )
                        + ":"
                        + str(
                            support_request.status
                        )
                    ),

                "event_time":
                    (
                        support_request.updated_at
                        or support_request.created_at
                    ),

                "time_label":
                    relative_time(
                        support_request.updated_at
                        or support_request.created_at
                    ),

                "category":
                    "support",

                "title":
                    "Support Request Update",

                "message":
                    (
                        "Your "
                        + support_request.support_category
                        + " support request is currently "
                        + support_request.status
                        + "."
                    ),

                "tag":
                    "Support",

                "action_label":
                    "View Support",

                "action_url":
                    url_for(
                        "personnel_support"
                    )
            })

        # ----------------------------------------------
        # LEAVE WORKFLOW UPDATES
        # ----------------------------------------------

        leave_requests = (
            LeaveRequest.query
            .filter_by(
                personnel_id=personnel_id
            )
            .order_by(
                LeaveRequest.updated_at.desc()
            )
            .limit(25)
            .all()
        )

        for leave_request in leave_requests:

            notification_items.append({
                "event_key":
                    (
                        "leave:"
                        + str(
                            leave_request.id
                        )
                        + ":"
                        + str(
                            leave_request.status
                        )
                    ),

                "event_time":
                    (
                        leave_request.updated_at
                        or leave_request.created_at
                    ),

                "time_label":
                    relative_time(
                        leave_request.updated_at
                        or leave_request.created_at
                    ),

                "category":
                    "leave",

                "title":
                    "Leave Request Update",

                "message":
                    (
                        "Your "
                        + leave_request.leave_type
                        + " leave request is currently "
                        + leave_request.status
                        + "."
                    ),

                "tag":
                    "Leave",

                "action_label":
                    "View Duty & Leave",

                "action_url":
                    url_for(
                        "personnel_duty_leave"
                    )
            })

        # ----------------------------------------------
        # GENERIC WELFARE FOLLOW-UP UPDATES
        # No officer notes or sensitive internal workflow
        # details are exposed to the personnel view.
        # ----------------------------------------------

        welfare_actions = (
            WelfareAction.query
            .filter(
                WelfareAction.personnel_id
                == personnel_id,

                WelfareAction.status.in_(
                    [
                        "Delivered",
                        "Follow-Up"
                    ]
                )
            )
            .order_by(
                WelfareAction.updated_at.desc()
            )
            .limit(25)
            .all()
        )

        for action in welfare_actions:

            notification_items.append({
                "event_key":
                    (
                        "welfare-action:"
                        + str(
                            action.id
                        )
                        + ":"
                        + action.status
                    ),

                "event_time":
                    (
                        action.updated_at
                        or action.created_at
                    ),

                "time_label":
                    relative_time(
                        action.updated_at
                        or action.created_at
                    ),

                "category":
                    "follow-up",

                "title":
                    (
                        "Welfare Follow-Up Update"
                    ),

                "message":
                    (
                        "A welfare-support follow-up update is available "
                        "for your case. Open Support or My Trends for your "
                        "current personal view."
                    ),

                "tag":
                    "Follow-Up",

                "action_label":
                    "View Support",

                "action_url":
                    url_for(
                        "personnel_support"
                    )
            })

        # Newest first.
        notification_items.sort(
            key=lambda item: (
                item["event_time"]
                or datetime.min
            ),
            reverse=True
        )

        notification_items = (
            notification_items[:75]
        )

        event_keys = [
            item["event_key"]
            for item in notification_items
        ]

        read_keys = set()

        if event_keys:

            read_rows = (
                PersonnelNotificationRead.query
                .filter(
                    PersonnelNotificationRead.personnel_id
                    == personnel_id,

                    PersonnelNotificationRead.event_key.in_(
                        event_keys
                    )
                )
                .all()
            )

            read_keys = {
                row.event_key
                for row in read_rows
            }

        for item in notification_items:

            item["is_read"] = (
                item["event_key"]
                in read_keys
            )

        unread_count = sum(
            1
            for item in notification_items
            if not item["is_read"]
        )

        notification_summary = {
            "unread":
                unread_count,

            "assessment_updates":
                sum(
                    1
                    for item in notification_items
                    if item["category"]
                    == "assessment"
                ),

            "support_updates":
                sum(
                    1
                    for item in notification_items
                    if item["category"]
                    == "support"
                ),

            "leave_updates":
                sum(
                    1
                    for item in notification_items
                    if item["category"]
                    == "leave"
                )
        }

        return render_template(
            "personnel_notifications.html",
            personnel=personnel,
            notification_items=
                notification_items,
            notification_summary=
                notification_summary
        )


    @app.route(
        "/api/personnel/notifications/read",
        methods=["POST"]
    )
    def mark_personnel_notification_read():

        if (
            session.get("role")
            != "personnel"
        ):

            return jsonify({
                "success": False,
                "message":
                    "Unauthorized."
            }), 401

        data = (
            request.get_json(
                silent=True
            )
            or {}
        )

        event_key = (
            data.get(
                "event_key"
            )
            or ""
        ).strip()

        if not event_key:

            return jsonify({
                "success": False,
                "message":
                    "Notification event key is required."
            }), 400

        personnel_id = session.get(
            "personnel_id"
        )

        existing = (
            PersonnelNotificationRead.query
            .filter_by(
                personnel_id=
                    personnel_id,
                event_key=
                    event_key
            )
            .first()
        )

        if existing is None:

            read_row = (
                PersonnelNotificationRead(
                    personnel_id=
                        personnel_id,
                    event_key=
                        event_key,
                    read_at=
                        datetime.now()
                )
            )

            db.session.add(
                read_row
            )

            db.session.commit()

        return jsonify({
            "success": True,
            "event_key":
                event_key
        })


    # --------------------------------------------------
    # PERSONNEL PRIVACY
    # --------------------------------------------------

    @app.route(
        "/personnel-privacy"
    )
    def personnel_privacy():

        if (
            session.get("role")
            != "personnel"
        ):

            return redirect(
                url_for(
                    "personnel_login"
                )
            )

        personnel_id = session.get(
            "personnel_id"
        )

        personnel = (
            Personnel.query
            .filter_by(
                personnel_id=personnel_id
            )
            .first()
        )

        data_inventory = {
            "operational_records":
                (
                    OperationalRecord.query
                    .filter_by(
                        personnel_id=personnel_id,
                        is_valid=True
                    )
                    .count()
                ),

            "wellness_checkins":
                (
                    WellnessCheckIn.query
                    .filter_by(
                        personnel_id=personnel_id
                    )
                    .count()
                ),

            "wellness_scans":
                (
                    WellnessScan.query
                    .filter_by(
                        personnel_id=personnel_id
                    )
                    .count()
                ),

            "assessments":
                (
                    WelfareAssessment.query
                    .filter_by(
                        personnel_id=personnel_id
                    )
                    .count()
                ),

            "support_requests":
                (
                    SupportRequest.query
                    .filter_by(
                        personnel_id=personnel_id
                    )
                    .count()
                ),

            "leave_requests":
                (
                    LeaveRequest.query
                    .filter_by(
                        personnel_id=personnel_id
                    )
                    .count()
                )
        }

        current_assessment = (
            WelfareAssessment.query
            .filter_by(
                personnel_id=personnel_id,
                is_current=True
            )
            .order_by(
                WelfareAssessment.assessment_time.desc()
            )
            .first()
        )

        latest_checkin = (
            WellnessCheckIn.query
            .filter_by(
                personnel_id=personnel_id
            )
            .order_by(
                WellnessCheckIn.submitted_at.desc()
            )
            .first()
        )

        latest_scan = (
            WellnessScan.query
            .filter_by(
                personnel_id=personnel_id
            )
            .order_by(
                WellnessScan.saved_at.desc()
            )
            .first()
        )

        voluntary_sources = {
            "checkin": {
                "status":
                    (
                        "Used Voluntarily"
                        if latest_checkin
                        else "Not Used"
                    ),
                "last_used":
                    (
                        latest_checkin.submitted_at
                        if latest_checkin
                        else None
                    ),
                "record_count":
                    data_inventory[
                        "wellness_checkins"
                    ]
            },

            "scan": {
                "status":
                    (
                        "Used Voluntarily"
                        if latest_scan
                        else "Not Used"
                    ),
                "last_used":
                    (
                        latest_scan.saved_at
                        if latest_scan
                        else None
                    ),
                "record_count":
                    data_inventory[
                        "wellness_scans"
                    ]
            }
        }

        safe_event_types = [
            "Personnel Created",
            "Personnel Updated",
            "Operational Record Created",
            "Operational Record Updated",
            "Leave Request Submitted",
            "Leave Request Cancelled",
            "Leave Request Approved",
            "Leave Request Denied",
            "Wellness Check-In Submitted",
            "Welfare Assessment Recalculated",
            "Decision Replay Accessed"
        ]

        audit_events = (
            AuditEvent.query
            .filter(
                AuditEvent.personnel_id
                == personnel_id,

                AuditEvent.event_type.in_(
                    safe_event_types
                )
            )
            .order_by(
                AuditEvent.created_at.desc()
            )
            .limit(30)
            .all()
        )

        privacy_summary = {
            "personnel_id":
                personnel_id,

            "current_attention":
                (
                    current_assessment.attention_level
                    if current_assessment
                    else "Insufficient Data"
                ),

            "stored_source_count":
                sum(
                    1
                    for value in data_inventory.values()
                    if value > 0
                ),

            "audit_event_count":
                len(audit_events),

            "voluntary_source_count":
                (
                    int(
                        latest_checkin
                        is not None
                    )
                    +
                    int(
                        latest_scan
                        is not None
                    )
                )
        }

        return render_template(
            "personnel_privacy.html",
            personnel=personnel,
            data_inventory=data_inventory,
            current_assessment=
                current_assessment,
            voluntary_sources=
                voluntary_sources,
            audit_events=
                audit_events,
            privacy_summary=
                privacy_summary
        )


    # --------------------------------------------------
    # SAVE WELLNESS SCAN
    # --------------------------------------------------

    @app.route(
        "/save-wellness-scan",
        methods=["POST"]
    )
    def save_wellness_scan():

        if (
            session.get("role")
            != "personnel"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        data = request.get_json(
            silent=True
        )

        if not data:

            return jsonify({
                "success": False,
                "message":
                    "No scan data received."
            }), 400

        heart_rate = data.get(
            "heart_rate"
        )

        breathing_rate = data.get(
            "breathing_rate"
        )

        signal_quality = data.get(
            "signal_quality"
        )

        confidence = data.get(
            "confidence"
        )

        measurement_time = data.get(
            "measurement_time"
        )

        scan_duration = data.get(
            "scan_duration"
        )

        if confidence not in [
            "High",
            "Medium"
        ]:

            return jsonify({
                "success": False,
                "message":
                    "Only usable scans can be saved."
            }), 400

        try:

            if heart_rate not in [
                None,
                "",
                "—"
            ]:

                heart_rate = int(
                    heart_rate
                )

            else:

                heart_rate = None


            if breathing_rate not in [
                None,
                "",
                "—"
            ]:

                breathing_rate = int(
                    breathing_rate
                )

            else:

                breathing_rate = None


            if (
                heart_rate is None
                and breathing_rate is None
            ):

                return jsonify({
                    "success": False,
                    "message":
                        "No usable heart-rate or breathing-rate measurement was provided. Nothing was saved."
                }), 400


        except (
            TypeError,
            ValueError
        ):

            return jsonify({
                "success": False,
                "message":
                    "Invalid scan measurement values."
            }), 400


        saved_time = datetime.now()


        new_scan = WellnessScan(

            personnel_id=session.get(
                "personnel_id"
            ),

            heart_rate=heart_rate,

            breathing_rate=breathing_rate,

            signal_quality=signal_quality,

            confidence=confidence,

            measurement_time=measurement_time,

            scan_duration=scan_duration,

            saved_at=saved_time
        )


        try:

            db.session.add(
                new_scan
            )

            db.session.commit()

        except Exception:

            db.session.rollback()

            return jsonify({
                "success": False,
                "message":
                    "Could not save the wellness scan."
            }), 500


        scan_record = {

            "id":
                new_scan.id,

            "personnel_id":
                new_scan.personnel_id,

            "heart_rate":
                new_scan.heart_rate,

            "breathing_rate":
                new_scan.breathing_rate,

            "signal_quality":
                new_scan.signal_quality,

            "confidence":
                new_scan.confidence,

            "measurement_time":
                new_scan.measurement_time,

            "scan_duration":
                new_scan.scan_duration,

            "saved_at":
                saved_time.strftime(
                    "%d %b %Y, %I:%M %p"
                )
        }


        return jsonify({

            "success": True,

            "message":
                "Wellness scan saved successfully.",

            "record":
                scan_record
        })


    # --------------------------------------------------
    # SEED 72-HOUR PROTOTYPE RECOVERY ROSTER
    # --------------------------------------------------

    @app.route(
        "/api/personnel/seed-recovery-roster",
        methods=["POST"]
    )
    def seed_recovery_roster():

        if (
            session.get("role")
            not in ["personnel", "welfare_officer"]
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        if (
            session.get("role")
            == "personnel"
        ):

            personnel_id = (
                session.get(
                    "personnel_id"
                )
                or ""
            ).strip()

        else:

            personnel_id = (
                request.args.get(
                    "personnel_id"
                )
                or ""
            ).strip()

        if not personnel_id:

            return jsonify({
                "success": False,
                "message":
                    "A Personnel ID is required for the prototype roster."
            }), 400

        personnel = (
            Personnel.query
            .filter_by(
                personnel_id=personnel_id
            )
            .first()
        )

        if personnel is None:

            return jsonify({
                "success": False,
                "message":
                    "Personnel record was not found."
            }), 404

        today = datetime.now().date()

        existing_slots = (
            RecoveryRosterSlot.query
            .filter_by(
                personnel_id=personnel_id,
                source_type="Prototype Demonstration Data"
            )
            .all()
        )

        for slot in existing_slots:
            db.session.delete(slot)

        prototype_slots = [

            RecoveryRosterSlot(
                personnel_id=personnel_id,
                start_time=datetime.combine(
                    today,
                    datetime.strptime(
                        "08:00",
                        "%H:%M"
                    ).time()
                ),
                end_time=datetime.combine(
                    today,
                    datetime.strptime(
                        "16:00",
                        "%H:%M"
                    ).time()
                ),
                commitment_type="Operational Duty",
                mandatory=True,
                reassignable=False,
                constraint_type="Minimum Staffing",
                constraint_note=(
                    "Current duty block is marked mandatory "
                    "for prototype staffing coverage."
                ),
                source_type="Prototype Demonstration Data"
            ),

            RecoveryRosterSlot(
                personnel_id=personnel_id,
                start_time=datetime.combine(
                    today + timedelta(days=1),
                    datetime.strptime(
                        "09:00",
                        "%H:%M"
                    ).time()
                ),
                end_time=datetime.combine(
                    today + timedelta(days=1),
                    datetime.strptime(
                        "13:00",
                        "%H:%M"
                    ).time()
                ),
                commitment_type="Mandatory Training",
                mandatory=True,
                reassignable=False,
                constraint_type="Mandatory Commitment",
                constraint_note=(
                    "Training is marked mandatory in the "
                    "prototype roster."
                ),
                source_type="Prototype Demonstration Data"
            ),

            RecoveryRosterSlot(
                personnel_id=personnel_id,
                start_time=datetime.combine(
                    today + timedelta(days=2),
                    datetime.strptime(
                        "10:00",
                        "%H:%M"
                    ).time()
                ),
                end_time=datetime.combine(
                    today + timedelta(days=2),
                    datetime.strptime(
                        "14:00",
                        "%H:%M"
                    ).time()
                ),
                commitment_type="Administrative Duty",
                mandatory=False,
                reassignable=True,
                constraint_type="Reassignable",
                constraint_note=(
                    "This non-critical block is marked "
                    "reassignable for planner demonstration."
                ),
                source_type="Prototype Demonstration Data"
            ),

            RecoveryRosterSlot(
                personnel_id=personnel_id,
                start_time=datetime.combine(
                    today + timedelta(days=3),
                    datetime.strptime(
                        "08:00",
                        "%H:%M"
                    ).time()
                ),
                end_time=datetime.combine(
                    today + timedelta(days=3),
                    datetime.strptime(
                        "12:00",
                        "%H:%M"
                    ).time()
                ),
                commitment_type="Operational Duty",
                mandatory=True,
                reassignable=False,
                constraint_type="Minimum Staffing",
                constraint_note=(
                    "Prototype future duty block retained "
                    "to bound the recovery window."
                ),
                source_type="Prototype Demonstration Data"
            )
        ]

        try:

            db.session.add_all(
                prototype_slots
            )

            db.session.commit()

        except Exception:

            db.session.rollback()

            return jsonify({
                "success": False,
                "message":
                    "Could not seed prototype recovery roster."
            }), 500

        return jsonify({

            "success": True,

            "message":
                "Prototype 72-hour recovery roster seeded successfully.",

            "personnel_id":
                personnel_id,

            "source_type":
                "Prototype Demonstration Data",

            "slot_count":
                len(prototype_slots)
        })


    # --------------------------------------------------
    # 72-HOUR RECOVERY WINDOW ANALYSIS
    # --------------------------------------------------

    @app.route(
        "/api/welfare-officer/recovery-window-analysis"
    )
    def recovery_window_analysis():

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        personnel_id = (
            request.args.get("personnel_id")
            or ""
        ).strip()

        if not personnel_id:

            return jsonify({
                "success": False,
                "message":
                    "Select a personnel case before running recovery analysis."
            }), 400

        window_start = datetime.now()
        window_end = (
            window_start
            + timedelta(hours=72)
        )

        roster_slots = (
            RecoveryRosterSlot.query
            .filter(
                RecoveryRosterSlot.personnel_id
                == personnel_id,
                RecoveryRosterSlot.is_valid
                == True,
                RecoveryRosterSlot.end_time
                > window_start,
                RecoveryRosterSlot.start_time
                < window_end
            )
            .order_by(
                RecoveryRosterSlot.start_time.asc()
            )
            .all()
        )

        if not roster_slots:

            return jsonify({
                "success": True,
                "personnel_id": personnel_id,
                "data_status":
                    "Insufficient Scheduling Data",
                "message":
                    "No valid roster commitments are available "
                    "inside the next 72 hours.",
                "largest_recovery_window_hours": None,
                "recovery_windows": []
            })

        recovery_windows = []

        cursor = window_start

        for slot in roster_slots:

            effective_start = max(
                slot.start_time,
                window_start
            )

            effective_end = min(
                slot.end_time,
                window_end
            )

            if effective_start > cursor:

                gap_hours = round(
                    (
                        effective_start
                        - cursor
                    ).total_seconds()
                    / 3600,
                    1
                )

                recovery_windows.append({
                    "start":
                        cursor.strftime(
                            "%d %b %Y, %I:%M %p"
                        ),
                    "end":
                        effective_start.strftime(
                            "%d %b %Y, %I:%M %p"
                        ),
                    "hours":
                        gap_hours
                })

            if effective_end > cursor:

                cursor = effective_end

        if cursor < window_end:

            gap_hours = round(
                (
                    window_end
                    - cursor
                ).total_seconds()
                / 3600,
                1
            )

            recovery_windows.append({
                "start":
                    cursor.strftime(
                        "%d %b %Y, %I:%M %p"
                    ),
                "end":
                    window_end.strftime(
                        "%d %b %Y, %I:%M %p"
                    ),
                "hours":
                    gap_hours
            })

        largest_window = (
            max(
                recovery_windows,
                key=lambda item: item["hours"]
            )
            if recovery_windows
            else None
        )

        roster_data = []

        for slot in roster_slots:

            roster_data.append({
                "id":
                    slot.id,

                "start_time":
                    slot.start_time.strftime(
                        "%d %b %Y, %I:%M %p"
                    ),

                "end_time":
                    slot.end_time.strftime(
                        "%d %b %Y, %I:%M %p"
                    ),

                "commitment_type":
                    slot.commitment_type,

                "mandatory":
                    slot.mandatory,

                "reassignable":
                    slot.reassignable,

                "constraint_type":
                    slot.constraint_type,

                "constraint_note":
                    slot.constraint_note,

                "source_type":
                    slot.source_type
            })

        baseline_largest_hours = (
            largest_window["hours"]
            if largest_window
            else 0
        )

        planner_options = []

        planner_options.append({
            "option_id":
                "A",
            "title":
                "Protect Existing Recovery Window",
            "status":
                "Feasible",
            "requires_human_approval":
                True,
            "schedule_change_required":
                False,
            "projected_largest_recovery_hours":
                baseline_largest_hours,
            "recovery_opportunity_change_hours":
                0,
            "reason":
                (
                    "Protect the largest already-available "
                    "uninterrupted recovery window. No roster "
                    "change is proposed."
                )
        })

        best_reassignable_option = None

        for candidate in roster_slots:

            if (
                candidate.mandatory
                or not candidate.reassignable
            ):
                continue

            simulated_slots = [
                slot
                for slot in roster_slots
                if slot.id != candidate.id
            ]

            simulated_windows = []
            simulated_cursor = window_start

            for slot in simulated_slots:

                effective_start = max(
                    slot.start_time,
                    window_start
                )

                effective_end = min(
                    slot.end_time,
                    window_end
                )

                if effective_start > simulated_cursor:

                    simulated_windows.append(
                        round(
                            (
                                effective_start
                                - simulated_cursor
                            ).total_seconds()
                            / 3600,
                            1
                        )
                    )

                if effective_end > simulated_cursor:
                    simulated_cursor = effective_end

            if simulated_cursor < window_end:

                simulated_windows.append(
                    round(
                        (
                            window_end
                            - simulated_cursor
                        ).total_seconds()
                        / 3600,
                        1
                    )
                )

            simulated_largest = (
                max(simulated_windows)
                if simulated_windows
                else 0
            )

            change_hours = round(
                simulated_largest
                - baseline_largest_hours,
                1
            )

            candidate_option = {
                "option_id":
                    "B",
                "title":
                    "Reassign Non-Critical Commitment",
                "status":
                    "Requires Human Approval",
                "requires_human_approval":
                    True,
                "schedule_change_required":
                    True,
                "candidate_slot_id":
                    candidate.id,
                "candidate_commitment":
                    candidate.commitment_type,
                "projected_largest_recovery_hours":
                    simulated_largest,
                "recovery_opportunity_change_hours":
                    max(change_hours, 0),
                "reason":
                    (
                        candidate.constraint_note
                        or
                        "This slot is marked reassignable "
                        "but requires authorized roster approval."
                    )
            }

            if (
                best_reassignable_option is None
                or
                candidate_option[
                    "projected_largest_recovery_hours"
                ]
                >
                best_reassignable_option[
                    "projected_largest_recovery_hours"
                ]
            ):
                best_reassignable_option = candidate_option

        if best_reassignable_option:

            planner_options.append(
                best_reassignable_option
            )

        else:

            planner_options.append({
                "option_id":
                    "B",
                "title":
                    "Reassign Non-Critical Commitment",
                "status":
                    "Rejected",
                "requires_human_approval":
                    True,
                "schedule_change_required":
                    True,
                "projected_largest_recovery_hours":
                    baseline_largest_hours,
                "recovery_opportunity_change_hours":
                    0,
                "reason":
                    "No valid non-mandatory reassignable slot "
                    "is available inside the 72-hour window."
            })

        blocked_slot = next(
            (
                slot
                for slot in roster_slots
                if (
                    slot.mandatory
                    and not slot.reassignable
                )
            ),
            None
        )

        if blocked_slot:

            planner_options.append({
                "option_id":
                    "C",
                "title":
                    "Move Mandatory Commitment",
                "status":
                    "Rejected",
                "requires_human_approval":
                    True,
                "schedule_change_required":
                    True,
                "candidate_slot_id":
                    blocked_slot.id,
                "candidate_commitment":
                    blocked_slot.commitment_type,
                "projected_largest_recovery_hours":
                    baseline_largest_hours,
                "recovery_opportunity_change_hours":
                    0,
                "reason":
                    (
                        blocked_slot.constraint_note
                        or
                        "The selected commitment is mandatory "
                        "and not marked reassignable."
                    )
            })

        else:

            planner_options.append({
                "option_id":
                    "C",
                "title":
                    "Move Mandatory Commitment",
                "status":
                    "Rejected",
                "requires_human_approval":
                    True,
                "schedule_change_required":
                    True,
                "projected_largest_recovery_hours":
                    baseline_largest_hours,
                "recovery_opportunity_change_hours":
                    0,
                "reason":
                    "No eligible mandatory commitment was found "
                    "for this alternative."
            })

        return jsonify({
            "success": True,
            "personnel_id":
                personnel_id,
            "data_status":
                "Available",
            "window_start":
                window_start.strftime(
                    "%d %b %Y, %I:%M %p"
                ),
            "window_end":
                window_end.strftime(
                    "%d %b %Y, %I:%M %p"
                ),
            "largest_recovery_window_hours":
                baseline_largest_hours,
            "largest_recovery_window":
                largest_window,
            "recovery_windows":
                recovery_windows,
            "roster_slots":
                roster_data,
            "planner_options":
                planner_options
        })


    # --------------------------------------------------
    # CREATE WELFARE ACTION
    # --------------------------------------------------

    @app.route(
        "/api/welfare-officer/welfare-actions",
        methods=["POST"]
    )
    def create_welfare_action():

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        data = request.get_json(
            silent=True
        ) or {}

        personnel_id = (
            data.get("personnel_id")
            or ""
        ).strip()

        if not personnel_id:

            return jsonify({
                "success": False,
                "message":
                    "Select a personnel case before recording a welfare action."
            }), 400

        action_type = (
            data.get("action_type")
            or ""
        ).strip()

        planner_option_id = (
            data.get("planner_option_id")
            or ""
        ).strip() or None

        officer_note = (
            data.get("officer_note")
            or ""
        ).strip() or None

        allowed_action_types = [
            "Welfare Conversation",
            "Duty / Rest Review",
            "Recovery Window",
            "Training Review",
            "Counselling Referral",
            "Support Follow-Up",
            "Continue Monitoring"
        ]

        if (
            action_type
            not in allowed_action_types
        ):

            return jsonify({
                "success": False,
                "message":
                    "Please select a valid welfare action."
            }), 400

        if (
            planner_option_id
            and planner_option_id
            not in ["A", "B", "C"]
        ):

            return jsonify({
                "success": False,
                "message":
                    "Planner option must be A, B or C."
            }), 400

        current_assessment = (
            WelfareAssessment.query
            .filter_by(
                personnel_id=personnel_id,
                is_current=True
            )
            .order_by(
                WelfareAssessment.assessment_time.desc()
            )
            .first()
        )

        welfare_action = WelfareAction(
            personnel_id=personnel_id,
            action_type=action_type,
            source_assessment_id=(
                current_assessment.id
                if current_assessment
                else None
            ),
            source_assessment_version=(
                current_assessment.assessment_version
                if current_assessment
                else None
            ),
            planner_option_id=planner_option_id,
            status="Planned",
            officer_note=officer_note
        )

        try:

            db.session.add(
                welfare_action
            )

            db.session.commit()

        except Exception:

            db.session.rollback()

            return jsonify({
                "success": False,
                "message":
                    "Could not create the welfare action."
            }), 500

        return jsonify({
            "success": True,
            "message":
                "Welfare action recorded as Planned.",
            "action": {
                "id":
                    welfare_action.id,
                "personnel_id":
                    welfare_action.personnel_id,
                "action_type":
                    welfare_action.action_type,
                "planner_option_id":
                    welfare_action.planner_option_id,
                "status":
                    welfare_action.status,
                "source_assessment_version":
                    welfare_action.source_assessment_version,
                "created_at":
                    welfare_action.created_at.strftime(
                        "%d %b %Y, %I:%M %p"
                    )
            }
        })


    # --------------------------------------------------
    # UPDATE WELFARE ACTION STATUS
    # --------------------------------------------------

    @app.route(
        "/api/welfare-officer/welfare-actions/<int:action_id>/status",
        methods=["POST"]
    )
    def update_welfare_action_status(action_id):

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        data = request.get_json(
            silent=True
        ) or {}

        new_status = (
            data.get("status")
            or ""
        ).strip()

        verification_note = (
            data.get("verification_note")
            or ""
        ).strip() or None

        allowed_statuses = [
            "Planned",
            "Approved",
            "Delivered",
            "Interrupted",
            "Unable to Verify",
            "Follow-Up"
        ]

        if (
            new_status
            not in allowed_statuses
        ):

            return jsonify({
                "success": False,
                "message":
                    "Please select a valid welfare action status."
            }), 400

        welfare_action = (
            WelfareAction.query
            .filter_by(
                id=action_id
            )
            .first()
        )

        if not welfare_action:

            return jsonify({
                "success": False,
                "message":
                    "Welfare action not found."
            }), 404

        now = datetime.now()

        welfare_action.status = new_status
        welfare_action.updated_at = now

        if (
            new_status == "Approved"
            and welfare_action.approved_at is None
        ):

            welfare_action.approved_at = now

        if (
            new_status == "Delivered"
            and welfare_action.delivered_at is None
        ):

            welfare_action.delivered_at = now

        if (
            new_status == "Interrupted"
            and welfare_action.interrupted_at is None
        ):

            welfare_action.interrupted_at = now

        if verification_note:

            welfare_action.verification_note = (
                verification_note
            )

        audit_event = AuditEvent(
            actor_role=
                session.get("role")
                or "welfare_officer",

            actor_id=
                session.get("officer_id")
                or session.get("user_id")
                or "authenticated_officer",

            event_type=
                "Welfare Action Status Updated",

            personnel_id=
                welfare_action.personnel_id,

            resource_type=
                "WelfareAction",

            resource_id=
                str(welfare_action.id),

            event_summary=
                (
                    "Welfare action "
                    + str(welfare_action.id)
                    + " updated to "
                    + welfare_action.status
                    + "."
                ),

            source_type=
                "VeerPulse Application",

            created_at=
                datetime.now()
        )

        db.session.add(
            audit_event
        )

        try:

            db.session.commit()

        except Exception:

            db.session.rollback()

            return jsonify({
                "success": False,
                "message":
                    "Could not update the welfare action."
            }), 500

        return jsonify({
            "success": True,
            "message":
                "Welfare action status updated successfully.",
            "action": {
                "id":
                    welfare_action.id,
                "personnel_id":
                    welfare_action.personnel_id,
                "action_type":
                    welfare_action.action_type,
                "planner_option_id":
                    welfare_action.planner_option_id,
                "status":
                    welfare_action.status,
                "source_assessment_version":
                    welfare_action.source_assessment_version,
                "approved_at":
                    (
                        welfare_action.approved_at.strftime(
                            "%d %b %Y, %I:%M %p"
                        )
                        if welfare_action.approved_at
                        else None
                    ),
                "delivered_at":
                    (
                        welfare_action.delivered_at.strftime(
                            "%d %b %Y, %I:%M %p"
                        )
                        if welfare_action.delivered_at
                        else None
                    ),
                "interrupted_at":
                    (
                        welfare_action.interrupted_at.strftime(
                            "%d %b %Y, %I:%M %p"
                        )
                        if welfare_action.interrupted_at
                        else None
                    ),
                "verification_note":
                    welfare_action.verification_note
            }
        })


    # --------------------------------------------------
    # LIST WELFARE ACTIONS
    # --------------------------------------------------

    @app.route(
        "/api/welfare-officer/welfare-actions",
        methods=["GET"]
    )
    def list_welfare_actions():

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        personnel_id = (
            request.args.get("personnel_id")
            or ""
        ).strip()

        if not personnel_id:

            return jsonify({
                "success": False,
                "message":
                    "Select a personnel case before viewing welfare actions."
            }), 400

        actions = (
            WelfareAction.query
            .filter_by(
                personnel_id=personnel_id
            )
            .order_by(
                WelfareAction.created_at.desc()
            )
            .all()
        )

        action_items = []

        for item in actions:

            action_items.append({
                "id":
                    item.id,

                "personnel_id":
                    item.personnel_id,

                "action_type":
                    item.action_type,

                "planner_option_id":
                    item.planner_option_id,

                "status":
                    item.status,

                "source_assessment_version":
                    item.source_assessment_version,

                "officer_note":
                    item.officer_note,

                "approved_at":
                    (
                        item.approved_at.strftime(
                            "%d %b %Y, %I:%M %p"
                        )
                        if item.approved_at
                        else None
                    ),

                "delivered_at":
                    (
                        item.delivered_at.strftime(
                            "%d %b %Y, %I:%M %p"
                        )
                        if item.delivered_at
                        else None
                    ),

                "interrupted_at":
                    (
                        item.interrupted_at.strftime(
                            "%d %b %Y, %I:%M %p"
                        )
                        if item.interrupted_at
                        else None
                    ),

                "verification_note":
                    item.verification_note,

                "created_at":
                    item.created_at.strftime(
                        "%d %b %Y, %I:%M %p"
                    ),

                "updated_at":
                    item.updated_at.strftime(
                        "%d %b %Y, %I:%M %p"
                    )
            })

        return jsonify({
            "success": True,
            "personnel_id":
                personnel_id,
            "action_count":
                len(action_items),
            "actions":
                action_items
        })


    # --------------------------------------------------
    # WELFARE OFFICER FOLLOW-UP ASSESSMENT
    # --------------------------------------------------

    @app.route(
        "/api/welfare-officer/personnel/<string:personnel_id>/follow-up-assessment",
        methods=["POST"]
    )
    def welfare_officer_follow_up_assessment(
        personnel_id
    ):

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        personnel_id = (
            personnel_id
            or ""
        ).strip()

        if not personnel_id:

            return jsonify({
                "success": False,
                "message":
                    "Personnel ID is required."
            }), 400

        latest_welfare_action = (
            WelfareAction.query
            .filter_by(
                personnel_id=personnel_id
            )
            .order_by(
                WelfareAction.created_at.desc()
            )
            .first()
        )

        if not latest_welfare_action:

            return jsonify({
                "success": False,
                "message":
                    "No welfare action is available for follow-up."
            }), 400

        if (
            latest_welfare_action.status
            != "Follow-Up"
        ):

            return jsonify({
                "success": False,
                "message":
                    "Move the latest welfare action to Follow-Up before recalculating the assessment."
            }), 400

        today = datetime.now().date()

        current_start = (
            today -
            timedelta(days=6)
        )

        current_records = (
            OperationalRecord.query
            .filter(
                OperationalRecord.personnel_id
                == personnel_id,

                OperationalRecord.is_valid
                == True,

                OperationalRecord.record_date
                >= current_start,

                OperationalRecord.record_date
                <= today
            )
            .order_by(
                OperationalRecord.record_date.asc()
            )
            .all()
        )

        historical_records = []

        # ----------------------------------------------
        # DECISION REPLAY SNAPSHOT FIELDS
        # ----------------------------------------------

        snapshot_current_record_count = (
            len(current_records)
        )

        snapshot_baseline_record_count = None
        snapshot_duty_delta = None
        snapshot_rest_delta = None
        snapshot_recovery_day_delta = None
        snapshot_factor_count = None
        snapshot_severity_score = None

        # ----------------------------------------------
        # EVIDENCE GATE 1:
        # enough recent validated operational records
        # ----------------------------------------------

        if len(current_records) < 4:

            attention_level = "Insufficient Data"

            evidence_status = (
                "Awaiting Sufficient Evidence"
            )

            confidence = None

            explanation_summary = (
                "Assessment withheld because fewer than four "
                "validated current-period operational records "
                "are available."
            )

            baseline_context = None

            baseline_version = None

        else:

            current_context = (
                current_records[-1].context_type
                or "Unknown"
            )

            comparison_record_limit = (
                len(current_records)
            )

            # Personal operational baseline:
            # use the immediately preceding equal-length block of
            # validated records. Recovery Day records are deliberately
            # retained so recovery-day availability can be compared.
            historical_records = (
                OperationalRecord.query
                .filter(
                    OperationalRecord.personnel_id
                    == personnel_id,

                    OperationalRecord.is_valid
                    == True,

                    OperationalRecord.record_date
                    < current_start
                )
                .order_by(
                    OperationalRecord.record_date.desc()
                )
                .limit(
                    comparison_record_limit
                )
                .all()
            )

            historical_records.reverse()

            snapshot_baseline_record_count = (
                len(historical_records)
            )

            current_worked_records = [
                record
                for record in current_records
                if (
                    record.duty_hours is not None
                    and record.duty_hours > 0
                )
            ]

            baseline_worked_records = [
                record
                for record in historical_records
                if (
                    record.duty_hours is not None
                    and record.duty_hours > 0
                )
            ]

            # ------------------------------------------
            # EVIDENCE GATE 2:
            # equal-length comparable context + enough
            # worked-day observations
            # ------------------------------------------

            evidence_is_sufficient = (
                len(historical_records)
                == comparison_record_limit

                and len(historical_records) >= 4

                and len(current_worked_records) >= 3

                and len(baseline_worked_records) >= 3
            )

            if not evidence_is_sufficient:

                attention_level = (
                    "Insufficient Data"
                )

                evidence_status = (
                    "Awaiting Comparable Evidence"
                )

                confidence = None

                baseline_context = (
                    "Personal Operational Baseline"
                )

                baseline_version = None

                explanation_summary = (
                    "Assessment withheld because there are not "
                    "enough equal-length validated historical "
                    "records for the personal operational baseline."
                )

            else:

                def average(values):

                    clean_values = [
                        value
                        for value in values
                        if value is not None
                    ]

                    if not clean_values:
                        return None

                    return round(
                        sum(clean_values)
                        / len(clean_values),
                        1
                    )

                current_duty_avg = average([
                    record.duty_hours
                    for record
                    in current_worked_records
                ])

                baseline_duty_avg = average([
                    record.duty_hours
                    for record
                    in baseline_worked_records
                ])

                current_rest_avg = average([
                    record.rest_hours
                    for record
                    in current_records
                ])

                baseline_rest_avg = average([
                    record.rest_hours
                    for record
                    in historical_records
                ])

                current_recovery_days = sum(
                    1
                    for record in current_records
                    if (
                        record.duty_hours is not None
                        and record.duty_hours == 0
                    )
                )

                baseline_recovery_days = sum(
                    1
                    for record in historical_records
                    if (
                        record.duty_hours is not None
                        and record.duty_hours == 0
                    )
                )

                _t_vals = [r.training_hours for r in current_records if r.training_hours is not None]
                current_training_hours = round(sum(_t_vals), 1) if _t_vals else None

                leave_cancellations = sum(
                    1
                    for record in current_records
                    if (
                        record.leave_status
                        == "Leave Cancelled"
                    )
                )

                duty_delta = round(
                    current_duty_avg
                    - baseline_duty_avg,
                    1
                )

                rest_delta = round(
                    current_rest_avg
                    - baseline_rest_avg,
                    1
                )

                recovery_delta = (
                    current_recovery_days
                    - baseline_recovery_days
                )

                snapshot_duty_delta = duty_delta
                snapshot_rest_delta = rest_delta
                snapshot_recovery_day_delta = (
                    recovery_delta
                )

                # --------------------------------------
                # TRANSPARENT PROTOTYPE RULES
                #
                # Important:
                # - no single factor can elevate attention
                # - transfer/deployment/training/leave are
                #   context and never standalone triggers
                # --------------------------------------

                concerning_factors = []

                severity_score = 0

                if duty_delta >= 2.0:

                    concerning_factors.append(
                        "Worked-day duty intensity is substantially above the comparable personal baseline."
                    )

                    severity_score += 2

                elif duty_delta >= 1.0:

                    concerning_factors.append(
                        "Worked-day duty intensity is moderately above the comparable personal baseline."
                    )

                    severity_score += 1

                if rest_delta <= -2.0:

                    concerning_factors.append(
                        "Average rest is substantially below the comparable personal baseline."
                    )

                    severity_score += 2

                elif rest_delta <= -1.0:

                    concerning_factors.append(
                        "Average rest is moderately below the comparable personal baseline."
                    )

                    severity_score += 1

                if recovery_delta <= -2:

                    concerning_factors.append(
                        "Recovery-day availability is substantially below the comparable personal baseline."
                    )

                    severity_score += 2

                elif recovery_delta <= -1:

                    concerning_factors.append(
                        "Recovery-day availability is below the comparable personal baseline."
                    )

                    severity_score += 1

                # --------------------------------------
                # MULTI-FACTOR GUARDRAIL
                # --------------------------------------

                factor_count = len(
                    concerning_factors
                )

                snapshot_factor_count = (
                    factor_count
                )

                snapshot_severity_score = (
                    severity_score
                )

                if (
                    factor_count >= 3
                    and severity_score >= 4
                ):

                    attention_level = (
                        "Elevated"
                    )

                elif factor_count >= 2:

                    attention_level = (
                        "Moderate"
                    )

                else:

                    attention_level = (
                        "Low"
                    )

                evidence_status = (
                    "Sufficient Operational Evidence"
                )

                confidence = (
                    "Medium"
                )

                baseline_context = (
                    "Personal Operational Baseline"
                )

                baseline_version = (
                    "VP-BASELINE-1.1"
                )

                # --------------------------------------
                # EXPLAIN WHY
                # --------------------------------------

                if factor_count == 0:

                    explanation_summary = (
                        "Current validated operational evidence does "
                        "not show a sustained multi-factor deterioration "
                        "relative to the person's comparable "
                        "personal operational baseline. "
                        f"Worked-day duty changed by {duty_delta:+.1f} hours, "
                        f"average rest changed by {rest_delta:+.1f} hours, "
                        f"and recovery-day availability changed by "
                        f"{recovery_delta:+d} day(s). "
                        "Training and leave records are retained as "
                        "supporting operational context and do not "
                        "independently increase welfare attention."
                    )

                elif factor_count == 1:

                    explanation_summary = (
                        "One operational deviation is present, but "
                        "VeerPulse does not elevate welfare attention "
                        "from a single factor alone. "
                        + concerning_factors[0]
                    )

                else:

                    explanation_summary = (
                        "Multiple sustained operational deviations are "
                        "present relative to the person's comparable "
                        "personal operational baseline: "
                        + " ".join(concerning_factors)
                    )

                if current_training_hours > 0:

                    explanation_summary += (
                        f" {current_training_hours:.1f} training hour(s) "
                        "are recorded in the current window as supporting "
                        "context."
                    )

                if leave_cancellations > 0:

                    explanation_summary += (
                        f" {leave_cancellations} leave cancellation(s) "
                        "are recorded as supporting context."
                    )


        # ----------------------------------------------
        # OPTIONAL VOLUNTARY CHECK-IN CONTEXT
        #
        # Guardrail:
        # - the check-in contributes only when the operational
        #   evidence gate is satisfied
        # - self-report alone cannot produce Elevated attention
        # - confidential free-text is not copied into the
        #   assessment explanation
        # - employment / leave / disciplinary decisions
        #   remain outside this assessment engine
        # ----------------------------------------------

        recent_checkin = (
            _recent_wellness_checkin(
                personnel_id,
                days=7
            )
        )

        if recent_checkin:

            if (
                evidence_status
                == "Sufficient Operational Evidence"
            ):

                evidence_status = (
                    "Operational + Check-In Context"
                )

            explanation_summary += (
                " A voluntary Wellness Check-In submitted within "
                "the last seven days is linked as supporting context. "
                + recent_checkin.summary_label
                + " were recorded. "
                "Self-report is weighted as supporting evidence and "
                "cannot produce Elevated attention without "
                "meaningful operational concern."
            )

            if (
                recent_checkin.support_request
                == "Yes"
            ):

                explanation_summary += (
                    " The personnel member also requested confidential "
                    "human welfare support; that request is routed separately "
                    "and does not change this assessment level."
                )


        # ----------------------------------------------
        # OPTIONAL CAMERA WELLNESS SCAN CONTEXT
        #
        # Guardrail:
        # - only a recently saved usable scan is linked
        # - HR / breathing are experimental wellness-support
        #   measurements, not clinical measurements
        # - the scan never independently sets or raises
        #   Low / Moderate / Elevated attention
        # ----------------------------------------------

        recent_scan = (
            _recent_usable_wellness_scan(
                personnel_id,
                hours=24
            )
        )

        if recent_scan:

            if (
                evidence_status
                == "Sufficient Operational Evidence"
            ):

                evidence_status = (
                    "Operational + Scan Context"
                )

            elif (
                evidence_status
                == "Operational + Check-In Context"
            ):

                evidence_status = (
                    "Operational + Check-In + Scan Context"
                )

            explanation_summary += (
                " A usable voluntary camera Wellness Scan saved within "
                "the last 24 hours is linked as optional supporting context. "
                "Camera-derived pulse and breathing estimates did not "
                "independently change the welfare-attention level."
            )

        # ----------------------------------------------
        # VERSIONED SINGLE SOURCE OF TRUTH
        # ----------------------------------------------

        latest_assessment = (
            WelfareAssessment.query
            .filter_by(
                personnel_id=personnel_id
            )
            .order_by(
                WelfareAssessment.assessment_version.desc()
            )
            .first()
        )

        next_version = (
            latest_assessment.assessment_version + 1
            if latest_assessment
            else 1
        )

        try:

            (
                WelfareAssessment.query
                .filter_by(
                    personnel_id=personnel_id,
                    is_current=True
                )
                .update({
                    "is_current": False
                })
            )

            new_assessment = WelfareAssessment(

                personnel_id=personnel_id,

                assessment_version=
                    next_version,

                attention_level=
                    attention_level,

                evidence_status=
                    evidence_status,

                confidence=
                    confidence,

                baseline_context=
                    baseline_context,

                baseline_version=
                    baseline_version,

                rules_version=
                    "VP-RULES-1.3",

                explanation_summary=
                    explanation_summary,

                is_current=True,

                assessment_time=
                    datetime.now(),

                created_at=
                    datetime.now()
            )

            db.session.add(
                new_assessment
            )

            # Flush first so the versioned assessment receives
            # its database ID before the replay snapshot is saved.
            db.session.flush()

            evidence_snapshot = (
                AssessmentEvidenceSnapshot(

                    assessment_id=
                        new_assessment.id,

                    personnel_id=
                        personnel_id,

                    assessment_version=
                        new_assessment.assessment_version,

                    rules_version=
                        new_assessment.rules_version,

                    baseline_version=
                        new_assessment.baseline_version,

                    baseline_context=
                        new_assessment.baseline_context,

                    evidence_status=
                        new_assessment.evidence_status,

                    current_record_count=
                        snapshot_current_record_count,

                    baseline_record_count=
                        snapshot_baseline_record_count,

                    duty_delta=
                        snapshot_duty_delta,

                    rest_delta=
                        snapshot_rest_delta,

                    recovery_day_delta=
                        snapshot_recovery_day_delta,

                    concerning_factor_count=
                        snapshot_factor_count,

                    severity_score=
                        snapshot_severity_score,

                    resulting_attention_level=
                        new_assessment.attention_level,

                    explanation_summary=
                        new_assessment.explanation_summary,

                    created_at=
                        datetime.now()
                )
            )

            db.session.add(
                evidence_snapshot
            )

            checkin_evidence = None

            if recent_checkin:

                checkin_evidence = (
                    AssessmentCheckInEvidence(

                        assessment_id=
                            new_assessment.id,

                        personnel_id=
                            personnel_id,

                        assessment_version=
                            new_assessment.assessment_version,

                        checkin_id=
                            recent_checkin.id,

                        checkin_submitted_at=
                            recent_checkin.submitted_at,

                        concern_score=
                            recent_checkin.concern_score,

                        concern_count=
                            recent_checkin.concern_count,

                        high_concern_count=
                            recent_checkin.high_concern_count,

                        support_requested=
                            (
                                recent_checkin.support_request
                                == "Yes"
                            ),

                        summary_label=
                            recent_checkin.summary_label,

                        created_at=
                            datetime.now()
                    )
                )

                db.session.add(
                    checkin_evidence
                )

            scan_evidence = None

            if recent_scan:

                scan_evidence = (
                    AssessmentScanEvidence(

                        assessment_id=
                            new_assessment.id,

                        personnel_id=
                            personnel_id,

                        assessment_version=
                            new_assessment.assessment_version,

                        scan_id=
                            recent_scan.id,

                        scan_saved_at=
                            recent_scan.saved_at,

                        heart_rate=
                            recent_scan.heart_rate,

                        breathing_rate=
                            recent_scan.breathing_rate,

                        signal_quality=
                            recent_scan.signal_quality,

                        confidence=
                            recent_scan.confidence,

                        created_at=
                            datetime.now()
                    )
                )

                db.session.add(
                    scan_evidence
                )

            # ------------------------------------------
            # VEERPULSE EARLY-WARNING ENGINE
            # ------------------------------------------

            analysis_result = (
                _build_veerpulse_analysis(
                    personnel_id=personnel_id,
                    assessment=new_assessment,
                    evidence_snapshot=evidence_snapshot,
                    recent_checkin=recent_checkin,
                    checkin_evidence=checkin_evidence,
                    scan_evidence=scan_evidence,
                    current_records=current_records,
                    historical_records=historical_records
                )
            )

            new_assessment.attention_level = (
                analysis_result.attention_level
            )

            new_assessment.confidence = (
                analysis_result.confidence
            )

            new_assessment.rules_version = (
                analysis_result.rules_version
            )

            new_assessment.explanation_summary = (
                analysis_result.analysis_reason
            )

            evidence_snapshot.rules_version = (
                analysis_result.rules_version
            )

            evidence_snapshot.resulting_attention_level = (
                analysis_result.attention_level
            )

            evidence_snapshot.explanation_summary = (
                analysis_result.analysis_reason
            )

            audit_event = AuditEvent(
                actor_role=
                    session.get("role")
                    or "personnel",

                actor_id=
                    session.get("officer_id")
                    or session.get("user_id")
                    or "officer",

                event_type=
                    "Welfare Assessment Recalculated",

                personnel_id=
                    personnel_id,

                resource_type=
                    "WelfareAssessment",

                resource_id=
                    str(new_assessment.id),

                event_summary=
                    (
                        "Assessment version "
                        + str(
                            new_assessment.assessment_version
                        )
                        + " recalculated with result "
                        + str(
                            new_assessment.attention_level
                        )
                        + " using rules "
                        + str(
                            new_assessment.rules_version
                        )
                        + "."
                    ),

                source_type=
                    "VeerPulse Application",

                created_at=
                    datetime.now()
            )

            db.session.add(
                audit_event
            )

            db.session.commit()

        except Exception:

            db.session.rollback()

            return jsonify({
                "success": False,
                "message":
                    "Assessment could not be recalculated."
            }), 500

        return jsonify({

            "success": True,

            "message":
                "Follow-up assessment recalculated successfully.",

            "assessment": {

                "id":
                    new_assessment.id,

                "personnel_id":
                    new_assessment.personnel_id,

                "assessment_version":
                    new_assessment.assessment_version,

                "attention_level":
                    new_assessment.attention_level,

                "evidence_status":
                    new_assessment.evidence_status,

                "confidence":
                    new_assessment.confidence,

                "baseline_context":
                    new_assessment.baseline_context,

                "baseline_version":
                    new_assessment.baseline_version,

                "rules_version":
                    new_assessment.rules_version,

                "evidence_snapshot_id":
                    evidence_snapshot.id,

                "checkin_evidence_id":
                    (
                        checkin_evidence.id
                        if checkin_evidence
                        else None
                    ),

                "recent_checkin_id":
                    (
                        recent_checkin.id
                        if recent_checkin
                        else None
                    ),

                "scan_evidence_id":
                    (
                        scan_evidence.id
                        if scan_evidence
                        else None
                    ),

                "recent_scan_id":
                    (
                        recent_scan.id
                        if recent_scan
                        else None
                    ),

                "explanation_summary":
                    new_assessment.explanation_summary,

                "analysis_result_id":
                    analysis_result.id,

                "early_warning_score":
                    analysis_result.overall_score,

                "operational_score":
                    analysis_result.operational_score,

                "checkin_score":
                    analysis_result.checkin_score,

                "trend_score":
                    analysis_result.trend_score,

                "trend_direction":
                    analysis_result.trend_direction,

                "trend_strength":
                    analysis_result.trend_strength,

                "signal_agreement":
                    analysis_result.signal_agreement,

                "signal_agreement_score":
                    (
                        agreement_result[
                            "agreement_score"
                        ]
                    ),

                "signal_agreement_direction":
                    (
                        agreement_result[
                            "direction"
                        ]
                    ),

                "signal_count":
                    (
                        agreement_result[
                            "signal_count"
                        ]
                    ),

                "analysis_confidence":
                    analysis_result.confidence,

                "recommendation":
                    analysis_result.recommendation,

                "assessment_time":
                    new_assessment.assessment_time.isoformat()
            }
        })


    # --------------------------------------------------
    # DECISION REPLAY
    # --------------------------------------------------

    @app.route(
        "/api/welfare-officer/decision-replay/<int:assessment_version>",
        methods=["GET"]
    )
    def decision_replay(assessment_version):

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        personnel_id = (
            request.args.get(
                "personnel_id"
            )
            or ""
        ).strip()

        if not personnel_id:

            return jsonify({
                "success": False,
                "message":
                    "Personnel ID is required for Decision Replay."
            }), 400

        assessment = (
            WelfareAssessment.query
            .filter_by(
                personnel_id=personnel_id,
                assessment_version=assessment_version
            )
            .order_by(
                WelfareAssessment.assessment_time.desc()
            )
            .first()
        )

        if not assessment:

            return jsonify({
                "success": False,
                "message":
                    "Assessment version not found."
            }), 404

        snapshot = (
            AssessmentEvidenceSnapshot.query
            .filter_by(
                personnel_id=personnel_id,
                assessment_version=assessment_version
            )
            .order_by(
                AssessmentEvidenceSnapshot.created_at.desc()
            )
            .first()
        )

        checkin_evidence = (
            AssessmentCheckInEvidence.query
            .filter_by(
                personnel_id=personnel_id,
                assessment_version=assessment_version
            )
            .order_by(
                AssessmentCheckInEvidence.created_at.desc()
            )
            .first()
        )

        scan_evidence = (
            AssessmentScanEvidence.query
            .filter_by(
                personnel_id=personnel_id,
                assessment_version=assessment_version
            )
            .order_by(
                AssessmentScanEvidence.created_at.desc()
            )
            .first()
        )

        if not snapshot:

            return jsonify({
                "success": True,
                "personnel_id":
                    personnel_id,
                "assessment_version":
                    assessment_version,
                "replay_status":
                    "Snapshot Unavailable",
                "message":
                    "This assessment predates Decision Replay snapshot storage.",
                "assessment": {
                    "attention_level":
                        assessment.attention_level,
                    "evidence_status":
                        assessment.evidence_status,
                    "confidence":
                        assessment.confidence,
                    "baseline_context":
                        assessment.baseline_context,
                    "baseline_version":
                        assessment.baseline_version,
                    "rules_version":
                        assessment.rules_version,
                    "explanation_summary":
                        assessment.explanation_summary,
                    "assessment_time":
                        assessment.assessment_time.isoformat()
                }
            })

        audit_event = AuditEvent(
            actor_role=
                session.get("role")
                or "welfare_officer",

            actor_id=
                session.get("officer_id")
                or session.get("user_id")
                or "authenticated_officer",

            event_type=
                "Decision Replay Accessed",

            personnel_id=
                personnel_id,

            resource_type=
                "AssessmentEvidenceSnapshot",

            resource_id=
                str(snapshot.id),

            event_summary=
                (
                    "Decision Replay accessed for assessment version "
                    + str(assessment_version)
                    + "."
                ),

            source_type=
                "VeerPulse Application",

            created_at=
                datetime.now()
        )

        db.session.add(
            audit_event
        )

        try:

            db.session.commit()

        except Exception:

            db.session.rollback()

            return jsonify({
                "success": False,
                "message":
                    "Decision Replay audit event could not be recorded."
            }), 500

        return jsonify({
            "success": True,
            "personnel_id":
                personnel_id,
            "assessment_version":
                assessment_version,
            "replay_status":
                "Reproducible Snapshot Available",
            "assessment": {
                "assessment_id":
                    assessment.id,
                "attention_level":
                    assessment.attention_level,
                "evidence_status":
                    assessment.evidence_status,
                "confidence":
                    assessment.confidence,
                "baseline_context":
                    assessment.baseline_context,
                "baseline_version":
                    assessment.baseline_version,
                "rules_version":
                    assessment.rules_version,
                "explanation_summary":
                    assessment.explanation_summary,
                "assessment_time":
                    assessment.assessment_time.isoformat()
            },
            "evidence_snapshot": {
                "snapshot_id":
                    snapshot.id,
                "current_record_count":
                    snapshot.current_record_count,
                "baseline_record_count":
                    snapshot.baseline_record_count,
                "duty_delta":
                    snapshot.duty_delta,
                "rest_delta":
                    snapshot.rest_delta,
                "recovery_day_delta":
                    snapshot.recovery_day_delta,
                "concerning_factor_count":
                    snapshot.concerning_factor_count,
                "severity_score":
                    snapshot.severity_score,
                "resulting_attention_level":
                    snapshot.resulting_attention_level,
                "evidence_status":
                    snapshot.evidence_status,
                "baseline_context":
                    snapshot.baseline_context,
                "baseline_version":
                    snapshot.baseline_version,
                "rules_version":
                    snapshot.rules_version,
                "explanation_summary":
                    snapshot.explanation_summary,
                "created_at":
                    snapshot.created_at.isoformat()
            },
            "voluntary_checkin_evidence": (
                {
                    "evidence_id":
                        checkin_evidence.id,
                    "checkin_id":
                        checkin_evidence.checkin_id,
                    "checkin_submitted_at":
                        checkin_evidence.checkin_submitted_at.isoformat(),
                    "summary_label":
                        checkin_evidence.summary_label,
                    "concern_score":
                        checkin_evidence.concern_score,
                    "concern_count":
                        checkin_evidence.concern_count,
                    "high_concern_count":
                        checkin_evidence.high_concern_count,
                    "support_requested":
                        checkin_evidence.support_requested,
                    "guardrail":
                        (
                            "Voluntary check-in context did not independently "
                            "set or raise the welfare-attention level."
                        )
                }
                if checkin_evidence
                else None
            ),
            "wellness_scan_evidence": (
                {
                    "evidence_id":
                        scan_evidence.id,
                    "scan_id":
                        scan_evidence.scan_id,
                    "scan_saved_at":
                        scan_evidence.scan_saved_at.isoformat(),
                    "heart_rate":
                        scan_evidence.heart_rate,
                    "breathing_rate":
                        scan_evidence.breathing_rate,
                    "signal_quality":
                        scan_evidence.signal_quality,
                    "confidence":
                        scan_evidence.confidence,
                    "guardrail":
                        (
                            "Experimental camera-derived wellness measurements "
                            "were supporting context only and did not independently "
                            "set or raise the welfare-attention level."
                        )
                }
                if scan_evidence
                else None
            )
        })


    # --------------------------------------------------
    # HRMS-STYLE OPERATIONAL RECORD IMPORT
    # --------------------------------------------------

    @app.route(
        "/api/welfare-officer/hrms/import-operational-records",
        methods=["POST"]
    )
    def import_hrms_operational_records():

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        uploaded_file = request.files.get(
            "file"
        )

        if (
            not uploaded_file
            or not uploaded_file.filename
        ):

            return jsonify({
                "success": False,
                "message":
                    "Please upload a CSV file."
            }), 400

        if (
            not uploaded_file.filename
            .lower()
            .endswith(".csv")
        ):

            return jsonify({
                "success": False,
                "message":
                    "Only CSV files are supported in this prototype import."
            }), 400

        try:

            raw_text = (
                uploaded_file
                .read()
                .decode("utf-8-sig")
            )

        except UnicodeDecodeError:

            return jsonify({
                "success": False,
                "message":
                    "The CSV must use UTF-8 text encoding."
            }), 400

        reader = csv.DictReader(
            StringIO(raw_text)
        )

        required_columns = {
            "personnel_id",
            "record_date",
            "context_type",
            "source_record_id"
        }

        available_columns = set(
            reader.fieldnames or []
        )

        missing_columns = sorted(
            required_columns
            - available_columns
        )

        if missing_columns:

            return jsonify({
                "success": False,
                "message":
                    "Missing required CSV columns.",
                "missing_columns":
                    missing_columns
            }), 400

        def parse_optional_float(
            value,
            field_name,
            row_number
        ):

            cleaned = (
                str(value).strip()
                if value is not None
                else ""
            )

            if cleaned == "":
                return None

            try:
                return float(cleaned)

            except ValueError:
                raise ValueError(
                    f"Row {row_number}: "
                    f"{field_name} must be numeric."
                )

        def parse_boolean(
            value,
            row_number
        ):

            cleaned = (
                str(value).strip().lower()
                if value is not None
                else ""
            )

            if cleaned in [
                "",
                "false",
                "0",
                "no",
                "n"
            ]:
                return False

            if cleaned in [
                "true",
                "1",
                "yes",
                "y"
            ]:
                return True

            raise ValueError(
                f"Row {row_number}: "
                "deployment_active must be true/false."
            )

        imported_count = 0
        updated_count = 0
        rejected_rows = []

        for row_number, row in enumerate(
            reader,
            start=2
        ):

            try:

                personnel_id = (
                    row.get("personnel_id")
                    or ""
                ).strip()

                record_date_text = (
                    row.get("record_date")
                    or ""
                ).strip()

                context_type = (
                    row.get("context_type")
                    or ""
                ).strip()

                source_record_id = (
                    row.get("source_record_id")
                    or ""
                ).strip()

                if (
                    not personnel_id
                    or not record_date_text
                    or not context_type
                    or not source_record_id
                ):

                    raise ValueError(
                        f"Row {row_number}: "
                        "required values cannot be blank."
                    )

                personnel_exists = (
                    Personnel.query
                    .filter_by(
                        personnel_id=personnel_id
                    )
                    .first()
                )

                if personnel_exists is None:

                    raise ValueError(
                        f"Row {row_number}: "
                        f"personnel_id {personnel_id} does not exist "
                        "in the Personnel master table."
                    )

                try:

                    record_date = datetime.strptime(
                        record_date_text,
                        "%Y-%m-%d"
                    ).date()

                except ValueError:

                    raise ValueError(
                        f"Row {row_number}: "
                        "record_date must use YYYY-MM-DD."
                    )

                duty_hours = parse_optional_float(
                    row.get("duty_hours"),
                    "duty_hours",
                    row_number
                )

                rest_hours = parse_optional_float(
                    row.get("rest_hours"),
                    "rest_hours",
                    row_number
                )

                training_hours = parse_optional_float(
                    row.get("training_hours"),
                    "training_hours",
                    row_number
                )

                for field_name, field_value in [
                    ("duty_hours", duty_hours),
                    ("rest_hours", rest_hours),
                    ("training_hours", training_hours)
                ]:

                    if (
                        field_value is not None
                        and field_value < 0
                    ):

                        raise ValueError(
                            f"Row {row_number}: "
                            f"{field_name} cannot be negative."
                        )

                deployment_active = parse_boolean(
                    row.get(
                        "deployment_active"
                    ),
                    row_number
                )

                leave_status = (
                    row.get("leave_status")
                    or ""
                ).strip() or None

                existing_record = (
                    OperationalRecord.query
                    .filter_by(
                        personnel_id=
                            personnel_id,
                        source_type=
                            "HRMS-Style Prototype Import",
                        source_record_id=
                            source_record_id
                    )
                    .first()
                )

                if existing_record:

                    existing_record.record_date = (
                        record_date
                    )

                    existing_record.duty_hours = (
                        duty_hours
                    )

                    existing_record.rest_hours = (
                        rest_hours
                    )

                    existing_record.context_type = (
                        context_type
                    )

                    existing_record.deployment_active = (
                        deployment_active
                    )

                    existing_record.training_hours = (
                        training_hours
                    )

                    existing_record.leave_status = (
                        leave_status
                    )

                    existing_record.is_valid = True

                    updated_count += 1

                else:

                    new_record = OperationalRecord(
                        personnel_id=
                            personnel_id,
                        record_date=
                            record_date,
                        duty_hours=
                            duty_hours,
                        rest_hours=
                            rest_hours,
                        context_type=
                            context_type,
                        deployment_active=
                            deployment_active,
                        training_hours=
                            training_hours,
                        leave_status=
                            leave_status,
                        source_type=
                            "HRMS-Style Prototype Import",
                        source_record_id=
                            source_record_id,
                        is_valid=True
                    )

                    db.session.add(
                        new_record
                    )

                    imported_count += 1

            except ValueError as error:

                rejected_rows.append({
                    "row":
                        row_number,
                    "reason":
                        str(error)
                })

        if (
            imported_count == 0
            and updated_count == 0
        ):

            db.session.rollback()

            return jsonify({
                "success": False,
                "message":
                    "No valid operational records were imported.",
                "rejected_rows":
                    rejected_rows
            }), 400

        audit_event = AuditEvent(
            actor_role=
                session.get("role")
                or "welfare_officer",

            actor_id=
                session.get("officer_id")
                or session.get("user_id")
                or "authenticated_officer",

            event_type=
                "HRMS-Style Operational Import",

            personnel_id=
                None,

            resource_type=
                "OperationalRecord",

            resource_id=
                None,

            event_summary=
                (
                    "HRMS-style CSV import completed: "
                    + str(imported_count)
                    + " imported, "
                    + str(updated_count)
                    + " updated, "
                    + str(len(rejected_rows))
                    + " rejected."
                ),

            source_type=
                "VeerPulse Application",

            created_at=
                datetime.now()
        )

        db.session.add(
            audit_event
        )

        try:

            db.session.commit()

        except Exception:

            db.session.rollback()

            return jsonify({
                "success": False,
                "message":
                    "The operational import could not be saved."
            }), 500

        return jsonify({
            "success": True,
            "message":
                "HRMS-style prototype operational import completed.",
            "source_type":
                "HRMS-Style Prototype Import",
            "imported_count":
                imported_count,
            "updated_count":
                updated_count,
            "rejected_count":
                len(rejected_rows),
            "rejected_rows":
                rejected_rows
        })


    # --------------------------------------------------
    # PRIVACY-GATED UNIT INSIGHTS
    # --------------------------------------------------

    @app.route(
        "/api/welfare-officer/unit-insights",
        methods=["GET"]
    )
    def welfare_officer_unit_insights_api():

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        minimum_group_size = 5

        current_assessments_raw = (
            WelfareAssessment.query
            .filter_by(
                is_current=True
            )
            .order_by(
                WelfareAssessment.assessment_time.desc()
            )
            .all()
        )

        # Defensive de-duplication by personnel_id.
        current_by_personnel = {}

        for assessment in current_assessments_raw:

            if (
                assessment.personnel_id
                not in current_by_personnel
            ):

                current_by_personnel[
                    assessment.personnel_id
                ] = assessment

        current_assessments = list(
            current_by_personnel.values()
        )

        personnel_ids = set(
            current_by_personnel.keys()
        )

        personnel_count = len(
            personnel_ids
        )

        if (
            personnel_count
            < minimum_group_size
        ):

            return jsonify({
                "success": True,
                "data_status":
                    "Insufficient Aggregate Data",
                "minimum_group_size":
                    minimum_group_size,
                "available_personnel_count":
                    personnel_count,
                "message":
                    (
                        "Unit-level analytics are withheld because the "
                        "minimum privacy-preserving group size has not "
                        "been reached."
                    ),
                "attention_distribution":
                    None,
                "operational_context":
                    None,
                "workflow_context":
                    None,
                "unit_breakdown":
                    [],
                "eligible_unit_count":
                    0,
                "suppressed_unit_count":
                    0
            })

        attention_counts = {
            "Low": 0,
            "Moderate": 0,
            "Elevated": 0,
            "Insufficient Data": 0
        }

        for assessment in current_assessments:

            level = (
                assessment.attention_level
                or "Insufficient Data"
            )

            if level not in attention_counts:
                level = "Insufficient Data"

            attention_counts[
                level
            ] += 1

        attention_distribution = {}

        for level, count in attention_counts.items():

            attention_distribution[
                level
            ] = {
                "count":
                    count,
                "percentage":
                    round(
                        (
                            count
                            / personnel_count
                        )
                        * 100,
                        1
                    )
            }

        personnel_rows = (
            Personnel.query
            .filter(
                Personnel.personnel_id.in_(
                    list(
                        personnel_ids
                    )
                )
            )
            .all()
        )

        personnel_map = {
            item.personnel_id:
                item
            for item in personnel_rows
        }

        today = datetime.now().date()

        period_start = (
            today
            - timedelta(days=6)
        )

        recent_records = (
            OperationalRecord.query
            .filter(
                OperationalRecord.personnel_id.in_(
                    list(
                        personnel_ids
                    )
                ),

                OperationalRecord.is_valid
                == True,

                OperationalRecord.record_date
                >= period_start,

                OperationalRecord.record_date
                <= today
            )
            .all()
        )

        duty_values = [
            float(item.duty_hours)
            for item in recent_records
            if item.duty_hours
            is not None
        ]

        rest_values = [
            float(item.rest_hours)
            for item in recent_records
            if item.rest_hours
            is not None
        ]

        training_values = [
            float(item.training_hours)
            for item in recent_records
            if item.training_hours
            is not None
        ]

        operational_personnel_ids = {
            item.personnel_id
            for item in recent_records
        }

        deployment_personnel_ids = {
            item.personnel_id
            for item in recent_records
            if item.deployment_active
        }

        recovery_record_count = sum(
            1
            for item in recent_records
            if (
                item.duty_hours
                is not None
                and float(
                    item.duty_hours
                ) == 0
            )
        )

        operational_context = {
            "period_start":
                period_start.isoformat(),

            "period_end":
                today.isoformat(),

            "valid_record_count":
                len(
                    recent_records
                ),

            "personnel_with_records":
                len(
                    operational_personnel_ids
                ),

            "coverage_percentage":
                round(
                    (
                        len(
                            operational_personnel_ids
                        )
                        / personnel_count
                    )
                    * 100,
                    1
                ),

            "average_duty_hours_per_record":
                (
                    round(
                        sum(
                            duty_values
                        )
                        / len(
                            duty_values
                        ),
                        1
                    )
                    if duty_values
                    else None
                ),

            "average_rest_hours_per_record":
                (
                    round(
                        sum(
                            rest_values
                        )
                        / len(
                            rest_values
                        ),
                        1
                    )
                    if rest_values
                    else None
                ),

            "total_training_hours":
                (
                    round(
                        sum(
                            training_values
                        ),
                        1
                    )
                    if training_values
                    else 0
                ),

            "deployment_active_personnel":
                len(
                    deployment_personnel_ids
                ),

            "recovery_record_count":
                recovery_record_count
        }

        open_support_count = (
            SupportRequest.query
            .filter(
                SupportRequest.personnel_id.in_(
                    list(
                        personnel_ids
                    )
                ),

                SupportRequest.status.in_(
                    [
                        "Open",
                        "In Review",
                        "In Progress"
                    ]
                )
            )
            .count()
        )

        follow_up_count = (
            WelfareAction.query
            .filter(
                WelfareAction.personnel_id.in_(
                    list(
                        personnel_ids
                    )
                ),

                WelfareAction.status
                == "Follow-Up"
            )
            .count()
        )

        workflow_context = {
            "open_support_requests":
                open_support_count,

            "follow_up_actions":
                follow_up_count,

            "current_priority_cases":
                (
                    attention_counts[
                        "Moderate"
                    ]
                    + attention_counts[
                        "Elevated"
                    ]
                )
        }

        unit_members = {}

        for personnel_id in personnel_ids:

            personnel = (
                personnel_map.get(
                    personnel_id
                )
            )

            unit_name = (
                (
                    personnel.unit
                    if personnel
                    else None
                )
                or "Unit Not Recorded"
            ).strip()

            unit_members.setdefault(
                unit_name,
                []
            ).append(
                personnel_id
            )

        eligible_units = []
        suppressed_unit_count = 0

        for unit_name, member_ids in sorted(
            unit_members.items()
        ):

            unit_size = len(
                member_ids
            )

            if (
                unit_size
                < minimum_group_size
            ):

                suppressed_unit_count += 1
                continue

            unit_counts = {
                "Low": 0,
                "Moderate": 0,
                "Elevated": 0,
                "Insufficient Data": 0
            }

            for personnel_id in member_ids:

                assessment = (
                    current_by_personnel.get(
                        personnel_id
                    )
                )

                level = (
                    assessment.attention_level
                    if assessment
                    else "Insufficient Data"
                )

                if (
                    level
                    not in unit_counts
                ):

                    level = (
                        "Insufficient Data"
                    )

                unit_counts[
                    level
                ] += 1

            eligible_units.append({
                "unit":
                    unit_name,

                "personnel_count":
                    unit_size,

                "attention_distribution": {
                    level: {
                        "count":
                            count,
                        "percentage":
                            round(
                                (
                                    count
                                    / unit_size
                                )
                                * 100,
                                1
                            )
                    }
                    for level, count
                    in unit_counts.items()
                }
            })

        return jsonify({
            "success": True,

            "data_status":
                "Aggregate Data Available",

            "minimum_group_size":
                minimum_group_size,

            "available_personnel_count":
                personnel_count,

            "attention_distribution":
                attention_distribution,

            "operational_context":
                operational_context,

            "workflow_context":
                workflow_context,

            "unit_breakdown":
                eligible_units,

            "eligible_unit_count":
                len(
                    eligible_units
                ),

            "suppressed_unit_count":
                suppressed_unit_count,

            "privacy_note":
                (
                    "Only aggregate statistics are returned. "
                    "Unit names and distributions are shown only "
                    "when that unit meets the minimum group size. "
                    "No individual ranking is produced."
                )
        })


    # --------------------------------------------------
    # WELFARE OFFICER AUDIT EVENTS
    # --------------------------------------------------

    @app.route(
        "/api/welfare-officer/audit-events",
        methods=["GET"]
    )
    def welfare_officer_audit_events():

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        personnel_id = (
            request.args.get("personnel_id")
            or ""
        ).strip()

        query = (
            AuditEvent.query
            .order_by(
                AuditEvent.created_at.desc()
            )
        )

        if personnel_id:
            query = query.filter_by(
                personnel_id=personnel_id
            )

        events = (
            query
            .limit(50)
            .all()
        )

        return jsonify({
            "success": True,
            "event_count":
                len(events),
            "events": [
                {
                    "id":
                        event.id,
                    "actor_role":
                        event.actor_role,
                    "actor_id":
                        event.actor_id,
                    "event_type":
                        event.event_type,
                    "personnel_id":
                        event.personnel_id,
                    "resource_type":
                        event.resource_type,
                    "resource_id":
                        event.resource_id,
                    "event_summary":
                        event.event_summary,
                    "source_type":
                        event.source_type,
                    "created_at":
                        event.created_at.isoformat()
                }
                for event in events
            ]
        })


    # --------------------------------------------------
    # WELFARE OFFICER DASHBOARD
    # --------------------------------------------------

    @app.route(
        "/welfare-officer-dashboard"
    )
    def welfare_officer_dashboard():

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return redirect(
                url_for(
                    "welfare_officer_login"
                )
            )

        total_personnel = (
            Personnel.query.count()
        )

        active_personnel = (
            Personnel.query
            .filter_by(
                employment_status="Active"
            )
            .count()
        )

        current_assessments = (
            WelfareAssessment.query
            .filter_by(
                is_current=True
            )
            .order_by(
                WelfareAssessment.assessment_time.desc()
            )
            .all()
        )

        attention_counts = {
            "Low": 0,
            "Moderate": 0,
            "Elevated": 0,
            "Insufficient Data": 0
        }

        assessed_personnel_ids = set()

        for assessment in current_assessments:

            assessed_personnel_ids.add(
                assessment.personnel_id
            )

            level = (
                assessment.attention_level
                or "Insufficient Data"
            )

            if level not in attention_counts:
                level = "Insufficient Data"

            attention_counts[level] += 1

        current_alerts = [
            item
            for item in current_assessments
            if item.attention_level
            in [
                "Moderate",
                "Elevated"
            ]
        ]

        level_order = {
            "Elevated": 0,
            "Moderate": 1
        }

        current_alerts.sort(
            key=lambda item: (
                level_order.get(
                    item.attention_level,
                    9
                ),
                -(
                    item.assessment_time.timestamp()
                    if item.assessment_time
                    else 0
                )
            )
        )

        open_support_requests = (
            SupportRequest.query
            .filter(
                SupportRequest.status.in_(
                    [
                        "Open",
                        "In Review",
                        "In Progress"
                    ]
                )
            )
            .order_by(
                SupportRequest.created_at.desc()
            )
            .all()
        )

        follow_up_actions = (
            WelfareAction.query
            .filter_by(
                status="Follow-Up"
            )
            .order_by(
                WelfareAction.updated_at.desc()
            )
            .all()
        )

        personnel_ids = set(
            assessed_personnel_ids
        )

        personnel_ids.update(
            item.personnel_id
            for item in open_support_requests
        )

        personnel_ids.update(
            item.personnel_id
            for item in follow_up_actions
        )

        personnel_map = {}

        if personnel_ids:

            personnel_records = (
                Personnel.query
                .filter(
                    Personnel.personnel_id.in_(
                        list(personnel_ids)
                    )
                )
                .all()
            )

            personnel_map = {
                item.personnel_id: item
                for item in personnel_records
            }

        recent_alerts = [
            {
                "assessment": item,
                "personnel":
                    personnel_map.get(
                        item.personnel_id
                    )
            }
            for item in current_alerts[:5]
        ]

        recent_support = [
            {
                "request": item,
                "personnel":
                    personnel_map.get(
                        item.personnel_id
                    )
            }
            for item in open_support_requests[:5]
        ]

        recent_followups = [
            {
                "action": item,
                "personnel":
                    personnel_map.get(
                        item.personnel_id
                    )
            }
            for item in follow_up_actions[:5]
        ]

        no_current_assessment = max(
            total_personnel
            - len(assessed_personnel_ids),
            0
        )

        dashboard_summary = {
            "total_personnel":
                total_personnel,

            "active_personnel":
                active_personnel,

            "current_assessments":
                len(current_assessments),

            "current_alerts":
                len(current_alerts),

            "priority_cases":
                len(current_alerts),

            "elevated":
                attention_counts["Elevated"],

            "open_support_requests":
                len(open_support_requests),

            "follow_ups":
                len(follow_up_actions),

            "no_current_assessment":
                no_current_assessment
        }

        return render_template(
            "welfare_officer_dashboard.html",
            dashboard_summary=
                dashboard_summary,
            attention_counts=
                attention_counts,
            recent_alerts=
                recent_alerts,
            recent_support=
                recent_support,
            recent_followups=
                recent_followups
        )


    # --------------------------------------------------
    # WELFARE OFFICER PERSONNEL OVERVIEW
    # --------------------------------------------------

    @app.route(
        "/welfare-officer-personnel-overview"
    )
    def welfare_officer_personnel_overview():

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return redirect(
                url_for(
                    "welfare_officer_login"
                )
            )

        personnel_records = (
            Personnel.query
            .order_by(
                Personnel.personnel_id.asc()
            )
            .all()
        )

        current_assessments = (
            WelfareAssessment.query
            .filter_by(
                is_current=True
            )
            .order_by(
                WelfareAssessment.assessment_time.desc()
            )
            .all()
        )

        assessment_map = {}
        for item in current_assessments:
            if item.personnel_id not in assessment_map:
                assessment_map[item.personnel_id] = item

        current_analyses = (
            WelfareAnalysisResult.query
            .filter_by(
                is_current=True
            )
            .order_by(
                WelfareAnalysisResult.updated_at.desc()
            )
            .all()
        )

        analysis_map = {}
        for item in current_analyses:
            if item.personnel_id not in analysis_map:
                analysis_map[item.personnel_id] = item

        analysis_ids = [
            item.id
            for item in analysis_map.values()
        ]

        factor_map = {}
        if analysis_ids:
            factor_rows = (
                AnalysisFactor.query
                .filter(
                    AnalysisFactor.analysis_result_id.in_(
                        analysis_ids
                    )
                )
                .all()
            )

            factor_map = {
                (
                    item.analysis_result_id,
                    item.factor_key
                ): item
                for item in factor_rows
            }

        personnel_rows = []

        for personnel in personnel_records:
            assessment = assessment_map.get(
                personnel.personnel_id
            )
            analysis = analysis_map.get(
                personnel.personnel_id
            )

            duty_factor = (
                factor_map.get((analysis.id, "duty_intensity"))
                if analysis
                else None
            )
            rest_factor = (
                factor_map.get((analysis.id, "average_rest"))
                if analysis
                else None
            )
            workload_factor = (
                factor_map.get((analysis.id, "checkin_workload"))
                if analysis
                else None
            )

            personnel_rows.append({
                "personnel": personnel,
                "assessment": assessment,
                "analysis": analysis,
                "attention_level": (
                    assessment.attention_level
                    if assessment
                    else "Insufficient Data"
                ),
                "duty_trend": (
                    duty_factor.direction
                    if duty_factor and duty_factor.direction
                    else "Insufficient Data"
                ),
                "rest_trend": (
                    rest_factor.direction
                    if rest_factor and rest_factor.direction
                    else "Insufficient Data"
                ),
                "workload": (
                    workload_factor.severity
                    if workload_factor and workload_factor.severity
                    else "Insufficient Data"
                ),
                "last_review": (
                    assessment.assessment_time
                    if assessment
                    else None
                )
            })

        attention_counts = {
            "Low": 0,
            "Moderate": 0,
            "Elevated": 0
        }

        for assessment in assessment_map.values():
            if assessment.attention_level in attention_counts:
                attention_counts[assessment.attention_level] += 1

        overview_summary = {
            "total_personnel": len(personnel_records),
            "active_personnel": sum(
                1
                for item in personnel_records
                if item.employment_status == "Active"
            ),
            "current_assessments": len(assessment_map),
            "without_current_assessment": max(
                len(personnel_records)
                - len(assessment_map),
                0
            ),
            "low": attention_counts["Low"],
            "moderate": attention_counts["Moderate"],
            "elevated": attention_counts["Elevated"]
        }

        return render_template(
            "welfare_officer_personnel_overview.html",
            personnel_rows=personnel_rows,
            overview_summary=overview_summary
        )



    # --------------------------------------------------
    # WELFARE OFFICER LEAVE MANAGEMENT PAGE
    # --------------------------------------------------

    @app.route(
        "/welfare-officer-leave-management"
    )
    def welfare_officer_leave_management():

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return redirect(
                url_for(
                    "welfare_officer_login"
                )
            )

        return render_template(
            "welfare_officer_leave_management.html"
        )


    # --------------------------------------------------
    # WELFARE OFFICER PERSONNEL MANAGEMENT API
    # --------------------------------------------------

    @app.route(
        "/api/welfare-officer/personnel",
        methods=["GET", "POST"]
    )
    def welfare_officer_personnel_api():

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        officer_id = (
            session.get("officer_id")
            or "officer"
        )

        if request.method == "GET":

            status_filter = (
                request.args.get("status")
                or ""
            ).strip()

            search_text = (
                request.args.get("q")
                or ""
            ).strip()

            raw_limit = (
                request.args.get(
                    "limit"
                )
            )

            limit = None

            if raw_limit not in (
                None,
                ""
            ):

                try:

                    limit = int(
                        raw_limit
                    )

                except (
                    TypeError,
                    ValueError
                ):

                    return jsonify({
                        "success": False,
                        "message":
                            "limit must be an integer."
                    }), 400

                limit = max(
                    1,
                    min(
                        limit,
                        100
                    )
                )

            query = Personnel.query

            if status_filter:
                query = query.filter_by(
                    employment_status=status_filter
                )

            if search_text:

                search_pattern = (
                    "%"
                    + search_text
                    + "%"
                )

                query = query.filter(
                    or_(
                        Personnel.personnel_id.ilike(
                            search_pattern
                        ),
                        Personnel.full_name.ilike(
                            search_pattern
                        ),
                        Personnel.unit.ilike(
                            search_pattern
                        ),
                        Personnel.department.ilike(
                            search_pattern
                        )
                    )
                )

            query = (
                query
                .order_by(
                    Personnel.personnel_id.asc()
                )
            )

            if limit is not None:

                query = (
                    query.limit(
                        limit
                    )
                )

            items = (
                query.all()
            )

            return jsonify({
                "success": True,
                "personnel_count": len(items),
                "personnel": [
                    _personnel_to_dict(item)
                    for item in items
                ]
            })

        data = request.get_json(
            silent=True
        ) or {}

        personnel_id = (
            data.get("personnel_id")
            or ""
        ).strip().upper()

        full_name = (
            data.get("full_name")
            or ""
        ).strip()

        initial_password = (
            data.get("initial_password")
            or ""
        )

        if (
            len(personnel_id) < 3
            or len(personnel_id) > 50
            or not personnel_id
                .replace("-", "")
                .replace("_", "")
                .isalnum()
        ):

            return jsonify({
                "success": False,
                "message":
                    "Personnel ID must use 3-50 letters, numbers, hyphens or underscores."
            }), 400

        if not full_name:

            return jsonify({
                "success": False,
                "message":
                    "Full name is required."
            }), 400

        if len(initial_password) < 8:

            return jsonify({
                "success": False,
                "message":
                    "Initial password must contain at least 8 characters."
            }), 400

        if (
            Personnel.query
            .filter_by(
                personnel_id=personnel_id
            )
            .first()
        ):

            return jsonify({
                "success": False,
                "message":
                    "That Personnel ID already exists."
            }), 409

        joining_date_value = (
            data.get("joining_date")
        )

        joining_date = (
            _parse_date(joining_date_value)
            if joining_date_value
            else None
        )

        if (
            joining_date_value
            and not joining_date
        ):

            return jsonify({
                "success": False,
                "message":
                    "Joining date must use YYYY-MM-DD."
            }), 400

        employment_status = (
            data.get("employment_status")
            or "Active"
        ).strip()

        allowed_statuses = [
            "Active",
            "Inactive",
            "Transferred",
            "Retired"
        ]

        if employment_status not in allowed_statuses:

            return jsonify({
                "success": False,
                "message":
                    "Please select a valid employment status."
            }), 400

        personnel = Personnel(
            personnel_id=personnel_id,
            full_name=full_name,
            rank_designation=(
                data.get("rank_designation")
                or ""
            ).strip() or None,
            unit=(
                data.get("unit")
                or ""
            ).strip() or None,
            department=(
                data.get("department")
                or ""
            ).strip() or None,
            posting_location=(
                data.get("posting_location")
                or ""
            ).strip() or None,
            joining_date=joining_date,
            email=(
                data.get("email")
                or ""
            ).strip() or None,
            phone=(
                data.get("phone")
                or ""
            ).strip() or None,
            password_hash=generate_password_hash(
                initial_password
            ),
            employment_status=employment_status,
            created_by=officer_id,
            updated_by=officer_id,
            created_at=datetime.now(),
            updated_at=datetime.now()
        )

        try:

            db.session.add(
                personnel
            )

            db.session.flush()

            _add_audit_event(
                actor_role="welfare_officer",
                actor_id=officer_id,
                event_type="Personnel Created",
                personnel_id=personnel_id,
                resource_type="Personnel",
                resource_id=personnel.id,
                event_summary=(
                    "Personnel record "
                    + personnel_id
                    + " created by authorized officer."
                )
            )

            db.session.commit()

        except Exception:

            db.session.rollback()

            return jsonify({
                "success": False,
                "message":
                    "The personnel record could not be created."
            }), 500

        return jsonify({
            "success": True,
            "message":
                "Personnel record created successfully.",
            "personnel":
                _personnel_to_dict(personnel)
        }), 201




    # --------------------------------------------------
    # WELFARE OFFICER PERSONNEL BULK EXCEL IMPORT
    # --------------------------------------------------

    @app.route(
        "/api/welfare-officer/personnel/bulk-template",
        methods=["GET"]
    )
    def welfare_officer_personnel_bulk_template():

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        if Workbook is None:

            return jsonify({
                "success": False,
                "message":
                    "Excel support is not installed. Install openpyxl."
            }), 500

        headers = [
            "personnel_id",
            "full_name",
            "rank_designation",
            "unit",
            "department",
            "posting_location",
            "joining_date",
            "email",
            "phone",
            "employment_status",
            "initial_password"
        ]

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Personnel"

        sheet.append(
            headers
        )

        for column_index, header in enumerate(
            headers,
            start=1
        ):
            cell = sheet.cell(
                row=1,
                column=column_index
            )
            cell.font = cell.font.copy(
                bold=True
            )

        sheet.freeze_panes = "A2"

        widths = {
            "A": 18,
            "B": 24,
            "C": 22,
            "D": 18,
            "E": 22,
            "F": 22,
            "G": 16,
            "H": 28,
            "I": 18,
            "J": 20,
            "K": 22
        }

        for column_letter, width in widths.items():
            sheet.column_dimensions[
                column_letter
            ].width = width

        instructions = workbook.create_sheet(
            "Instructions"
        )

        instruction_rows = [
            [
                "Column",
                "Requirement"
            ],
            [
                "personnel_id",
                "Required. Unique ID using letters, numbers, hyphens or underscores."
            ],
            [
                "full_name",
                "Required."
            ],
            [
                "rank_designation",
                "Optional."
            ],
            [
                "unit",
                "Optional."
            ],
            [
                "department",
                "Optional."
            ],
            [
                "posting_location",
                "Optional."
            ],
            [
                "joining_date",
                "Optional. Use YYYY-MM-DD or an Excel date."
            ],
            [
                "email",
                "Optional."
            ],
            [
                "phone",
                "Optional."
            ],
            [
                "employment_status",
                "Required/optional. Active, Inactive, Transferred or Retired. Blank defaults to Active."
            ],
            [
                "initial_password",
                "Required. Minimum 8 characters. It is hashed before storage."
            ]
        ]

        for row in instruction_rows:
            instructions.append(row)

        instructions.freeze_panes = "A2"
        instructions.column_dimensions["A"].width = 24
        instructions.column_dimensions["B"].width = 85

        output = BytesIO()
        workbook.save(
            output
        )
        output.seek(0)

        return send_file(
            output,
            as_attachment=True,
            download_name=
                "veerpulse_personnel_bulk_template.xlsx",
            mimetype=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            )
        )


    @app.route(
        "/api/welfare-officer/personnel/bulk-import",
        methods=["POST"]
    )
    def welfare_officer_personnel_bulk_import():

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        if load_workbook is None:

            return jsonify({
                "success": False,
                "message":
                    "Excel support is not installed. Install openpyxl."
            }), 500

        officer_id = (
            session.get("officer_id")
            or "officer"
        )

        upload = request.files.get(
            "file"
        )

        if (
            upload is None
            or not upload.filename
        ):

            return jsonify({
                "success": False,
                "message":
                    "Please select an Excel .xlsx file."
            }), 400

        filename = upload.filename.lower()

        if not filename.endswith(
            ".xlsx"
        ):

            return jsonify({
                "success": False,
                "message":
                    "Only .xlsx Excel files are supported."
            }), 400

        if (
            request.content_length
            and request.content_length
            > 5 * 1024 * 1024
        ):

            return jsonify({
                "success": False,
                "message":
                    "Excel file is too large. Maximum upload size is 5 MB."
            }), 400

        required_headers = [
            "personnel_id",
            "full_name",
            "rank_designation",
            "unit",
            "department",
            "posting_location",
            "joining_date",
            "email",
            "phone",
            "employment_status",
            "initial_password"
        ]

        try:

            workbook = load_workbook(
                upload.stream,
                read_only=True,
                data_only=True
            )

            sheet = workbook.active

        except Exception:

            return jsonify({
                "success": False,
                "message":
                    "The Excel file could not be read. Please use a valid .xlsx workbook."
            }), 400

        header_row = next(
            sheet.iter_rows(
                min_row=1,
                max_row=1,
                values_only=True
            ),
            None
        )

        if not header_row:

            return jsonify({
                "success": False,
                "message":
                    "The Excel file does not contain a header row."
            }), 400

        normalized_headers = [
            (
                str(value).strip().lower()
                if value is not None
                else ""
            )
            for value in header_row
        ]

        missing_headers = [
            header
            for header in required_headers
            if header not in normalized_headers
        ]

        if missing_headers:

            return jsonify({
                "success": False,
                "message":
                    "Excel columns do not match Personnel Management. Missing: "
                    + ", ".join(
                        missing_headers
                    ),
                "required_columns":
                    required_headers
            }), 400

        header_index = {
            header: normalized_headers.index(
                header
            )
            for header in required_headers
        }

        allowed_statuses = {
            "active": "Active",
            "inactive": "Inactive",
            "transferred": "Transferred",
            "retired": "Retired"
        }

        parsed_rows = []
        rejected_rows = []
        skipped_rows = []
        seen_ids = set()

        def cell_value(
            row,
            column_name
        ):

            index = header_index[
                column_name
            ]

            if index >= len(row):
                return None

            return row[index]


        def clean_text(
            value
        ):

            if value is None:
                return ""

            return str(
                value
            ).strip()


        for excel_row_number, row in enumerate(
            sheet.iter_rows(
                min_row=2,
                values_only=True
            ),
            start=2
        ):

            if excel_row_number > 2001:

                rejected_rows.append({
                    "row":
                        excel_row_number,
                    "personnel_id": "",
                    "reason":
                        "Maximum 2000 personnel rows are supported per import."
                })

                break

            if not any(
                value is not None
                and str(value).strip() != ""
                for value in row
            ):
                continue

            personnel_id = clean_text(
                cell_value(
                    row,
                    "personnel_id"
                )
            ).upper()

            full_name = clean_text(
                cell_value(
                    row,
                    "full_name"
                )
            )

            initial_password = clean_text(
                cell_value(
                    row,
                    "initial_password"
                )
            )

            if (
                len(personnel_id) < 3
                or len(personnel_id) > 50
                or not personnel_id
                    .replace("-", "")
                    .replace("_", "")
                    .isalnum()
            ):

                rejected_rows.append({
                    "row":
                        excel_row_number,
                    "personnel_id":
                        personnel_id,
                    "reason":
                        "Invalid Personnel ID."
                })

                continue

            if personnel_id in seen_ids:

                rejected_rows.append({
                    "row":
                        excel_row_number,
                    "personnel_id":
                        personnel_id,
                    "reason":
                        "Duplicate Personnel ID inside the Excel file."
                })

                continue

            seen_ids.add(
                personnel_id
            )

            if not full_name:

                rejected_rows.append({
                    "row":
                        excel_row_number,
                    "personnel_id":
                        personnel_id,
                    "reason":
                        "Full name is required."
                })

                continue

            if len(initial_password) < 8:

                rejected_rows.append({
                    "row":
                        excel_row_number,
                    "personnel_id":
                        personnel_id,
                    "reason":
                        "Initial password must contain at least 8 characters."
                })

                continue

            raw_status = clean_text(
                cell_value(
                    row,
                    "employment_status"
                )
            )

            status_key = (
                raw_status.lower()
                if raw_status
                else "active"
            )

            if status_key not in allowed_statuses:

                rejected_rows.append({
                    "row":
                        excel_row_number,
                    "personnel_id":
                        personnel_id,
                    "reason":
                        "Employment status must be Active, Inactive, Transferred or Retired."
                })

                continue

            joining_value = cell_value(
                row,
                "joining_date"
            )

            joining_date = None

            if isinstance(
                joining_value,
                datetime
            ):
                joining_date = (
                    joining_value.date()
                )

            elif isinstance(
                joining_value,
                date
            ):
                joining_date = (
                    joining_value
                )

            elif joining_value not in (
                None,
                ""
            ):

                joining_date = _parse_date(
                    clean_text(
                        joining_value
                    )
                )

                if joining_date is None:

                    rejected_rows.append({
                        "row":
                            excel_row_number,
                        "personnel_id":
                            personnel_id,
                        "reason":
                            "Joining date must use YYYY-MM-DD or a valid Excel date."
                    })

                    continue

            parsed_rows.append({
                "row":
                    excel_row_number,
                "personnel_id":
                    personnel_id,
                "full_name":
                    full_name,
                "rank_designation":
                    clean_text(
                        cell_value(
                            row,
                            "rank_designation"
                        )
                    ) or None,
                "unit":
                    clean_text(
                        cell_value(
                            row,
                            "unit"
                        )
                    ) or None,
                "department":
                    clean_text(
                        cell_value(
                            row,
                            "department"
                        )
                    ) or None,
                "posting_location":
                    clean_text(
                        cell_value(
                            row,
                            "posting_location"
                        )
                    ) or None,
                "joining_date":
                    joining_date,
                "email":
                    clean_text(
                        cell_value(
                            row,
                            "email"
                        )
                    ) or None,
                "phone":
                    clean_text(
                        cell_value(
                            row,
                            "phone"
                        )
                    ) or None,
                "employment_status":
                    allowed_statuses[
                        status_key
                    ],
                "initial_password":
                    initial_password
            })

        if not parsed_rows:

            return jsonify({
                "success": True,
                "message":
                    "No new personnel records were available to import.",
                "created_count": 0,
                "skipped_count":
                    len(skipped_rows),
                "rejected_count":
                    len(rejected_rows),
                "skipped_rows":
                    skipped_rows,
                "rejected_rows":
                    rejected_rows
            })

        candidate_ids = [
            item["personnel_id"]
            for item in parsed_rows
        ]

        existing_ids = {
            row[0]
            for row in (
                db.session.query(
                    Personnel.personnel_id
                )
                .filter(
                    Personnel.personnel_id.in_(
                        candidate_ids
                    )
                )
                .all()
            )
        }

        import_rows = []

        for item in parsed_rows:

            if (
                item["personnel_id"]
                in existing_ids
            ):

                skipped_rows.append({
                    "row":
                        item["row"],
                    "personnel_id":
                        item["personnel_id"],
                    "reason":
                        "Personnel ID already exists in PostgreSQL."
                })

                continue

            import_rows.append(
                item
            )

        created_personnel = []

        try:

            for item in import_rows:

                personnel = Personnel(
                    personnel_id=
                        item["personnel_id"],
                    full_name=
                        item["full_name"],
                    rank_designation=
                        item["rank_designation"],
                    unit=
                        item["unit"],
                    department=
                        item["department"],
                    posting_location=
                        item["posting_location"],
                    joining_date=
                        item["joining_date"],
                    email=
                        item["email"],
                    phone=
                        item["phone"],
                    password_hash=
                        generate_password_hash(
                            item[
                                "initial_password"
                            ]
                        ),
                    employment_status=
                        item[
                            "employment_status"
                        ],
                    created_by=
                        officer_id,
                    updated_by=
                        officer_id,
                    created_at=
                        datetime.now(),
                    updated_at=
                        datetime.now()
                )

                db.session.add(
                    personnel
                )

                db.session.flush()

                _add_audit_event(
                    actor_role=
                        "welfare_officer",
                    actor_id=
                        officer_id,
                    event_type=
                        "Personnel Bulk Imported",
                    personnel_id=
                        personnel.personnel_id,
                    resource_type=
                        "Personnel",
                    resource_id=
                        personnel.id,
                    event_summary=(
                        "Personnel record "
                        + personnel.personnel_id
                        + " created through authorized Excel bulk import."
                    )
                )

                created_personnel.append(
                    personnel.personnel_id
                )

            db.session.commit()

        except Exception:

            db.session.rollback()

            return jsonify({
                "success": False,
                "message":
                    "The Excel import could not be completed. No new bulk records were saved."
            }), 500

        return jsonify({
            "success": True,
            "message":
                "Excel personnel import completed.",
            "created_count":
                len(created_personnel),
            "skipped_count":
                len(skipped_rows),
            "rejected_count":
                len(rejected_rows),
            "created_personnel":
                created_personnel,
            "skipped_rows":
                skipped_rows,
            "rejected_rows":
                rejected_rows,
            "required_columns":
                required_headers
        })


    @app.route(
        "/api/welfare-officer/personnel/<string:personnel_id>",
        methods=["GET", "PUT"]
    )
    def welfare_officer_personnel_detail_api(
        personnel_id
    ):

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        officer_id = (
            session.get("officer_id")
            or "officer"
        )

        personnel = (
            Personnel.query
            .filter_by(
                personnel_id=personnel_id
            )
            .first()
        )

        if not personnel:

            return jsonify({
                "success": False,
                "message":
                    "Personnel record not found."
            }), 404

        if request.method == "GET":

            return jsonify({
                "success": True,
                "personnel":
                    _personnel_to_dict(personnel)
            })

        data = request.get_json(
            silent=True
        ) or {}

        allowed_statuses = [
            "Active",
            "Inactive",
            "Transferred",
            "Retired"
        ]

        changed_fields = []

        text_fields = [
            "full_name",
            "rank_designation",
            "unit",
            "department",
            "posting_location",
            "email",
            "phone"
        ]

        for field_name in text_fields:

            if field_name not in data:
                continue

            value = (
                data.get(field_name)
                or ""
            ).strip() or None

            if (
                field_name == "full_name"
                and not value
            ):

                return jsonify({
                    "success": False,
                    "message":
                        "Full name cannot be blank."
                }), 400

            if getattr(
                personnel,
                field_name
            ) != value:

                setattr(
                    personnel,
                    field_name,
                    value
                )

                changed_fields.append(
                    field_name
                )

        if "joining_date" in data:

            raw_joining_date = data.get(
                "joining_date"
            )

            joining_date = (
                _parse_date(raw_joining_date)
                if raw_joining_date
                else None
            )

            if (
                raw_joining_date
                and not joining_date
            ):

                return jsonify({
                    "success": False,
                    "message":
                        "Joining date must use YYYY-MM-DD."
                }), 400

            if personnel.joining_date != joining_date:
                personnel.joining_date = joining_date
                changed_fields.append(
                    "joining_date"
                )

        if "employment_status" in data:

            employment_status = (
                data.get("employment_status")
                or ""
            ).strip()

            if employment_status not in allowed_statuses:

                return jsonify({
                    "success": False,
                    "message":
                        "Please select a valid employment status."
                }), 400

            if (
                personnel.employment_status
                != employment_status
            ):

                personnel.employment_status = (
                    employment_status
                )

                changed_fields.append(
                    "employment_status"
                )

        if "new_password" in data:

            new_password = (
                data.get("new_password")
                or ""
            )

            if new_password:

                if len(new_password) < 8:

                    return jsonify({
                        "success": False,
                        "message":
                            "New password must contain at least 8 characters."
                    }), 400

                personnel.password_hash = (
                    generate_password_hash(
                        new_password
                    )
                )

                changed_fields.append(
                    "password_reset"
                )

        personnel.updated_by = officer_id
        personnel.updated_at = datetime.now()

        if not changed_fields:

            return jsonify({
                "success": True,
                "message":
                    "No personnel fields changed.",
                "personnel":
                    _personnel_to_dict(personnel)
            })

        _add_audit_event(
            actor_role="welfare_officer",
            actor_id=officer_id,
            event_type="Personnel Updated",
            personnel_id=personnel.personnel_id,
            resource_type="Personnel",
            resource_id=personnel.id,
            event_summary=(
                "Personnel record updated. Changed fields: "
                + ", ".join(changed_fields)
                + "."
            )
        )

        try:
            db.session.commit()
        except Exception:
            db.session.rollback()
            return jsonify({
                "success": False,
                "message":
                    "The personnel record could not be updated."
            }), 500

        return jsonify({
            "success": True,
            "message":
                "Personnel record updated successfully.",
            "personnel":
                _personnel_to_dict(personnel)
        })


    # --------------------------------------------------
    # WELFARE OFFICER OPERATIONAL RECORD MANAGEMENT
    # --------------------------------------------------

    @app.route(
        "/api/welfare-officer/personnel/<string:personnel_id>/operational-records",
        methods=["GET", "POST"]
    )
    def welfare_officer_operational_records_api(
        personnel_id
    ):

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        personnel = (
            Personnel.query
            .filter_by(
                personnel_id=personnel_id
            )
            .first()
        )

        if not personnel:

            return jsonify({
                "success": False,
                "message":
                    "Personnel record not found."
            }), 404

        if request.method == "GET":

            records = (
                OperationalRecord.query
                .filter_by(
                    personnel_id=personnel_id
                )
                .order_by(
                    OperationalRecord.record_date.desc()
                )
                .limit(120)
                .all()
            )

            return jsonify({
                "success": True,
                "personnel_id": personnel_id,
                "records": [
                    {
                        "id": item.id,
                        "record_date": item.record_date.isoformat(),
                        "duty_hours": item.duty_hours,
                        "rest_hours": item.rest_hours,
                        "context_type": item.context_type,
                        "deployment_active": item.deployment_active,
                        "training_hours": item.training_hours,
                        "leave_status": item.leave_status,
                        "source_type": item.source_type,
                        "source_record_id": item.source_record_id,
                        "is_valid": item.is_valid
                    }
                    for item in records
                ]
            })

        data = request.get_json(
            silent=True
        ) or {}

        record_date = _parse_date(
            data.get("record_date")
        )

        if not record_date:

            return jsonify({
                "success": False,
                "message":
                    "Record date must use YYYY-MM-DD."
            }), 400

        existing = (
            OperationalRecord.query
            .filter_by(
                personnel_id=personnel_id,
                record_date=record_date,
                is_valid=True
            )
            .first()
        )

        if existing:

            return jsonify({
                "success": False,
                "message":
                    "A valid operational record already exists for that date. Edit the existing record instead."
            }), 409

        def optional_number(name):
            value = data.get(name)
            if value in [None, ""]:
                return None
            try:
                number = float(value)
            except (TypeError, ValueError):
                raise ValueError(
                    name + " must be numeric."
                )
            if number < 0 or number > 24:
                raise ValueError(
                    name + " must be between 0 and 24."
                )
            return number

        try:
            duty_hours = optional_number(
                "duty_hours"
            )
            rest_hours = optional_number(
                "rest_hours"
            )
            training_hours = optional_number(
                "training_hours"
            )
        except ValueError as error:
            return jsonify({
                "success": False,
                "message": str(error)
            }), 400

        if (
            duty_hours is not None
            and rest_hours is not None
            and duty_hours + rest_hours > 24.01
        ):

            return jsonify({
                "success": False,
                "message":
                    "Duty hours plus rest hours cannot exceed 24 hours."
            }), 400

        context_type = (
            data.get("context_type")
            or "Station Duty"
        ).strip()

        deployment_active = data.get(
            "deployment_active",
            False
        )

        if not isinstance(
            deployment_active,
            bool
        ):

            return jsonify({
                "success": False,
                "message":
                    "deployment_active must be true or false."
            }), 400

        record = OperationalRecord(
            personnel_id=personnel_id,
            record_date=record_date,
            duty_hours=duty_hours,
            rest_hours=rest_hours,
            context_type=context_type,
            deployment_active=deployment_active,
            training_hours=training_hours,
            leave_status=(
                data.get("leave_status")
                or ""
            ).strip() or None,
            source_type="Officer Entry",
            source_record_id=(
                "OFFICER-"
                + personnel_id
                + "-"
                + record_date.strftime("%Y%m%d")
            ),
            is_valid=True,
            created_at=datetime.now()
        )

        try:

            db.session.add(
                record
            )

            db.session.flush()

            _add_audit_event(
                actor_role="welfare_officer",
                actor_id=(
                    session.get("officer_id")
                    or "officer"
                ),
                event_type="Operational Record Created",
                personnel_id=personnel_id,
                resource_type="OperationalRecord",
                resource_id=record.id,
                event_summary=(
                    "Operational record created for "
                    + record_date.isoformat()
                    + "."
                )
            )

            db.session.commit()

        except Exception:

            db.session.rollback()

            return jsonify({
                "success": False,
                "message":
                    "The operational record could not be created."
            }), 500

        return jsonify({
            "success": True,
            "message":
                "Operational record created successfully.",
            "record_id": record.id
        }), 201


    @app.route(
        "/api/welfare-officer/operational-records/<int:record_id>",
        methods=["PUT"]
    )
    def welfare_officer_update_operational_record(
        record_id
    ):

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        record = (
            OperationalRecord.query
            .filter_by(
                id=record_id
            )
            .first()
        )

        if not record:

            return jsonify({
                "success": False,
                "message":
                    "Operational record not found."
            }), 404

        data = request.get_json(
            silent=True
        ) or {}

        original_date = record.record_date
        changed_fields = []

        if "record_date" in data:
            new_date = _parse_date(
                data.get("record_date")
            )
            if not new_date:
                return jsonify({
                    "success": False,
                    "message":
                        "Record date must use YYYY-MM-DD."
                }), 400
            if new_date != record.record_date:
                duplicate = (
                    OperationalRecord.query
                    .filter(
                        OperationalRecord.id != record.id,
                        OperationalRecord.personnel_id
                        == record.personnel_id,
                        OperationalRecord.record_date
                        == new_date,
                        OperationalRecord.is_valid
                        == True
                    )
                    .first()
                )
                if duplicate:
                    return jsonify({
                        "success": False,
                        "message":
                            "Another valid operational record already exists on that date."
                    }), 409
                record.record_date = new_date
                changed_fields.append(
                    "record_date"
                )

        for field_name in [
            "duty_hours",
            "rest_hours",
            "training_hours"
        ]:

            if field_name not in data:
                continue

            raw_value = data.get(
                field_name
            )

            if raw_value in [None, ""]:
                value = None
            else:
                try:
                    value = float(
                        raw_value
                    )
                except (TypeError, ValueError):
                    return jsonify({
                        "success": False,
                        "message":
                            field_name + " must be numeric."
                    }), 400

                if value < 0 or value > 24:
                    return jsonify({
                        "success": False,
                        "message":
                            field_name + " must be between 0 and 24."
                    }), 400

            if getattr(record, field_name) != value:
                setattr(record, field_name, value)
                changed_fields.append(
                    field_name
                )

        if (
            record.duty_hours is not None
            and record.rest_hours is not None
            and record.duty_hours
                + record.rest_hours
                > 24.01
        ):

            return jsonify({
                "success": False,
                "message":
                    "Duty hours plus rest hours cannot exceed 24 hours."
            }), 400

        for field_name in [
            "context_type",
            "leave_status"
        ]:

            if field_name in data:
                value = (
                    data.get(field_name)
                    or ""
                ).strip() or None
                if (
                    field_name == "context_type"
                    and not value
                ):
                    return jsonify({
                        "success": False,
                        "message":
                            "Context type cannot be blank."
                    }), 400
                if getattr(record, field_name) != value:
                    setattr(record, field_name, value)
                    changed_fields.append(
                        field_name
                    )

        if "deployment_active" in data:
            deployment_active = data.get(
                "deployment_active"
            )
            if not isinstance(
                deployment_active,
                bool
            ):
                return jsonify({
                    "success": False,
                    "message":
                        "deployment_active must be true or false."
                }), 400
            if (
                record.deployment_active
                != deployment_active
            ):
                record.deployment_active = (
                    deployment_active
                )
                changed_fields.append(
                    "deployment_active"
                )

        if not changed_fields:
            return jsonify({
                "success": True,
                "message":
                    "No operational fields changed.",
                "record_id": record.id
            })

        record.source_type = "Officer Edit"

        _add_audit_event(
            actor_role="welfare_officer",
            actor_id=(
                session.get("officer_id")
                or "officer"
            ),
            event_type="Operational Record Updated",
            personnel_id=record.personnel_id,
            resource_type="OperationalRecord",
            resource_id=record.id,
            event_summary=(
                "Operational record "
                + str(record.id)
                + " updated. Changed fields: "
                + ", ".join(changed_fields)
                + ". Previous record date: "
                + original_date.isoformat()
                + "."
            )
        )

        try:
            db.session.commit()
        except Exception:
            db.session.rollback()
            return jsonify({
                "success": False,
                "message":
                    "The operational record could not be updated."
            }), 500

        return jsonify({
            "success": True,
            "message":
                "Operational record updated successfully.",
            "record_id": record.id
        })


    # --------------------------------------------------
    # WELFARE OFFICER LEAVE MANAGEMENT API
    # --------------------------------------------------

    @app.route(
        "/api/welfare-officer/leave-requests",
        methods=["GET"]
    )
    def welfare_officer_leave_requests_api():

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        status_filter = (
            request.args.get("status")
            or ""
        ).strip()

        personnel_filter = (
            request.args.get("personnel_id")
            or ""
        ).strip()

        query = LeaveRequest.query

        if status_filter:
            query = query.filter_by(
                status=status_filter
            )

        if personnel_filter:
            query = query.filter_by(
                personnel_id=personnel_filter
            )

        items = (
            query
            .order_by(
                LeaveRequest.created_at.desc()
            )
            .all()
        )

        counts = {
            "Pending": 0,
            "Approved": 0,
            "Denied": 0,
            "Cancelled": 0
        }

        all_items = LeaveRequest.query.all()

        for item in all_items:
            if item.status in counts:
                counts[item.status] += 1

        return jsonify({
            "success": True,
            "counts": counts,
            "leave_requests": [
                _leave_to_dict(item)
                for item in items
            ]
        })


    @app.route(
        "/api/welfare-officer/leave-requests/<int:leave_id>/decision",
        methods=["POST"]
    )
    def welfare_officer_leave_decision(
        leave_id
    ):

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        leave_item = (
            LeaveRequest.query
            .filter_by(
                id=leave_id
            )
            .first()
        )

        if not leave_item:

            return jsonify({
                "success": False,
                "message":
                    "Leave request not found."
            }), 404

        if leave_item.status != "Pending":

            return jsonify({
                "success": False,
                "message":
                    "Only a pending leave request can receive a new decision."
            }), 409

        data = request.get_json(
            silent=True
        ) or {}

        decision = (
            data.get("decision")
            or ""
        ).strip()

        decision_reason = (
            data.get("decision_reason")
            or ""
        ).strip()

        decision_note = (
            data.get("decision_note")
            or ""
        ).strip() or None

        if decision not in [
            "Approved",
            "Denied"
        ]:

            return jsonify({
                "success": False,
                "message":
                    "Decision must be Approved or Denied."
            }), 400

        if (
            decision == "Denied"
            and len(decision_reason) < 3
        ):

            return jsonify({
                "success": False,
                "message":
                    "A clear denial reason is required."
            }), 400

        officer_id = (
            session.get("officer_id")
            or "officer"
        )

        leave_item.status = decision
        leave_item.decision_reason = (
            decision_reason
            or (
                "Approved by authorized officer"
                if decision == "Approved"
                else None
            )
        )
        leave_item.decision_note = (
            decision_note
        )
        leave_item.decided_by = officer_id
        leave_item.decided_at = datetime.now()
        leave_item.updated_at = datetime.now()

        try:

            if decision == "Approved":
                _approved_leave_to_operational_records(
                    leave_item
                )

            _add_audit_event(
                actor_role="welfare_officer",
                actor_id=officer_id,
                event_type=(
                    "Leave Request Approved"
                    if decision == "Approved"
                    else "Leave Request Denied"
                ),
                personnel_id=leave_item.personnel_id,
                resource_type="LeaveRequest",
                resource_id=leave_item.id,
                event_summary=(
                    leave_item.request_code
                    + " decision: "
                    + decision
                    + "."
                )
            )

            db.session.commit()

        except Exception:

            db.session.rollback()

            return jsonify({
                "success": False,
                "message":
                    "The leave decision could not be saved."
            }), 500

        return jsonify({
            "success": True,
            "message":
                "Leave decision saved successfully.",
            "leave_request":
                _leave_to_dict(leave_item)
        })


    @app.route(
        "/api/welfare-officer/personnel/<string:personnel_id>/leave-metrics",
        methods=["GET"]
    )
    def welfare_officer_personnel_leave_metrics(
        personnel_id
    ):

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        try:
            window_days = int(
                request.args.get(
                    "days",
                    90
                )
            )
        except ValueError:
            window_days = 90

        window_days = max(
            30,
            min(window_days, 365)
        )

        window_start = (
            datetime.now()
            - timedelta(days=window_days)
        )

        items = (
            LeaveRequest.query
            .filter(
                LeaveRequest.personnel_id
                == personnel_id,
                LeaveRequest.created_at
                >= window_start
            )
            .order_by(
                LeaveRequest.created_at.asc()
            )
            .all()
        )

        counts = {
            "Pending": 0,
            "Approved": 0,
            "Denied": 0,
            "Cancelled": 0
        }

        for item in items:
            if item.status in counts:
                counts[item.status] += 1

        total = len(items)

        approved_days = sum(
            item.requested_days
            for item in items
            if item.status == "Approved"
        )

        gaps = [
            (
                items[index].created_at
                - items[index - 1].created_at
            ).total_seconds()
            / 86400
            for index in range(1, len(items))
        ]

        return jsonify({
            "success": True,
            "metrics": {
                "personnel_id": personnel_id,
                "window_days": window_days,
                "applications": total,
                "pending": counts["Pending"],
                "approved": counts["Approved"],
                "denied": counts["Denied"],
                "cancelled": counts["Cancelled"],
                "approved_leave_days": approved_days,
                "applications_per_30_days": (
                    round(
                        total * 30 / window_days,
                        2
                    )
                    if window_days
                    else 0
                ),
                "average_days_between_requests": (
                    round(
                        sum(gaps) / len(gaps),
                        1
                    )
                    if gaps
                    else None
                ),
                "denial_rate_percent": (
                    round(
                        counts["Denied"]
                        * 100
                        / total,
                        1
                    )
                    if total
                    else 0
                ),
                "cancellation_rate_percent": (
                    round(
                        counts["Cancelled"]
                        * 100
                        / total,
                        1
                    )
                    if total
                    else 0
                ),
                "interpretation_note": (
                    "Leave frequency is operational context only and is not an automatic welfare-risk or disciplinary score."
                )
            }
        })


    # --------------------------------------------------
    # WELFARE OFFICER ALERTS
    # --------------------------------------------------

    @app.route(
        "/welfare-officer-alerts"
    )
    def welfare_officer_alerts():

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return redirect(
                url_for(
                    "welfare_officer_login"
                )
            )

        # STEP 1: Multi-person live alert queue.
        # Every current assessment is read from PostgreSQL and alerts are
        # shown only for Moderate / Elevated welfare-attention levels.
        current_assessments = (
            WelfareAssessment.query
            .filter_by(
                is_current=True
            )
            .order_by(
                WelfareAssessment.assessment_time.desc()
            )
            .all()
        )

        assessment_counts = {
            "Low": 0,
            "Moderate": 0,
            "Elevated": 0,
            "Insufficient Data": 0
        }

        for assessment in current_assessments:

            level = (
                assessment.attention_level
                or "Insufficient Data"
            )

            if level not in assessment_counts:
                level = "Insufficient Data"

            assessment_counts[level] += 1

        alert_assessments = [
            assessment
            for assessment in current_assessments
            if assessment.attention_level
            in ["Moderate", "Elevated"]
        ]

        # Elevated first, then Moderate; newest assessment first inside
        # each attention level.
        alert_assessments.sort(
            key=lambda assessment: (
                0
                if assessment.attention_level
                == "Elevated"
                else 1,
                -assessment.assessment_time.timestamp()
                if assessment.assessment_time
                else 0
            )
        )

        personnel_ids = list({
            assessment.personnel_id
            for assessment in alert_assessments
            if assessment.personnel_id
        })

        if personnel_ids:
            personnel_items = (
                Personnel.query
                .filter(
                    Personnel.personnel_id.in_(
                        personnel_ids
                    )
                )
                .all()
            )
        else:
            personnel_items = []

        personnel_map = {
            item.personnel_id: item
            for item in personnel_items
        }

        alert_counts = {
            "total": len(alert_assessments),
            "Moderate": assessment_counts["Moderate"],
            "Elevated": assessment_counts["Elevated"],
            "current_assessments": len(current_assessments)
        }

        return render_template(
            "welfare_officer_alerts.html",
            current_assessments=current_assessments,
            alert_assessments=alert_assessments,
            assessment_counts=assessment_counts,
            alert_counts=alert_counts,
            personnel_map=personnel_map
        )


    # --------------------------------------------------
    # WELFARE OFFICER SUPPORT REQUESTS
    # --------------------------------------------------

    @app.route(
        "/welfare-officer-support-requests"
    )
    def welfare_officer_support_requests():

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return redirect(
                url_for(
                    "welfare_officer_login"
                )
            )

        support_requests = (
            SupportRequest.query
            .order_by(
                SupportRequest.created_at.desc()
            )
            .all()
        )

        open_count = sum(
            1
            for item in support_requests
            if item.status == "Open"
        )

        in_progress_count = sum(
            1
            for item in support_requests
            if item.status == "In Progress"
        )

        closed_count = sum(
            1
            for item in support_requests
            if item.status == "Closed"
        )

        return render_template(
            "welfare_officer_support_requests.html",
            support_requests=support_requests,
            open_count=open_count,
            in_progress_count=in_progress_count,
            closed_count=closed_count
        )


    # --------------------------------------------------
    # UPDATE SUPPORT REQUEST STATUS
    # --------------------------------------------------

    @app.route(
        "/api/welfare-officer/support-request/<int:request_id>/status",
        methods=["POST"]
    )
    def update_support_request_status(
        request_id
    ):

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized"
            }), 401

        data = request.get_json(
            silent=True
        ) or {}

        new_status = (
            data.get("status")
            or ""
        ).strip()

        allowed_statuses = [
            "Open",
            "In Progress",
            "Closed"
        ]

        if new_status not in allowed_statuses:

            return jsonify({
                "success": False,
                "message":
                    "Please select a valid support request status."
            }), 400

        support_request = (
            SupportRequest.query
            .filter_by(
                id=request_id
            )
            .first()
        )

        if not support_request:

            return jsonify({
                "success": False,
                "message":
                    "Support request not found."
            }), 404

        support_request.status = (
            new_status
        )

        support_request.updated_at = (
            datetime.now()
        )

        try:

            db.session.commit()

        except Exception:

            db.session.rollback()

            return jsonify({
                "success": False,
                "message":
                    "Could not update the support request."
            }), 500

        return jsonify({

            "success": True,

            "message":
                "Support request status updated successfully.",

            "request": {

                "id":
                    support_request.id,

                "personnel_id":
                    support_request.personnel_id,

                "status":
                    support_request.status,

                "updated_at":
                    support_request.updated_at.strftime(
                        "%d %b %Y, %I:%M %p"
                    )
            }
        })


    # --------------------------------------------------
    # WELFARE OFFICER CASE REVIEW
    # --------------------------------------------------

    @app.route(
        "/welfare-officer-case-review"
    )
    @app.route(
        "/welfare-officer-case-review/<string:personnel_id>"
    )
    def welfare_officer_case_review(
        personnel_id=None
    ):

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return redirect(
                url_for(
                    "welfare_officer_login"
                )
            )

        # Case Review must always be opened for an explicitly
        # selected personnel record. Never silently substitute another
        # person's ID.
        if not personnel_id:

            personnel_id = (
                request.args.get(
                    "personnel_id"
                )
                or ""
            ).strip()

        if not personnel_id:

            return redirect(
                url_for(
                    "welfare_officer_priority_review"
                )
            )

        selected_personnel = (
            Personnel.query
            .filter_by(
                personnel_id=personnel_id
            )
            .first()
        )

        if selected_personnel is None:

            return redirect(
                url_for(
                    "welfare_officer_priority_review"
                )
            )

        personnel_id = str(
            personnel_id
        ).strip()

        personnel = (
            Personnel.query
            .filter_by(
                personnel_id=personnel_id
            )
            .first()
        )

        current_assessment = (
            WelfareAssessment.query
            .filter_by(
                personnel_id=personnel_id,
                is_current=True
            )
            .order_by(
                WelfareAssessment.assessment_time.desc()
            )
            .first()
        )

        latest_checkin = (
            WellnessCheckIn.query
            .filter_by(
                personnel_id=personnel_id
            )
            .order_by(
                WellnessCheckIn.submitted_at.desc()
            )
            .first()
        )

        latest_scan = (
            WellnessScan.query
            .filter(
                WellnessScan.personnel_id
                == personnel_id,

                WellnessScan.confidence.in_(
                    [
                        "High",
                        "Medium"
                    ]
                )
            )
            .order_by(
                WellnessScan.saved_at.desc()
            )
            .first()
        )

        assessment_checkin_evidence = None
        assessment_scan_evidence = None

        if current_assessment:

            assessment_checkin_evidence = (
                AssessmentCheckInEvidence.query
                .filter_by(
                    assessment_id=
                        current_assessment.id,
                    personnel_id=
                        personnel_id
                )
                .order_by(
                    AssessmentCheckInEvidence.created_at.desc()
                )
                .first()
            )

            if assessment_checkin_evidence is None:

                assessment_checkin_evidence = (
                    AssessmentCheckInEvidence.query
                    .filter_by(
                        personnel_id=
                            personnel_id,
                        assessment_version=
                            current_assessment.assessment_version
                    )
                    .order_by(
                        AssessmentCheckInEvidence.created_at.desc()
                    )
                    .first()
                )

            assessment_scan_evidence = (
                AssessmentScanEvidence.query
                .filter_by(
                    assessment_id=
                        current_assessment.id,
                    personnel_id=
                        personnel_id
                )
                .order_by(
                    AssessmentScanEvidence.created_at.desc()
                )
                .first()
            )

            if assessment_scan_evidence is None:

                assessment_scan_evidence = (
                    AssessmentScanEvidence.query
                    .filter_by(
                        personnel_id=
                            personnel_id,
                        assessment_version=
                            current_assessment.assessment_version
                    )
                    .order_by(
                        AssessmentScanEvidence.created_at.desc()
                    )
                    .first()
                )

        return render_template(
            "welfare_case_review.html",
            current_assessment=current_assessment,
            selected_personnel_id=personnel_id,
            personnel=personnel,
            latest_checkin=latest_checkin,
            assessment_checkin_evidence=
                assessment_checkin_evidence,
            latest_scan=latest_scan,
            assessment_scan_evidence=
                assessment_scan_evidence
        )


    # --------------------------------------------------
    # WELFARE OFFICER PRIORITY REVIEW
    # --------------------------------------------------

    @app.route(
        "/welfare-officer-priority-review"
    )
    def welfare_officer_priority_review():

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return redirect(
                url_for(
                    "welfare_officer_login"
                )
            )

        priority_assessments = (
            WelfareAssessment.query
            .filter(
                WelfareAssessment.is_current == True,
                WelfareAssessment.attention_level.in_(
                    [
                        "Moderate",
                        "Elevated"
                    ]
                )
            )
            .order_by(
                WelfareAssessment.assessment_time.desc()
            )
            .all()
        )

        level_order = {
            "Elevated": 0,
            "Moderate": 1
        }

        priority_assessments.sort(
            key=lambda item: (
                level_order.get(
                    item.attention_level,
                    9
                ),
                -(
                    item.assessment_time.timestamp()
                    if item.assessment_time
                    else 0
                )
            )
        )

        personnel_ids = [
            item.personnel_id
            for item in priority_assessments
        ]

        personnel_map = {}

        if personnel_ids:

            personnel_records = (
                Personnel.query
                .filter(
                    Personnel.personnel_id.in_(
                        personnel_ids
                    )
                )
                .all()
            )

            personnel_map = {
                item.personnel_id: item
                for item in personnel_records
            }

        priority_cases = []

        for position, assessment in enumerate(
            priority_assessments,
            start=1
        ):

            priority_cases.append({
                "priority_number":
                    position,

                "assessment":
                    assessment,

                "personnel":
                    personnel_map.get(
                        assessment.personnel_id
                    )
            })

        priority_counts = {
            "total":
                len(priority_cases),

            "elevated":
                sum(
                    1
                    for item in priority_assessments
                    if item.attention_level
                    == "Elevated"
                ),

            "moderate":
                sum(
                    1
                    for item in priority_assessments
                    if item.attention_level
                    == "Moderate"
                )
        }

        return render_template(
            "welfare_priority_review.html",
            priority_cases=priority_cases,
            priority_counts=priority_counts
        )


    # --------------------------------------------------
    # WELFARE OFFICER OPERATIONAL REVIEW
    # --------------------------------------------------

    @app.route(
        "/welfare-officer-operational-review"
    )
    def welfare_officer_operational_review():

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return redirect(
                url_for(
                    "welfare_officer_login"
                )
            )

        personnel_id = (
            request.args.get(
                "personnel_id"
            )
            or ""
        ).strip()

        if not personnel_id:

            personnel_records = (
                Personnel.query
                .order_by(
                    Personnel.personnel_id.asc()
                )
                .all()
            )

            return render_template(
                "welfare_operational_review.html",
                personnel_records=personnel_records,
                selection_error=None,
                selector_mode=True
            )

        selected_personnel = (
            Personnel.query
            .filter_by(
                personnel_id=personnel_id
            )
            .first()
        )

        if selected_personnel is None:

            personnel_records = (
                Personnel.query
                .order_by(
                    Personnel.personnel_id.asc()
                )
                .all()
            )

            return render_template(
                "welfare_operational_review.html",
                personnel_records=personnel_records,
                selection_error=
                    "Personnel ID was not found. Select a valid personnel record.",
                selector_mode=True
            )

        period_end = datetime.now().date()
        period_start = (
            period_end
            - timedelta(days=6)
        )

        operational_records = (
            OperationalRecord.query
            .filter(
                OperationalRecord.personnel_id
                == personnel_id,
                OperationalRecord.is_valid
                == True,
                OperationalRecord.record_date
                >= period_start,
                OperationalRecord.record_date
                <= period_end
            )
            .order_by(
                OperationalRecord.record_date.desc()
            )
            .all()
        )

        current_assessment = (
            WelfareAssessment.query
            .filter_by(
                personnel_id=personnel_id,
                is_current=True
            )
            .order_by(
                WelfareAssessment.assessment_time.desc()
            )
            .first()
        )

        worked_records = [
            item
            for item in operational_records
            if (
                item.duty_hours is not None
                and item.duty_hours > 0
            )
        ]

        _d_vals = [i.duty_hours for i in operational_records if i.duty_hours is not None]
        total_duty_hours = round(sum(_d_vals), 1) if _d_vals else None

        rest_values = [
            item.rest_hours
            for item in operational_records
            if item.rest_hours is not None
        ]

        average_rest_hours = (
            round(
                sum(rest_values)
                / len(rest_values),
                1
            )
            if rest_values
            else None
        )

        _tr_vals = [i.training_hours for i in operational_records if i.training_hours is not None]
        total_training_hours = round(sum(_tr_vals), 1) if _tr_vals else None

        recovery_day_count = sum(
            1
            for item in operational_records
            if (
                item.duty_hours is not None
                and item.duty_hours == 0
            )
        )

        leave_cancellations = sum(
            1
            for item in operational_records
            if (
                item.leave_status
                and "cancel" in item.leave_status.lower()
            )
        )

        latest_record = (
            operational_records[0]
            if operational_records
            else None
        )

        operational_summary = {
            "personnel_id":
                personnel_id,

            "period_start":
                period_start,

            "period_end":
                period_end,

            "record_count":
                len(operational_records),

            "worked_day_count":
                len(worked_records),

            "total_duty_hours":
                total_duty_hours,

            "average_rest_hours":
                average_rest_hours,

            "total_training_hours":
                total_training_hours,

            "recovery_day_count":
                recovery_day_count,

            "leave_cancellations":
                leave_cancellations,

            "current_context":
                (
                    latest_record.context_type
                    if latest_record
                    else None
                ),

            "deployment_active":
                (
                    latest_record.deployment_active
                    if latest_record
                    else False
                ),

            "data_status":
                (
                    "Available"
                    if operational_records
                    else "Insufficient Data"
                )
        }

        return render_template(
            "welfare_operational_review.html",
            operational_records=operational_records,
            operational_summary=operational_summary,
            current_assessment=current_assessment,
            selector_mode=False
        )


    # --------------------------------------------------
    # WELFARE OFFICER INTERVENTION & RECOMMENDATIONS
    # --------------------------------------------------

    @app.route(
        "/welfare-officer-intervention-recommendations"
    )
    def welfare_officer_intervention_recommendations():

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return redirect(
                url_for(
                    "welfare_officer_login"
                )
            )

        personnel_id = (
            request.args.get("personnel_id")
            or ""
        ).strip()

        if not personnel_id:

            return render_template(
                "welfare_intervention_recommendations.html",
                selected_personnel_id=None,
                selected_personnel=None,
                selection_error=None
            )

        selected_personnel = (
            Personnel.query
            .filter_by(
                personnel_id=personnel_id
            )
            .first()
        )

        if selected_personnel is None:

            return render_template(
                "welfare_intervention_recommendations.html",
                selected_personnel_id=None,
                selected_personnel=None,
                selection_error=
                    "Personnel ID was not found. Search and select a valid personnel record."
            )

        current_assessment = (
            WelfareAssessment.query
            .filter_by(
                personnel_id=personnel_id,
                is_current=True
            )
            .order_by(
                WelfareAssessment.assessment_time.desc()
            )
            .first()
        )

        today = datetime.now().date()
        projection_end = (
            today
            + timedelta(days=3)
        )

        upcoming_records = (
            OperationalRecord.query
            .filter(
                OperationalRecord.personnel_id
                == personnel_id,
                OperationalRecord.is_valid
                == True,
                OperationalRecord.record_date
                >= today,
                OperationalRecord.record_date
                <= projection_end
            )
            .order_by(
                OperationalRecord.record_date.asc()
            )
            .all()
        )

        current_window_start = (
            today
            - timedelta(days=6)
        )

        current_records = (
            OperationalRecord.query
            .filter(
                OperationalRecord.personnel_id
                == personnel_id,
                OperationalRecord.is_valid
                == True,
                OperationalRecord.record_date
                >= current_window_start,
                OperationalRecord.record_date
                <= today
            )
            .order_by(
                OperationalRecord.record_date.asc()
            )
            .all()
        )

        _c_d_vals = [i.duty_hours for i in current_records if i.duty_hours is not None]
        current_total_duty = round(sum(_c_d_vals), 1) if _c_d_vals else None

        _c_t_vals = [i.training_hours for i in current_records if i.training_hours is not None]
        current_training_hours = round(sum(_c_t_vals), 1) if _c_t_vals else None

        recovery_day_count = sum(
            1
            for item in current_records
            if (
                item.duty_hours is not None
                and item.duty_hours == 0
            )
        )

        latest_welfare_action = (
            WelfareAction.query
            .filter_by(
                personnel_id=personnel_id
            )
            .order_by(
                WelfareAction.created_at.desc()
            )
            .first()
        )

        source_assessment = None

        if latest_welfare_action:

            if latest_welfare_action.source_assessment_id:

                source_assessment = (
                    WelfareAssessment.query
                    .filter_by(
                        id=latest_welfare_action.source_assessment_id,
                        personnel_id=personnel_id
                    )
                    .first()
                )

            if (
                source_assessment is None
                and latest_welfare_action.source_assessment_version
            ):

                source_assessment = (
                    WelfareAssessment.query
                    .filter_by(
                        personnel_id=personnel_id,
                        assessment_version=
                            latest_welfare_action.source_assessment_version
                    )
                    .first()
                )

        follow_up_comparison = {
            "available": False,
            "status":
                "Awaiting Follow-Up",
            "outcome":
                "Not Available",
            "source_version":
                (
                    source_assessment.assessment_version
                    if source_assessment
                    else None
                ),
            "source_attention":
                (
                    source_assessment.attention_level
                    if source_assessment
                    else None
                ),
            "current_version":
                (
                    current_assessment.assessment_version
                    if current_assessment
                    else None
                ),
            "current_attention":
                (
                    current_assessment.attention_level
                    if current_assessment
                    else None
                ),
            "action_type":
                (
                    latest_welfare_action.action_type
                    if latest_welfare_action
                    else None
                ),
            "action_status":
                (
                    latest_welfare_action.status
                    if latest_welfare_action
                    else None
                ),
            "message":
                (
                    "No welfare action has been recorded for this personnel yet."
                    if latest_welfare_action is None
                    else
                    "Move the welfare action to Follow-Up and create the next assessment to compare outcomes."
                )
        }

        comparable_levels = {
            "Low": 1,
            "Moderate": 2,
            "Elevated": 3
        }

        if (
            latest_welfare_action
            and latest_welfare_action.status == "Follow-Up"
            and source_assessment
            and current_assessment
            and current_assessment.assessment_version
                > source_assessment.assessment_version
        ):

            source_rank = comparable_levels.get(
                source_assessment.attention_level
            )

            current_rank = comparable_levels.get(
                current_assessment.attention_level
            )

            if (
                source_rank is not None
                and current_rank is not None
            ):

                follow_up_comparison["available"] = True

                if current_rank < source_rank:

                    follow_up_comparison["status"] = (
                        "Lower Welfare Attention"
                    )

                    follow_up_comparison["outcome"] = (
                        "Improved"
                    )

                    follow_up_comparison["message"] = (
                        "The latest welfare-attention category is lower than "
                        "the assessment version linked to the welfare action. "
                        "This indicates improvement in the current stored "
                        "operational pattern, not a medical outcome."
                    )

                elif current_rank == source_rank:

                    follow_up_comparison["status"] = (
                        "No Category Change"
                    )

                    follow_up_comparison["outcome"] = (
                        "Stable"
                    )

                    follow_up_comparison["message"] = (
                        "The latest welfare-attention category is unchanged. "
                        "Continue human review using the current evidence and "
                        "operational context."
                    )

                else:

                    follow_up_comparison["status"] = (
                        "Higher Welfare Attention"
                    )

                    follow_up_comparison["outcome"] = (
                        "Needs Further Review"
                    )

                    follow_up_comparison["message"] = (
                        "The latest welfare-attention category is higher than "
                        "the assessment version linked to the welfare action. "
                        "Further authorized human review is recommended."
                    )

            else:

                follow_up_comparison["status"] = (
                    "Not Comparable"
                )

                follow_up_comparison["outcome"] = (
                    "Insufficient Data"
                )

                follow_up_comparison["message"] = (
                    "One of the assessment versions is Insufficient Data or "
                    "otherwise not comparable. VeerPulse does not invent an "
                    "improvement result."
                )

        projection_summary = {
            "personnel_id":
                personnel_id,

            "projection_start":
                today,

            "projection_end":
                projection_end,

            "upcoming_record_count":
                len(upcoming_records),

            "current_total_duty":
                current_total_duty,

            "current_training_hours":
                current_training_hours,

            "current_recovery_days":
                recovery_day_count,

            "data_status":
                (
                    "Available"
                    if upcoming_records
                    else "No Upcoming Roster Data"
                )
        }

        return render_template(
            "welfare_intervention_recommendations.html",
            current_assessment=current_assessment,
            upcoming_records=upcoming_records,
            projection_summary=projection_summary,
            latest_welfare_action=latest_welfare_action,
            source_assessment=source_assessment,
            follow_up_comparison=follow_up_comparison,
            selected_personnel=selected_personnel,
            selected_personnel_id=personnel_id,
            selection_error=None
        )


    # --------------------------------------------------
    # HRMS / DATA INTEGRATION STATUS HELPER
    # --------------------------------------------------

    def _build_hrms_integration_status():

        source_type = (
            "HRMS-Style Prototype Import"
        )

        source_records = (
            OperationalRecord.query
            .filter_by(
                source_type=source_type,
                is_valid=True
            )
            .order_by(
                OperationalRecord.record_date.desc()
            )
            .all()
        )

        source_personnel_ids = {
            item.personnel_id
            for item in source_records
        }

        record_dates = [
            item.record_date
            for item in source_records
            if item.record_date
        ]

        import_events = (
            AuditEvent.query
            .filter_by(
                event_type=
                    "HRMS-Style Operational Import"
            )
            .order_by(
                AuditEvent.created_at.desc()
            )
            .limit(5)
            .all()
        )

        latest_import = (
            import_events[0]
            if import_events
            else None
        )

        integration_summary = {
            "source_type":
                source_type,

            "valid_record_count":
                len(source_records),

            "personnel_covered":
                len(
                    source_personnel_ids
                ),

            "import_event_count":
                (
                    AuditEvent.query
                    .filter_by(
                        event_type=
                            "HRMS-Style Operational Import"
                    )
                    .count()
                ),

            "latest_import_time":
                (
                    latest_import.created_at
                    if latest_import
                    else None
                ),

            "latest_import_actor":
                (
                    latest_import.actor_id
                    if latest_import
                    else None
                ),

            "latest_import_summary":
                (
                    latest_import.event_summary
                    if latest_import
                    else "No HRMS-style import has been recorded yet."
                ),

            "earliest_record_date":
                (
                    min(record_dates)
                    if record_dates
                    else None
                ),

            "latest_record_date":
                (
                    max(record_dates)
                    if record_dates
                    else None
                )
        }

        return (
            integration_summary,
            import_events
        )


    @app.route(
        "/api/welfare-officer/hrms/template",
        methods=["GET"]
    )
    def download_hrms_operational_template():

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return redirect(
                url_for(
                    "welfare_officer_login"
                )
            )

        csv_text = (
            "personnel_id,record_date,context_type,"
            "source_record_id,duty_hours,rest_hours,"
            "deployment_active,training_hours,leave_status\n"
            "VP0001,2026-09-15,Station Duty,"
            "HRMS-DEMO-001,8,8,false,0,\n"
        )

        buffer = BytesIO(
            csv_text.encode(
                "utf-8"
            )
        )

        buffer.seek(0)

        return send_file(
            buffer,
            mimetype="text/csv",
            as_attachment=True,
            download_name=
                "veerpulse_hrms_operational_template.csv"
        )


    # --------------------------------------------------
    # WELFARE OFFICER HRMS / DATA INTEGRATION
    # --------------------------------------------------

    @app.route(
        "/welfare-officer-hrms-integration"
    )
    def welfare_officer_hrms_integration():

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return redirect(
                url_for(
                    "welfare_officer_login"
                )
            )

        (
            integration_summary,
            recent_import_events
        ) = (
            _build_hrms_integration_status()
        )

        return render_template(
            "welfare_hrms_integration.html",
            integration_summary=
                integration_summary,
            recent_import_events=
                recent_import_events
        )


    # --------------------------------------------------
    # WELFARE OFFICER REPORTS
    # --------------------------------------------------

    @app.route(
        "/welfare-officer-reports"
    )
    def welfare_officer_reports():

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return redirect(
                url_for(
                    "welfare_officer_login"
                )
            )

        personnel_id = (
            request.args.get("personnel_id")
            or ""
        ).strip()

        if not personnel_id:

            return render_template(
                "welfare_officer_reports.html",
                selected_personnel_id=None,
                selected_personnel=None,
                selection_error=None
            )

        selected_personnel = (
            Personnel.query
            .filter_by(
                personnel_id=personnel_id
            )
            .first()
        )

        if selected_personnel is None:

            return render_template(
                "welfare_officer_reports.html",
                selected_personnel_id=None,
                selected_personnel=None,
                selection_error=
                    "Personnel ID was not found. Search and select a valid personnel record."
            )

        current_assessment = (
            WelfareAssessment.query
            .filter_by(
                personnel_id=personnel_id,
                is_current=True
            )
            .order_by(
                WelfareAssessment.assessment_time.desc()
            )
            .first()
        )

        assessment_history = (
            WelfareAssessment.query
            .filter_by(
                personnel_id=personnel_id
            )
            .order_by(
                WelfareAssessment.assessment_version.desc()
            )
            .all()
        )

        welfare_actions = (
            WelfareAction.query
            .filter_by(
                personnel_id=personnel_id
            )
            .order_by(
                WelfareAction.created_at.desc()
            )
            .all()
        )

        operational_record_count = (
            OperationalRecord.query
            .filter_by(
                personnel_id=personnel_id,
                is_valid=True
            )
            .count()
        )

        replay_snapshot_count = (
            AssessmentEvidenceSnapshot.query
            .filter_by(
                personnel_id=personnel_id
            )
            .count()
        )

        action_status_counts = {
            "Planned": 0,
            "Approved": 0,
            "Delivered": 0,
            "Interrupted": 0,
            "Unable to Verify": 0,
            "Follow-Up": 0
        }

        for action in welfare_actions:

            if (
                action.status
                in action_status_counts
            ):

                action_status_counts[
                    action.status
                ] += 1

        # --------------------------------------------------
        # SAME-PERSON JOURNEY TIMELINE
        # --------------------------------------------------

        report_timeline = []

        for assessment in assessment_history:

            report_timeline.append({
                "event_time":
                    assessment.assessment_time,

                "event_type":
                    "Assessment",

                "title":
                    (
                        "Assessment v"
                        + str(
                            assessment.assessment_version
                        )
                        + " — "
                        + assessment.attention_level
                    ),

                "detail":
                    (
                        assessment.explanation_summary
                        or (
                            "Evidence: "
                            + assessment.evidence_status
                        )
                    ),

                "version":
                    assessment.assessment_version,

                "status":
                    assessment.attention_level
            })

        for action in welfare_actions:

            report_timeline.append({
                "event_time":
                    action.created_at,

                "event_type":
                    "Welfare Action",

                "title":
                    (
                        "Welfare Action — "
                        + action.action_type
                    ),

                "detail":
                    (
                        "Linked to assessment v"
                        + str(
                            action.source_assessment_version
                            or "—"
                        )
                        + ". Current delivery state: "
                        + action.status
                        + "."
                    ),

                "version":
                    action.source_assessment_version,

                "status":
                    action.status
            })

            if (
                action.updated_at
                and action.created_at
                and action.updated_at
                    > action.created_at
            ):

                report_timeline.append({
                    "event_time":
                        action.updated_at,

                    "event_type":
                        "Action Update",

                    "title":
                        (
                            "Action Status — "
                            + action.status
                        ),

                    "detail":
                        (
                            action.verification_note
                            or action.officer_note
                            or (
                                action.action_type
                                + " workflow updated."
                            )
                        ),

                    "version":
                        action.source_assessment_version,

                    "status":
                        action.status
                })

        report_timeline.sort(
            key=lambda item: (
                item["event_time"]
                or datetime.min
            )
        )

        report_period_start = (
            min(
                (item["event_time"] for item in report_timeline if item.get("event_time")),
                default=None
            )
        )

        report_period_end = (
            max(
                (item["event_time"] for item in report_timeline if item.get("event_time")),
                default=None
            )
        )

        report_summary = {
            "personnel_id":
                personnel_id,

            "period_start":
                report_period_start,

            "period_end":
                report_period_end,

            "assessment_count":
                len(assessment_history),

            "operational_record_count":
                operational_record_count,

            "welfare_action_count":
                len(welfare_actions),

            "replay_snapshot_count":
                replay_snapshot_count,

            "timeline_event_count":
                len(report_timeline),

            "current_attention":
                (
                    current_assessment.attention_level
                    if current_assessment
                    else "Insufficient Data"
                ),

            "current_evidence_status":
                (
                    current_assessment.evidence_status
                    if current_assessment
                    else "Insufficient Data"
                ),

            "current_assessment_version":
                (
                    current_assessment.assessment_version
                    if current_assessment
                    else None
                )
        }

        return render_template(
            "welfare_officer_reports.html",
            selected_personnel_id=personnel_id,
            selected_personnel=selected_personnel,
            selection_error=None,
            current_assessment=current_assessment,
            assessment_history=assessment_history,
            welfare_actions=welfare_actions,
            action_status_counts=action_status_counts,
            report_summary=report_summary,
            report_timeline=report_timeline
        )


    # --------------------------------------------------
    # WELFARE OFFICER NOTIFICATIONS
    # --------------------------------------------------

    @app.route(
        "/welfare-officer-notifications"
    )
    def welfare_officer_notifications():

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return redirect(
                url_for(
                    "welfare_officer_login"
                )
            )

        officer_id = (
            session.get("officer_id")
            or "welfare_officer"
        )

        now = datetime.now()

        def relative_time(value):

            if value is None:
                return "Time unavailable"

            delta = now - value

            seconds = max(
                0,
                int(delta.total_seconds())
            )

            if seconds < 60:
                return "Just now"

            minutes = seconds // 60

            if minutes < 60:
                return (
                    str(minutes)
                    + (
                        " minute ago"
                        if minutes == 1
                        else " minutes ago"
                    )
                )

            hours = minutes // 60

            if hours < 24:
                return (
                    str(hours)
                    + (
                        " hour ago"
                        if hours == 1
                        else " hours ago"
                    )
                )

            days = hours // 24

            if days == 1:
                return "Yesterday"

            if days < 7:
                return (
                    str(days)
                    + " days ago"
                )

            return value.strftime(
                "%d %b %Y"
            )

        notification_items = []

        # ----------------------------------------------
        # Welfare assessment alerts
        # ----------------------------------------------

        alert_assessments = (
            WelfareAssessment.query
            .filter(
                WelfareAssessment.attention_level.in_(
                    [
                        "Moderate",
                        "Elevated"
                    ]
                )
            )
            .order_by(
                WelfareAssessment.assessment_time.desc()
            )
            .limit(50)
            .all()
        )

        for assessment in alert_assessments:

            notification_items.append({
                "event_key":
                    (
                        "assessment:"
                        + str(assessment.id)
                    ),

                "event_time":
                    assessment.assessment_time,

                "time_label":
                    relative_time(
                        assessment.assessment_time
                    ),

                "category":
                    "welfare-alert",

                "title":
                    (
                        assessment.attention_level
                        + " Welfare Alert"
                    ),

                "message":
                    (
                        "Personnel "
                        + assessment.personnel_id
                        + " has a "
                        + assessment.attention_level
                        + " welfare assessment requiring authorized human review."
                    ),

                "tags":
                    [
                        "Welfare Alert",
                        (
                            "Assessment v"
                            + str(
                                assessment.assessment_version
                            )
                        )
                    ],

                "action_label":
                    "Open Case",

                "action_url":
                    url_for(
                        "welfare_officer_case_review",
                        personnel_id=
                            assessment.personnel_id
                    )
            })

        # ----------------------------------------------
        # Voluntary support requests
        # ----------------------------------------------

        recent_support_requests = (
            SupportRequest.query
            .order_by(
                SupportRequest.created_at.desc()
            )
            .limit(50)
            .all()
        )

        for support_request in recent_support_requests:

            notification_items.append({
                "event_key":
                    (
                        "support:"
                        + str(support_request.id)
                        + ":created"
                    ),

                "event_time":
                    support_request.created_at,

                "time_label":
                    relative_time(
                        support_request.created_at
                    ),

                "category":
                    "support-request",

                "title":
                    "New Support Request",

                "message":
                    (
                        "Personnel "
                        + support_request.personnel_id
                        + " submitted a voluntary "
                        + support_request.support_category
                        + " support request."
                    ),

                "tags":
                    [
                        "Support Request",
                        support_request.priority
                    ],

                "action_label":
                    "View Request",

                "action_url":
                    url_for(
                        "welfare_officer_support_requests"
                    )
            })

            if (
                support_request.updated_at
                and support_request.created_at
                and support_request.updated_at
                    > support_request.created_at
            ):

                notification_items.append({
                    "event_key":
                        (
                            "support:"
                            + str(support_request.id)
                            + ":status:"
                            + str(
                                support_request.status
                            )
                        ),

                    "event_time":
                        support_request.updated_at,

                    "time_label":
                        relative_time(
                            support_request.updated_at
                        ),

                    "category":
                        "support-request",

                    "title":
                        "Support Request Updated",

                    "message":
                        (
                            "Personnel "
                            + support_request.personnel_id
                            + " support request is now "
                            + support_request.status
                            + "."
                        ),

                    "tags":
                        [
                            "Support Request",
                            support_request.status
                        ],

                    "action_label":
                        "View Request",

                    "action_url":
                        url_for(
                            "welfare_officer_support_requests"
                        )
                })

        # ----------------------------------------------
        # Welfare-action / follow-up case updates
        # ----------------------------------------------

        recent_actions = (
            WelfareAction.query
            .order_by(
                WelfareAction.updated_at.desc()
            )
            .limit(50)
            .all()
        )

        for action in recent_actions:

            notification_items.append({
                "event_key":
                    (
                        "welfare-action:"
                        + str(action.id)
                        + ":created"
                    ),

                "event_time":
                    action.created_at,

                "time_label":
                    relative_time(
                        action.created_at
                    ),

                "category":
                    "case-update",

                "title":
                    "Welfare Action Recorded",

                "message":
                    (
                        action.action_type
                        + " was recorded for personnel "
                        + action.personnel_id
                        + "."
                    ),

                "tags":
                    [
                        "Case Update",
                        action.status
                    ],

                "action_label":
                    "Open Intervention",

                "action_url":
                    url_for(
                        "welfare_officer_intervention_recommendations",
                        personnel_id=
                            action.personnel_id
                    )
            })

            if (
                action.updated_at
                and action.created_at
                and action.updated_at
                    > action.created_at
            ):

                notification_items.append({
                    "event_key":
                        (
                            "welfare-action:"
                            + str(action.id)
                            + ":status:"
                            + action.status
                        ),

                    "event_time":
                        action.updated_at,

                    "time_label":
                        relative_time(
                            action.updated_at
                        ),

                    "category":
                        "case-update",

                    "title":
                        (
                            "Welfare Action — "
                            + action.status
                        ),

                    "message":
                        (
                            action.action_type
                            + " for personnel "
                            + action.personnel_id
                            + " is now "
                            + action.status
                            + "."
                        ),

                    "tags":
                        [
                            "Case Update",
                            action.status
                        ],

                    "action_label":
                        (
                            "Open Follow-Up"
                            if action.status == "Follow-Up"
                            else "Open Intervention"
                        ),

                    "action_url":
                        url_for(
                            "welfare_officer_intervention_recommendations",
                            personnel_id=
                                action.personnel_id
                        )
                })

        officer_preference = (
            WelfareOfficerPreference.query
            .filter_by(
                officer_id=officer_id
            )
            .first()
        )

        if officer_preference:

            notification_items = [
                item
                for item in notification_items
                if (
                    (
                        item["category"]
                        != "welfare-alert"
                    )
                    or
                    officer_preference.notify_welfare_alerts
                )
                and
                (
                    (
                        item["category"]
                        != "support-request"
                    )
                    or
                    officer_preference.notify_support_requests
                )
                and
                (
                    (
                        item["category"]
                        != "case-update"
                    )
                    or
                    officer_preference.notify_case_updates
                )
            ]

        # Newest events first; keep the visible feed bounded.
        notification_items.sort(
            key=lambda item: (
                item["event_time"]
                or datetime.min
            ),
            reverse=True
        )

        notification_items = (
            notification_items[:100]
        )

        event_keys = [
            item["event_key"]
            for item in notification_items
        ]

        read_keys = set()

        if event_keys:

            read_rows = (
                OfficerNotificationRead.query
                .filter(
                    OfficerNotificationRead.officer_id
                    == officer_id,
                    OfficerNotificationRead.event_key.in_(
                        event_keys
                    )
                )
                .all()
            )

            read_keys = {
                row.event_key
                for row in read_rows
            }

        for item in notification_items:

            item["is_read"] = (
                item["event_key"]
                in read_keys
            )

        unread_count = sum(
            1
            for item in notification_items
            if not item["is_read"]
        )

        current_alert_count = (
            WelfareAssessment.query
            .filter(
                WelfareAssessment.is_current == True,
                WelfareAssessment.attention_level.in_(
                    [
                        "Moderate",
                        "Elevated"
                    ]
                )
            )
            .count()
        )

        open_support_count = (
            SupportRequest.query
            .filter(
                SupportRequest.status.in_(
                    [
                        "Open",
                        "In Review",
                        "In Progress"
                    ]
                )
            )
            .count()
        )

        active_case_update_count = (
            WelfareAction.query
            .filter(
                WelfareAction.status.in_(
                    [
                        "Planned",
                        "Approved",
                        "Delivered",
                        "Follow-Up"
                    ]
                )
            )
            .count()
        )

        notification_summary = {
            "unread":
                unread_count,

            "alerts":
                current_alert_count,

            "support_requests":
                open_support_count,

            "case_updates":
                active_case_update_count
        }

        return render_template(
            "welfare_officer_notifications.html",
            notification_items=
                notification_items,
            notification_summary=
                notification_summary
        )


    @app.route(
        "/api/welfare-officer/notifications/read",
        methods=["POST"]
    )
    def mark_welfare_officer_notification_read():

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return jsonify({
                "success": False,
                "message":
                    "Unauthorized."
            }), 401

        data = (
            request.get_json(
                silent=True
            )
            or {}
        )

        event_key = (
            data.get("event_key")
            or ""
        ).strip()

        if not event_key:

            return jsonify({
                "success": False,
                "message":
                    "Notification event key is required."
            }), 400

        officer_id = (
            session.get("officer_id")
            or ""
        ).strip()

        if not officer_id:

            return jsonify({
                "success": False,
                "message":
                    "Unauthorized."
            }), 401

        existing = (
            OfficerNotificationRead.query
            .filter_by(
                officer_id=officer_id,
                event_key=event_key
            )
            .first()
        )

        if existing is None:

            read_row = (
                OfficerNotificationRead(
                    officer_id=officer_id,
                    event_key=event_key,
                    read_at=datetime.now()
                )
            )

            db.session.add(
                read_row
            )

            db.session.commit()

        return jsonify({
            "success": True,
            "event_key":
                event_key
        })


    # --------------------------------------------------
    # WELFARE OFFICER PRIVACY & AUDIT
    # --------------------------------------------------

    @app.route(
        "/welfare-officer-privacy-audit"
    )
    def welfare_officer_privacy_audit():

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return redirect(
                url_for(
                    "welfare_officer_login"
                )
            )

        selected_personnel_id = (
            request.args.get(
                "personnel_id"
            )
            or ""
        ).strip()

        selected_event_type = (
            request.args.get(
                "event_type"
            )
            or ""
        ).strip()

        audit_query = (
            AuditEvent.query
            .order_by(
                AuditEvent.created_at.desc()
            )
        )

        if selected_personnel_id:

            audit_query = (
                audit_query
                .filter(
                    AuditEvent.personnel_id
                    == selected_personnel_id
                )
            )

        if selected_event_type:

            audit_query = (
                audit_query
                .filter(
                    AuditEvent.event_type
                    == selected_event_type
                )
            )

        audit_events = (
            audit_query
            .limit(100)
            .all()
        )

        total_audit_events = (
            AuditEvent.query.count()
        )

        last_24_hours = (
            datetime.now()
            - timedelta(hours=24)
        )

        recent_24h_count = (
            AuditEvent.query
            .filter(
                AuditEvent.created_at
                >= last_24_hours
            )
            .count()
        )

        distinct_personnel_count = (
            AuditEvent.query
            .with_entities(
                AuditEvent.personnel_id
            )
            .filter(
                AuditEvent.personnel_id
                .isnot(None)
            )
            .distinct()
            .count()
        )

        hrms_import_count = (
            AuditEvent.query
            .filter_by(
                event_type=
                    "HRMS-Style Operational Import"
            )
            .count()
        )

        assessment_recalc_count = (
            AuditEvent.query
            .filter_by(
                event_type=
                    "Welfare Assessment Recalculated"
            )
            .count()
        )

        decision_replay_count = (
            AuditEvent.query
            .filter_by(
                event_type=
                    "Decision Replay Accessed"
            )
            .count()
        )

        welfare_action_update_count = (
            AuditEvent.query
            .filter_by(
                event_type=
                    "Welfare Action Status Updated"
            )
            .count()
        )

        event_type_rows = (
            AuditEvent.query
            .with_entities(
                AuditEvent.event_type
            )
            .distinct()
            .order_by(
                AuditEvent.event_type.asc()
            )
            .all()
        )

        available_event_types = [
            row[0]
            for row in event_type_rows
            if row[0]
        ]

        recent_import_events = (
            AuditEvent.query
            .filter_by(
                event_type=
                    "HRMS-Style Operational Import"
            )
            .order_by(
                AuditEvent.created_at.desc()
            )
            .limit(5)
            .all()
        )

        privacy_audit_summary = {
            "total_events":
                total_audit_events,

            "recent_24h":
                recent_24h_count,

            "personnel_referenced":
                distinct_personnel_count,

            "hrms_imports":
                hrms_import_count,

            "assessment_recalculations":
                assessment_recalc_count,

            "decision_replays":
                decision_replay_count,

            "welfare_action_updates":
                welfare_action_update_count,

            "visible_rows":
                len(
                    audit_events
                )
        }

        return render_template(
            "welfare_officer_privacy_audit.html",

            audit_events=
                audit_events,

            privacy_audit_summary=
                privacy_audit_summary,

            available_event_types=
                available_event_types,

            selected_personnel_id=
                selected_personnel_id,

            selected_event_type=
                selected_event_type,

            recent_import_events=
                recent_import_events
        )


    # --------------------------------------------------
    # WELFARE OFFICER UNIT INSIGHTS
    # --------------------------------------------------

    @app.route(
        "/welfare-officer-unit-insights"
    )
    def welfare_officer_unit_insights():

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return redirect(
                url_for(
                    "welfare_officer_login"
                )
            )

        return render_template(
            "welfare_officer_unit_insights.html"
        )


    # --------------------------------------------------
    # WELFARE OFFICER DUTY UPLOAD
    # --------------------------------------------------

    @app.route(
        "/welfare-officer-duty-upload"
    )
    def welfare_officer_duty_upload():

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return redirect(
                url_for(
                    "welfare_officer_login"
                )
            )

        (
            integration_summary,
            recent_import_events
        ) = (
            _build_hrms_integration_status()
        )

        return render_template(
            "welfare_officer_duty_upload.html",
            integration_summary=
                integration_summary,
            recent_import_events=
                recent_import_events
        )


    # --------------------------------------------------
    # WELFARE OFFICER SETTINGS
    # --------------------------------------------------

    @app.route(
        "/welfare-officer-settings"
    )
    def welfare_officer_settings():

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return redirect(
                url_for(
                    "welfare_officer_login"
                )
            )

        officer_id = session.get(
            "officer_id"
        )

        officer = (
            WelfareOfficer.query
            .filter_by(
                officer_id=officer_id
            )
            .first()
        )

        preference = (
            WelfareOfficerPreference.query
            .filter_by(
                officer_id=officer_id
            )
            .first()
        )

        if preference is None:

            preference = (
                WelfareOfficerPreference(
                    officer_id=officer_id,
                    notify_welfare_alerts=True,
                    notify_support_requests=True,
                    notify_case_updates=True,
                    updated_at=datetime.now()
                )
            )

            db.session.add(
                preference
            )

            db.session.commit()

        return render_template(
            "welfare_officer_settings.html",
            officer=officer,
            preference=preference
        )


    @app.route(
        "/api/welfare-officer/settings/profile",
        methods=["POST"]
    )
    def update_welfare_officer_profile():

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized."
            }), 401

        officer_id = session.get(
            "officer_id"
        )

        officer = (
            WelfareOfficer.query
            .filter_by(
                officer_id=officer_id
            )
            .first()
        )

        if officer is None:

            return jsonify({
                "success": False,
                "message":
                    "Officer account was not found."
            }), 404

        data = (
            request.get_json(
                silent=True
            )
            or {}
        )

        full_name = (
            str(
                data.get(
                    "full_name",
                    ""
                )
            )
            .strip()
        )

        designation = (
            str(
                data.get(
                    "designation",
                    ""
                )
            )
            .strip()
        )

        assigned_unit = (
            str(
                data.get(
                    "assigned_unit",
                    ""
                )
            )
            .strip()
        )

        email = (
            str(
                data.get(
                    "email",
                    ""
                )
            )
            .strip()
        )

        if not full_name:

            return jsonify({
                "success": False,
                "message":
                    "Officer name is required."
            }), 400

        if len(full_name) > 120:

            return jsonify({
                "success": False,
                "message":
                    "Officer name is too long."
            }), 400

        if len(designation) > 120:

            return jsonify({
                "success": False,
                "message":
                    "Designation is too long."
            }), 400

        if len(assigned_unit) > 120:

            return jsonify({
                "success": False,
                "message":
                    "Assigned unit is too long."
            }), 400

        if len(email) > 160:

            return jsonify({
                "success": False,
                "message":
                    "Email is too long."
            }), 400

        officer.full_name = (
            full_name
        )

        officer.designation = (
            designation
            or None
        )

        officer.assigned_unit = (
            assigned_unit
            or None
        )

        officer.email = (
            email
            or None
        )

        officer.updated_at = (
            datetime.now()
        )

        session[
            "officer_name"
        ] = officer.full_name

        db.session.add(
            AuditEvent(
                actor_role=
                    "welfare_officer",
                actor_id=
                    officer_id,
                event_type=
                    "Welfare Officer Profile Updated",
                personnel_id=
                    None,
                resource_type=
                    "WelfareOfficer",
                resource_id=
                    str(officer.id),
                event_summary=
                    "Welfare Officer account profile was updated.",
                source_type=
                    "VeerPulse Application",
                created_at=
                    datetime.now()
            )
        )

        db.session.commit()

        return jsonify({
            "success": True,
            "message":
                "Officer profile saved."
        })


    @app.route(
        "/api/welfare-officer/settings/preferences",
        methods=["POST"]
    )
    def update_welfare_officer_preferences():

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized."
            }), 401

        officer_id = session.get(
            "officer_id"
        )

        data = (
            request.get_json(
                silent=True
            )
            or {}
        )

        preference = (
            WelfareOfficerPreference.query
            .filter_by(
                officer_id=officer_id
            )
            .first()
        )

        if preference is None:

            preference = (
                WelfareOfficerPreference(
                    officer_id=officer_id
                )
            )

            db.session.add(
                preference
            )

        preference.notify_welfare_alerts = bool(
            data.get(
                "notify_welfare_alerts",
                True
            )
        )

        preference.notify_support_requests = bool(
            data.get(
                "notify_support_requests",
                True
            )
        )

        preference.notify_case_updates = bool(
            data.get(
                "notify_case_updates",
                True
            )
        )

        preference.updated_at = (
            datetime.now()
        )

        db.session.add(
            AuditEvent(
                actor_role=
                    "welfare_officer",
                actor_id=
                    officer_id,
                event_type=
                    "Welfare Officer Preferences Updated",
                personnel_id=
                    None,
                resource_type=
                    "WelfareOfficerPreference",
                resource_id=
                    None,
                event_summary=
                    "Notification preferences were updated.",
                source_type=
                    "VeerPulse Application",
                created_at=
                    datetime.now()
            )
        )

        db.session.commit()

        return jsonify({
            "success": True,
            "message":
                "Notification preferences saved."
        })


    @app.route(
        "/api/welfare-officer/settings/password",
        methods=["POST"]
    )
    def change_welfare_officer_password():

        if (
            session.get("role")
            != "welfare_officer"
        ):

            return jsonify({
                "success": False,
                "message": "Unauthorized."
            }), 401

        officer_id = session.get(
            "officer_id"
        )

        officer = (
            WelfareOfficer.query
            .filter_by(
                officer_id=officer_id
            )
            .first()
        )

        if officer is None:

            return jsonify({
                "success": False,
                "message":
                    "Officer account was not found."
            }), 404

        data = (
            request.get_json(
                silent=True
            )
            or {}
        )

        current_password = (
            str(
                data.get(
                    "current_password",
                    ""
                )
            )
        )

        new_password = (
            str(
                data.get(
                    "new_password",
                    ""
                )
            )
        )

        confirm_password = (
            str(
                data.get(
                    "confirm_password",
                    ""
                )
            )
        )

        if not check_password_hash(
            officer.password_hash,
            current_password
        ):

            return jsonify({
                "success": False,
                "message":
                    "Current password is incorrect."
            }), 400

        if len(new_password) < 8:

            return jsonify({
                "success": False,
                "message":
                    "New password must contain at least 8 characters."
            }), 400

        if new_password != confirm_password:

            return jsonify({
                "success": False,
                "message":
                    "New password and confirmation do not match."
            }), 400

        if check_password_hash(
            officer.password_hash,
            new_password
        ):

            return jsonify({
                "success": False,
                "message":
                    "New password must be different from the current password."
            }), 400

        officer.password_hash = (
            generate_password_hash(
                new_password
            )
        )

        officer.updated_at = (
            datetime.now()
        )

        db.session.add(
            AuditEvent(
                actor_role=
                    "welfare_officer",
                actor_id=
                    officer_id,
                event_type=
                    "Welfare Officer Password Changed",
                personnel_id=
                    None,
                resource_type=
                    "WelfareOfficer",
                resource_id=
                    str(officer.id),
                event_summary=
                    "Welfare Officer password was changed.",
                source_type=
                    "VeerPulse Application",
                created_at=
                    datetime.now()
            )
        )

        db.session.commit()

        return jsonify({
            "success": True,
            "message":
                "Password changed successfully."
        })


    # --------------------------------------------------
    # LOGOUT
    # --------------------------------------------------

    @app.route(
        "/logout"
    )
    def logout():

        session.clear()

        return redirect(
            url_for("home")
        )


    return app